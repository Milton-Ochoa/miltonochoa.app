"""Lectura del Excel de estudiantes (openpyxl, sin pandas).

Réplica del patrón de `programacion/monitores/views.py`: encabezados tolerantes
a orden/mayúsculas (y tildes), `load_workbook(read_only=True, data_only=True)`,
`iter_rows(values_only=True)`. Las columnas dependen del tipo de plantilla:
SIMULACRO/PENSAR esperan `Nombres`, `Grado`, `Usuario`; MP añade `Código`,
`Año` y `Estudiante` (los trae el export grdGeneral del portal). Devuelve
`list[dict]`; omite las filas sin `Nombres`. Si faltan columnas →
`ExcelInvalido`.
"""
import unicodedata

from openpyxl import load_workbook

# Encabezado normalizado (lower, sin tildes) → clave de salida.
_COLUMNAS_BASE = {
    'nombres': 'nombre',
    'grado':   'grado',
    'usuario': 'usuario',
}
_COLUMNAS_MP = {
    **_COLUMNAS_BASE,
    'codigo':     'codigo_colegio',
    'ano':        'anio',
    'estudiante': 'codigo_estudiante',
}
COLUMNAS_POR_TIPO = {
    'SIMULACRO': _COLUMNAS_BASE,
    'PENSAR':    _COLUMNAS_BASE,
    'MP':        _COLUMNAS_MP,
}
# Encabezado normalizado → nombre "bonito" para los mensajes de error.
_DISPLAY = {
    'nombres': 'Nombres', 'grado': 'Grado', 'usuario': 'Usuario',
    'codigo': 'Código', 'ano': 'Año', 'estudiante': 'Estudiante',
}


class ExcelInvalido(Exception):
    """El Excel no se pudo leer o le faltan columnas obligatorias."""


def _norm(valor):
    """None→'', número entero→int, y str().strip()."""
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _norm_encabezado(valor):
    """Normaliza un encabezado para compararlo: lower + sin tildes ('Código'→
    'codigo', 'Año'→'ano'). Tolera que el export venga con o sin acentos."""
    texto = _norm(valor).lower()
    return (unicodedata.normalize('NFKD', texto)
            .encode('ascii', 'ignore').decode())


def _entero_texto(valor):
    """Como `_norm` pero además quita el sufijo '.0' que openpyxl deja al leer
    un entero como float (5 → 5.0 → '5'). Para grado/usuario/códigos numéricos."""
    texto = _norm(valor)
    if texto.endswith('.0') and texto[:-2].isdigit():
        return texto[:-2]
    return texto


def leer_estudiantes(archivo, tipo='SIMULACRO'):
    """Lee el .xlsx subido y devuelve `list[dict]` de estudiantes.

    `archivo` es un file-like (p. ej. `request.FILES['excel']`); `tipo` decide
    las columnas obligatorias (ver `COLUMNAS_POR_TIPO`). Lanza `ExcelInvalido`
    si el archivo no se puede abrir o faltan columnas.
    """
    columnas = COLUMNAS_POR_TIPO[tipo]
    try:
        wb = load_workbook(archivo, read_only=True, data_only=True)
        ws = wb.active
        filas = ws.iter_rows(values_only=True)
        encabezados = next(filas, None)
    except Exception:
        raise ExcelInvalido('No se pudo leer el Excel: ¿está dañado o vacío?')

    if not encabezados:
        raise ExcelInvalido('El Excel está vacío.')

    # Mapa encabezado-normalizado → índice de columna (tolera orden/mayúsculas/tildes).
    indices = {}
    for i, cab in enumerate(encabezados):
        clave = _norm_encabezado(cab)
        if clave in columnas:
            indices[columnas[clave]] = i

    faltan = [_DISPLAY[nom] for nom, clave in columnas.items() if clave not in indices]
    if faltan:
        raise ExcelInvalido(
            'Al Excel le faltan columnas obligatorias: '
            + ', '.join(sorted(faltan))
            + '. Se esperan: ' + ', '.join(_DISPLAY[n] for n in columnas) + '.')

    def celda(fila, clave):
        idx = indices[clave]
        return fila[idx] if idx < len(fila) else ''

    estudiantes = []
    for fila in filas:
        if not fila:
            continue
        nombre = _norm(celda(fila, 'nombre'))
        if not nombre:  # fila sin nombre = fila vacía → se omite
            continue
        est = {'nombre': nombre}
        for clave in columnas.values():
            if clave != 'nombre':
                est[clave] = _entero_texto(celda(fila, clave))
        estudiantes.append(est)
    return estudiantes
