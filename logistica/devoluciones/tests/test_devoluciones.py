"""Tests del módulo de devoluciones de colegios (F5): gates de área y de
módulo (LECTURA vs COMPLETO), alta con captura por grados, detalle pivotado,
export con el layout de la hoja del usuario y autocompletado de colegios desde
el ERP de despachos.

Mismo arnés que el resto de logística: `Client(HTTP_HOST='logistica.testserver')`
y datos sembrados por los servicios (única puerta de escritura al stock).
"""
import io
from datetime import date

from django.contrib.auth.models import Group, User
from django.test import Client, TestCase
from openpyxl import load_workbook

from core.areas import GRUPO_STAFF_FINANCIERA, GRUPO_STAFF_LOGISTICA
from logistica.despachos.models import OrdenDespacho
from logistica.inventario.models import (Bodega, Categoria, DevolucionColegio,
                                         Movimiento, Stock)
from logistica.inventario.services import registrar_devolucion_colegio
from logistica.inventario.tests.utils import crear_material, lineas_post
from usuarios.models import ModuloUsuario

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _filas_xlsx(response):
    ws = load_workbook(io.BytesIO(response.content)).active
    return list(ws.iter_rows(min_row=2, values_only=True))


def _encabezados_xlsx(response):
    ws = load_workbook(io.BytesIO(response.content)).active
    return list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))


class _BaseDevolucionesTest(TestCase):

    def setUp(self):
        self.client = Client(HTTP_HOST='logistica.testserver')
        self.user = User.objects.create_user(username='logis', password='pass')
        self.user.groups.add(Group.objects.get(name=GRUPO_STAFF_LOGISTICA))
        self.client.login(username='logis', password='pass')

        self.categoria = Categoria.objects.create(nombre='Simulacros')
        self.bodega = Bodega.objects.create(nombre='Principal')
        self.material = crear_material(categoria=self.categoria,
                                       referencia='Cuadernillo A')
        self.otro = crear_material(categoria=self.categoria,
                                   referencia='Cuadernillo B')

    def _mensajes(self, response):
        return [str(m) for m in response.context['messages']]

    def _stock(self, item):
        fila = Stock.objects.filter(item=item, bodega=self.bodega).first()
        return fila.cantidad if fila else 0

    def _cabecera(self, **kw):
        datos = {
            'fecha_recibido': '2026-07-20',
            'colegio': 'Colegio Norte',
            'codigo_colegio': 'C-01',
            'regional': 'Nororiente',
            'ejecutivo': 'Ana Ruiz',
            'bodega': self.bodega.pk,
            'observaciones': 'Sobrantes del simulacro',
        }
        datos.update(kw)
        return datos

    def _registrar(self, **kw):
        """Una devolución por servicio (para los tests que solo la consultan)."""
        datos = {'fecha_recibido': date(2026, 7, 20), 'colegio': 'Colegio Norte',
                 'codigo_colegio': 'C-01', 'regional': 'Nororiente',
                 'ejecutivo': 'Ana Ruiz', 'observaciones': 'Sobrantes'}
        lineas = kw.pop('lineas', [(self.material[3], 5), (self.material[4], 2)])
        datos.update(kw)
        return registrar_devolucion_colegio(bodega=self.bodega, lineas=lineas,
                                            usuario=self.user, **datos)


class GatesDevolucionesTest(_BaseDevolucionesTest):

    def _urls(self):
        return ['/devoluciones/', '/devoluciones/nueva/',
                '/devoluciones/detalle/']

    def test_anonimo_redirigido(self):
        c = Client(HTTP_HOST='logistica.testserver')
        for url in self._urls():
            self.assertEqual(c.get(url).status_code, 302, url)

    def test_usuario_de_otra_area_no_entra(self):
        u = User.objects.create_user(username='finan', password='pass')
        u.groups.add(Group.objects.get(name=GRUPO_STAFF_FINANCIERA))
        c = Client(HTTP_HOST='logistica.testserver')
        c.login(username='finan', password='pass')
        for url in self._urls():
            self.assertEqual(c.get(url).status_code, 302, url)

    def test_staff_logistica_ve_las_vistas(self):
        dev = self._registrar()
        for url in self._urls() + [f'/devoluciones/{dev.pk}/']:
            self.assertEqual(self.client.get(url).status_code, 200, url)

    def test_menu_muestra_el_item(self):
        html = self.client.get('/devoluciones/').content.decode()
        self.assertIn('Devoluciones', html)

    def test_lectura_ve_pero_no_registra(self):
        ModuloUsuario.objects.create(user=self.user, area='logistica',
                                     modulo='devoluciones',
                                     nivel=ModuloUsuario.Nivel.LECTURA)
        self.assertEqual(self.client.get('/devoluciones/').status_code, 200)
        resp = self.client.post('/devoluciones/nueva/', {
            **self._cabecera(),
            **lineas_post([(self.material, {0: 3})]),
        })
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(DevolucionColegio.objects.count(), 0)

    def test_lectura_si_puede_exportar(self):
        """El export es un POST "de lectura" (declarado en posts_lectura)."""
        ModuloUsuario.objects.create(user=self.user, area='logistica',
                                     modulo='devoluciones',
                                     nivel=ModuloUsuario.Nivel.LECTURA)
        resp = self.client.post('/devoluciones/exportar/', {})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], XLSX_MIME)

    def test_sin_acceso_al_modulo_no_entra(self):
        ModuloUsuario.objects.create(user=self.user, area='logistica',
                                     modulo='devoluciones',
                                     nivel=ModuloUsuario.Nivel.SIN_ACCESO)
        resp = self.client.get('/devoluciones/')
        self.assertEqual(resp.status_code, 302)


class AltaDevolucionTest(_BaseDevolucionesTest):

    def test_alta_expande_por_grados_y_suma_stock(self):
        resp = self.client.post('/devoluciones/nueva/', {
            **self._cabecera(),
            **lineas_post([(self.material, {0: 3, 5: 7}),
                           (self.otro, {11: 1})]),
        })
        dev = DevolucionColegio.objects.get()
        self.assertRedirects(resp, f'/devoluciones/{dev.pk}/')
        # Una línea por grado con cantidad, no una por fila del formulario.
        self.assertEqual(dev.lineas.count(), 3)
        self.assertEqual(dev.colegio, 'Colegio Norte')
        self.assertEqual(dev.ejecutivo, 'Ana Ruiz')
        self.assertEqual(dev.creado_por, self.user)
        self.assertEqual(self._stock(self.material[0]), 3)
        self.assertEqual(self._stock(self.material[5]), 7)
        self.assertEqual(self._stock(self.otro[11]), 1)
        self.assertEqual(
            Movimiento.objects.filter(tipo=Movimiento.Tipo.DEV_COLEGIO).count(), 3)

    def test_fila_sin_cantidades_da_error_y_no_escribe(self):
        resp = self.client.post('/devoluciones/nueva/', {
            **self._cabecera(),
            **lineas_post([(self.material, {})]),
        })
        self.assertEqual(resp.status_code, 200)  # re-render, no redirect
        self.assertEqual(DevolucionColegio.objects.count(), 0)
        self.assertEqual(Movimiento.objects.count(), 0)
        self.assertTrue(any('no tiene cantidades' in m
                            for m in self._mensajes(resp)))

    def test_cabecera_invalida_conserva_las_lineas(self):
        resp = self.client.post('/devoluciones/nueva/', {
            **self._cabecera(colegio=''),
            **lineas_post([(self.material, {2: 4})]),
        })
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(DevolucionColegio.objects.count(), 0)
        previas = resp.context['lineas_previas']
        self.assertEqual(previas[0]['material'],
                         self.material[0].clave_material)
        self.assertEqual(previas[0]['celdas'][2]['valor'], '4')

    def test_sin_lineas_da_error(self):
        resp = self.client.post('/devoluciones/nueva/', {
            **self._cabecera(),
            **lineas_post([('', {})]),
        })
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(DevolucionColegio.objects.count(), 0)
        self.assertTrue(any('al menos una línea' in m
                            for m in self._mensajes(resp)))


class ColegiosDelErpTest(_BaseDevolucionesTest):
    """El campo de colegio se autocompleta con los clientes del ERP, pero
    acepta texto libre (el que devuelve puede no estar en despachos)."""

    def test_datalist_trae_centro_de_costos_y_cliente(self):
        OrdenDespacho.objects.create(id_orden='PPAL-1',
                                     centro_costos='COLEGIO ALFA',
                                     cliente='CLIENTE ALFA S.A.')
        # Sin centro de costos cae al cliente (criterio de `colegio_erp`).
        OrdenDespacho.objects.create(id_orden='PPAL-2', centro_costos='',
                                     cliente='COLEGIO BETA')
        html = self.client.get('/devoluciones/nueva/').content.decode()
        self.assertIn('COLEGIO ALFA', html)
        self.assertIn('COLEGIO BETA', html)
        self.assertNotIn('CLIENTE ALFA S.A.', html)

    def test_colegio_libre_se_acepta(self):
        self.client.post('/devoluciones/nueva/', {
            **self._cabecera(colegio='Colegio que no está en el ERP'),
            **lineas_post([(self.material, {1: 2})]),
        })
        self.assertEqual(DevolucionColegio.objects.get().colegio,
                         'Colegio que no está en el ERP')


class DetalleDevolucionTest(_BaseDevolucionesTest):

    def test_detalle_pivota_por_material(self):
        dev = self._registrar(lineas=[(self.material[3], 5),
                                      (self.material[4], 2),
                                      (self.otro[0], 1)])
        resp = self.client.get(f'/devoluciones/{dev.pk}/')
        materiales = resp.context['materiales']
        self.assertEqual(len(materiales), 2)  # dos materiales, no tres líneas
        fila = next(m for m in materiales if m['referencia'] == 'Cuadernillo A')
        self.assertEqual(fila['celdas'][3]['cantidad'], 5)
        self.assertEqual(fila['celdas'][4]['cantidad'], 2)
        self.assertIsNone(fila['celdas'][6]['item'])  # grado no devuelto
        self.assertEqual(fila['total'], 7)


class ExportDevolucionesTest(_BaseDevolucionesTest):

    def test_layout_de_la_hoja(self):
        self._registrar(lineas=[(self.material[3], 5), (self.material[4], 2)])
        resp = self.client.post('/devoluciones/exportar/', {})
        self.assertEqual(resp['Content-Type'], XLSX_MIME)
        cabeceras = _encabezados_xlsx(resp)
        self.assertEqual(cabeceras[:7],
                         ['Fecha de recibido', 'Colegio', 'Código', 'Regional',
                          'Ejecutivo', 'Categoría', 'Referencia'])
        self.assertEqual(cabeceras[7:19], [f'{g}°' for g in range(12)])
        self.assertEqual(cabeceras[19:], ['Total', 'Estado',
                                          'Motivo del rechazo', 'Observaciones'])
        # "Registro Effi" se omite por decisión del usuario.
        self.assertNotIn('Registro Effi', cabeceras)

        fila = _filas_xlsx(resp)[0]
        self.assertEqual(fila[0], '20/07/2026')
        self.assertEqual(fila[1], 'Colegio Norte')
        self.assertEqual(fila[5], 'Simulacros')
        self.assertEqual(fila[6], 'Cuadernillo A')
        self.assertEqual(fila[7 + 3], 5)
        self.assertEqual(fila[7 + 4], 2)
        # Grado no devuelto: celda vacía (openpyxl la relee como None), no 0.
        self.assertIsNone(fila[7 + 6])
        self.assertEqual(fila[19], 7)

    def test_una_fila_por_devolucion_y_material(self):
        self._registrar(lineas=[(self.material[0], 1), (self.otro[0], 2)])
        self._registrar(colegio='Colegio Sur', lineas=[(self.material[1], 3)])
        filas = _filas_xlsx(self.client.post('/devoluciones/exportar/', {}))
        self.assertEqual(len(filas), 3)

    def test_filtros_de_fecha_y_colegio(self):
        self._registrar(colegio='Colegio Norte',
                        fecha_recibido=date(2026, 7, 1),
                        lineas=[(self.material[0], 1)])
        self._registrar(colegio='Colegio Sur',
                        fecha_recibido=date(2026, 7, 20),
                        lineas=[(self.material[1], 1)])

        filas = _filas_xlsx(self.client.post('/devoluciones/exportar/',
                                             {'desde': '2026-07-10'}))
        self.assertEqual([f[1] for f in filas], ['Colegio Sur'])

        filas = _filas_xlsx(self.client.post('/devoluciones/exportar/',
                                             {'hasta': '2026-07-10'}))
        self.assertEqual([f[1] for f in filas], ['Colegio Norte'])

        filas = _filas_xlsx(self.client.post('/devoluciones/exportar/',
                                             {'colegio': 'sur'}))
        self.assertEqual([f[1] for f in filas], ['Colegio Sur'])

    def test_get_no_permitido(self):
        self.assertEqual(self.client.get('/devoluciones/exportar/').status_code,
                         405)


class ListaDevolucionesTest(_BaseDevolucionesTest):

    def test_lista_totaliza_materiales_y_unidades(self):
        self._registrar(lineas=[(self.material[0], 4), (self.otro[0], 6)])
        resp = self.client.get('/devoluciones/')
        fila = resp.context['devoluciones'][0]
        self.assertEqual(fila.n_materiales, 2)
        self.assertEqual(fila.unidades, 10)
        self.assertContains(resp, 'Colegio Norte')

    def test_lista_marca_el_estado(self):
        self._registrar(lineas=[(self.material[0], 4)])
        self._registrar(colegio='Colegio Sur', lineas=[(self.otro[0], 1)],
                        valida=False, motivo_no_valida='Cuadernillos rayados')
        html = self.client.get('/devoluciones/').content.decode()
        self.assertIn('No válida', html)
        self.assertIn('Cuadernillos rayados', html)


class DevolucionNoValidaTest(_BaseDevolucionesTest):
    """Se recibió pero no se acepta: queda el registro, no entra al inventario."""

    def test_alta_no_valida_registra_sin_tocar_stock(self):
        resp = self.client.post('/devoluciones/nueva/', {
            **self._cabecera(),
            'no_valida': 'on',
            'motivo_no_valida': 'Material mojado',
            **lineas_post([(self.material, {0: 3, 5: 7})]),
        })
        dev = DevolucionColegio.objects.get()
        self.assertRedirects(resp, f'/devoluciones/{dev.pk}/')
        self.assertFalse(dev.valida)
        self.assertEqual(dev.motivo_no_valida, 'Material mojado')
        self.assertEqual(dev.lineas.count(), 2)  # se sabe qué llegó
        self.assertEqual(self._stock(self.material[0]), 0)
        self.assertEqual(self._stock(self.material[5]), 0)
        self.assertEqual(Movimiento.objects.count(), 0)

    def test_motivo_obligatorio(self):
        resp = self.client.post('/devoluciones/nueva/', {
            **self._cabecera(),
            'no_valida': 'on',
            **lineas_post([(self.material, {0: 3})]),
        })
        self.assertEqual(resp.status_code, 200)  # re-render, no redirect
        self.assertEqual(DevolucionColegio.objects.count(), 0)
        self.assertTrue(any('por qué la devolución no es válida' in m
                            for m in self._mensajes(resp)))

    def test_valida_por_defecto_sigue_sumando(self):
        """Sin el checkbox el comportamiento histórico no cambia."""
        self.client.post('/devoluciones/nueva/', {
            **self._cabecera(),
            **lineas_post([(self.material, {0: 3})]),
        })
        dev = DevolucionColegio.objects.get()
        self.assertTrue(dev.valida)
        self.assertEqual(self._stock(self.material[0]), 3)

    def test_detalle_avisa_y_no_enlaza_kardex(self):
        dev = self._registrar(lineas=[(self.material[0], 2)], valida=False,
                              motivo_no_valida='Faltaban cuadernillos')
        resp = self.client.get(f'/devoluciones/{dev.pk}/')
        self.assertContains(resp, 'Faltaban cuadernillos')
        # Sin movimientos no hay kardex que enlazar.
        self.assertNotContains(
            resp, f'/articulos/{self.material[0].pk}/kardex/')

    def test_detalle_valido_si_enlaza_kardex(self):
        dev = self._registrar(lineas=[(self.material[0], 2)])
        resp = self.client.get(f'/devoluciones/{dev.pk}/')
        self.assertContains(resp, f'/articulos/{self.material[0].pk}/kardex/')

    def test_export_filtra_por_estado(self):
        self._registrar(colegio='Colegio Norte', lineas=[(self.material[0], 1)])
        self._registrar(colegio='Colegio Sur', lineas=[(self.otro[0], 1)],
                        valida=False, motivo_no_valida='Rayados')

        filas = _filas_xlsx(self.client.post('/devoluciones/exportar/',
                                             {'estado': 'validas'}))
        self.assertEqual([f[1] for f in filas], ['Colegio Norte'])

        filas = _filas_xlsx(self.client.post('/devoluciones/exportar/',
                                             {'estado': 'no_validas'}))
        self.assertEqual([f[1] for f in filas], ['Colegio Sur'])
        self.assertEqual(filas[0][20], 'No válida')
        self.assertEqual(filas[0][21], 'Rayados')

        # Sin filtro salen las dos: el default no esconde registros.
        filas = _filas_xlsx(self.client.post('/devoluciones/exportar/', {}))
        self.assertEqual(len(filas), 2)


class DetallePorMaterialTest(_BaseDevolucionesTest):
    """La vista que evita tener que bajar el Excel para ver el detalle."""

    def test_una_fila_por_devolucion_y_material(self):
        self._registrar(lineas=[(self.material[3], 5), (self.material[4], 2),
                                (self.otro[0], 1)])
        self._registrar(colegio='Colegio Sur', lineas=[(self.material[1], 3)])
        resp = self.client.get('/devoluciones/detalle/')
        self.assertEqual(resp.status_code, 200)
        filas = resp.context['filas']
        self.assertEqual(len(filas), 3)  # 2 materiales + 1
        fila = next(f for f in filas
                    if f['devolucion'].colegio == 'Colegio Norte'
                    and f['referencia'] == 'Cuadernillo A')
        self.assertEqual(fila['celdas'][3]['cantidad'], 5)
        self.assertEqual(fila['total'], 7)

    def test_pinta_las_columnas_de_la_hoja(self):
        self._registrar(lineas=[(self.material[3], 5)])
        resp = self.client.get('/devoluciones/detalle/')
        for texto in ['Ejecutivo', 'Categoría', 'Referencia', 'Ana Ruiz',
                      'Cuadernillo A', 'Colegio Norte']:
            self.assertContains(resp, texto)

    def test_marca_las_no_validas(self):
        self._registrar(colegio='Colegio Sur', lineas=[(self.otro[0], 1)],
                        valida=False, motivo_no_valida='Cuadernillos rayados')
        resp = self.client.get('/devoluciones/detalle/')
        self.assertContains(resp, 'No válida')
        self.assertContains(resp, 'Cuadernillos rayados')

    def test_gate_de_area(self):
        c = Client(HTTP_HOST='logistica.testserver')
        self.assertEqual(c.get('/devoluciones/detalle/').status_code, 302)
