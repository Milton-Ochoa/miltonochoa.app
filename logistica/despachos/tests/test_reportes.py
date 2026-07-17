"""Tests de la F5: badge (context processor), export a Excel del tablero con
filtros vigentes, bodega por defecto por usuario y su gestión (superusuario).

Arnés `Client(HTTP_HOST='logistica.testserver')`. Las órdenes se crean por ORM
directo para controlar cada estado con precisión (patrón `test_tablero.py`).
"""
import io
from datetime import datetime, timedelta

from django.contrib.auth.models import Group, User
from django.test import Client, RequestFactory, TestCase
from django.utils import timezone

from openpyxl import load_workbook

from core.areas import GRUPO_STAFF_LOGISTICA
from logistica.despachos.context_processors import alertas_despachos
from logistica.despachos.models import AsignacionBodega, OrdenDespacho
from usuarios.models import ModuloUsuario

_TABLERO = '/despachos/'
_EXPORTAR = '/despachos/exportar/'
_BODEGAS = '/despachos/bodegas/'


def _orden(id_orden, *, estado=OrdenDespacho.Estado.PENDIENTE, es_despachable=True,
           fecha_entrega=None, fecha_orden=None, resumen='2× 727', **kw):
    return OrdenDespacho.objects.create(
        id_orden=id_orden, estado=estado, es_despachable=es_despachable,
        fecha_entrega=fecha_entrega, fecha_orden=fecha_orden,
        resumen_articulos=resumen, cliente=kw.pop('cliente', 'Colegio X'),
        bodega=kw.pop('bodega', 'BUCARAMANGA'), ciudad=kw.pop('ciudad', 'Bucaramanga'),
        departamento=kw.pop('departamento', 'Santander'), **kw)


class _BaseLogistica(TestCase):
    def setUp(self):
        self.client = Client(HTTP_HOST='logistica.testserver')
        self.user = User.objects.create_user(username='logis', password='pass')
        self.user.groups.add(Group.objects.get(name=GRUPO_STAFF_LOGISTICA))
        self.client.login(username='logis', password='pass')


# ---------------------------------------------------------------------------
# Badge (context processor)
# ---------------------------------------------------------------------------

class BadgeTest(TestCase):
    def setUp(self):
        self.rf = RequestFactory()
        hoy = timezone.localdate()
        # Vencida (cuenta): abierta, despachable, entrega pasada.
        _orden('PPAL-1', fecha_entrega=hoy - timedelta(days=1))
        # Próxima (no cuenta): entrega futura.
        _orden('PPAL-2', fecha_entrega=hoy + timedelta(days=1))
        # Vencida pero NO despachable (100% formación) → no cuenta.
        _orden('PPAL-3', es_despachable=False, fecha_entrega=hoy - timedelta(days=1))
        # Vencida pero DESPACHADA (no abierta) → no cuenta.
        _orden('PPAL-4', estado=OrdenDespacho.Estado.DESPACHADA,
               fecha_entrega=hoy - timedelta(days=1))

    def _request(self, area='logistica', personal=True):
        req = self.rf.get('/')
        req.area = area
        req.es_personal_logistica = personal
        return req

    def test_cuenta_solo_vencidas_abiertas_despachables(self):
        self.assertEqual(alertas_despachos(self._request()),
                         {'desp_vencidas_count': 1})

    def test_vacio_fuera_del_area(self):
        self.assertEqual(alertas_despachos(self._request(area='financiera')), {})

    def test_vacio_sin_personal_logistica(self):
        self.assertEqual(alertas_despachos(self._request(personal=False)), {})


# ---------------------------------------------------------------------------
# Export a Excel
# ---------------------------------------------------------------------------

class ExportTest(_BaseLogistica):
    def setUp(self):
        super().setUp()
        self.o1 = _orden('PPAL-10', bodega='BUCARAMANGA', cliente='Colegio Uno')
        self.o2 = _orden('PPAL-11', bodega='MONTERIA', cliente='Colegio Dos')

    def _leer(self, resp):
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        return [[c.value for c in fila] for fila in ws.iter_rows()]

    def test_export_abiertas_incluye_ambas(self):
        resp = self.client.post(_EXPORTAR, {'tab': 'abiertas'})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml', resp['Content-Type'])
        filas = self._leer(resp)
        # fila 0 = cabecera; las órdenes en la primera columna.
        ordenes = {f[0] for f in filas[1:]}
        self.assertEqual(ordenes, {'PPAL-10', 'PPAL-11'})

    def test_export_respeta_filtro_de_columna(self):
        resp = self.client.post(_EXPORTAR, {'tab': 'abiertas', 'f_bodega': 'monteria'})
        ordenes = {f[0] for f in self._leer(resp)[1:]}
        self.assertEqual(ordenes, {'PPAL-11'})

    def test_export_respeta_rango_de_entrega(self):
        hoy = timezone.localdate()
        self.o1.fecha_entrega = hoy
        self.o1.save(update_fields=['fecha_entrega'])
        self.o2.fecha_entrega = hoy + timedelta(days=30)
        self.o2.save(update_fields=['fecha_entrega'])
        resp = self.client.post(_EXPORTAR, {
            'tab': 'abiertas', 'ent_hasta': hoy.strftime('%Y-%m-%d')})
        ordenes = {f[0] for f in self._leer(resp)[1:]}
        self.assertEqual(ordenes, {'PPAL-10'})

    def test_export_columna_colegio_usa_centro_costos_con_fallback(self):
        # centro_costos presente → se exporta ese; vacío → cae al cliente.
        self.o1.centro_costos = 'COLEGIO REAL'
        self.o1.save(update_fields=['centro_costos'])
        filas = self._leer(self.client.post(_EXPORTAR, {'tab': 'abiertas'}))
        self.assertEqual(filas[0][1], 'Colegio')  # cabecera renombrada
        por_orden = {f[0]: f[1] for f in filas[1:]}
        self.assertEqual(por_orden['PPAL-10'], 'COLEGIO REAL')   # centro de costos
        self.assertEqual(por_orden['PPAL-11'], 'Colegio Dos')    # fallback a cliente

    def test_export_filtro_colegio_precede_centro_costos_y_cae_a_cliente(self):
        self.o1.centro_costos = 'COLEGIO REAL'
        self.o1.save(update_fields=['centro_costos'])

        def _ordenes(**post):
            return {f[0] for f in self._leer(
                self.client.post(_EXPORTAR, {'tab': 'abiertas', **post}))[1:]}

        # Casa por centro_costos.
        self.assertEqual(_ordenes(f_cliente='real'), {'PPAL-10'})
        # o1 tiene centro_costos no vacío → NO casa por su cliente.
        self.assertEqual(_ordenes(f_cliente='colegio uno'), set())
        # o2 sin centro_costos → casa por cliente (fallback).
        self.assertEqual(_ordenes(f_cliente='colegio dos'), {'PPAL-11'})

    def test_export_requiere_post(self):
        self.assertEqual(self.client.get(_EXPORTAR).status_code, 405)

    def test_export_permitido_en_lectura(self):
        # `/despachos/exportar/` está en posts_lectura → LECTURA puede exportar.
        ModuloUsuario.objects.create(user=self.user, area='logistica',
                                     modulo='despachos',
                                     nivel=ModuloUsuario.Nivel.LECTURA)
        self.assertEqual(
            self.client.post(_EXPORTAR, {'tab': 'abiertas'}).status_code, 200)


# ---------------------------------------------------------------------------
# Bodega por defecto + gestión (superusuario)
# ---------------------------------------------------------------------------

class BodegaDefaultTest(_BaseLogistica):
    def test_tablero_expone_la_bodega_del_usuario(self):
        AsignacionBodega.objects.create(usuario=self.user, bodega='MONTERIA')
        ctx = self.client.get(_TABLERO).context
        self.assertEqual(ctx['bodega_default'], 'MONTERIA')

    def test_tablero_sin_asignacion_bodega_vacia(self):
        self.assertEqual(self.client.get(_TABLERO).context['bodega_default'], '')


class BodegasGestionTest(_BaseLogistica):
    def setUp(self):
        super().setUp()
        self.admin = User.objects.create_superuser(
            username='root', password='pass', email='r@r.com')
        self.admin_client = Client(HTTP_HOST='logistica.testserver')
        self.admin_client.login(username='root', password='pass')

    def test_no_superusuario_no_ve_la_pagina(self):
        # El staff de logística (no superusuario) es rechazado al tablero.
        resp = self.client.get(_BODEGAS)
        self.assertRedirects(resp, _TABLERO)

    def test_superusuario_ve_la_pagina(self):
        self.assertEqual(self.admin_client.get(_BODEGAS).status_code, 200)

    def test_superusuario_asigna_bodega(self):
        self.admin_client.post(_BODEGAS, {'user_id': self.user.pk,
                                          'bodega': 'BUCARAMANGA'})
        asignacion = AsignacionBodega.objects.get(usuario=self.user)
        self.assertEqual(asignacion.bodega, 'BUCARAMANGA')
        self.assertEqual(asignacion.asignado_por, self.admin)

    def test_asignar_vacio_borra_la_asignacion(self):
        AsignacionBodega.objects.create(usuario=self.user, bodega='MONTERIA')
        self.admin_client.post(_BODEGAS, {'user_id': self.user.pk, 'bodega': ''})
        self.assertFalse(
            AsignacionBodega.objects.filter(usuario=self.user).exists())

    def test_no_superusuario_no_puede_asignar(self):
        resp = self.client.post(_BODEGAS, {'user_id': self.user.pk,
                                           'bodega': 'BUCARAMANGA'})
        self.assertRedirects(resp, _TABLERO)
        self.assertFalse(
            AsignacionBodega.objects.filter(usuario=self.user).exists())

    def test_reasignar_actualiza_en_vez_de_duplicar(self):
        self.admin_client.post(_BODEGAS, {'user_id': self.user.pk,
                                          'bodega': 'BUCARAMANGA'})
        self.admin_client.post(_BODEGAS, {'user_id': self.user.pk,
                                          'bodega': 'MONTERIA'})
        self.assertEqual(
            AsignacionBodega.objects.filter(usuario=self.user).count(), 1)
        self.assertEqual(
            AsignacionBodega.objects.get(usuario=self.user).bodega, 'MONTERIA')
