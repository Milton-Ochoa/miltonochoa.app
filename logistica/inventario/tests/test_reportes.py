"""Tests de la Fase 6 (dashboard, badges y exports): context processor de
alertas fuera/dentro del área, badges del menú, tarjetas del dashboard y los
tres exports a Excel con sus filtros.

Mismo arnés que el resto de fases: `Client(HTTP_HOST='logistica.testserver')`
y datos sembrados SIEMPRE vía services (única puerta de escritura al stock).
"""
import io
from datetime import timedelta

from django.contrib.auth.models import Group, User
from django.test import Client, TestCase
from django.utils import timezone
from openpyxl import load_workbook

from core.areas import GRUPO_STAFF_FINANCIERA, GRUPO_STAFF_LOGISTICA

from logistica.inventario.models import (GRADOS, Bodega, Categoria, Item,
                                         Movimiento, Prestamo, Tercero)
from logistica.inventario.services import crear_prestamo, registrar_entrada
from logistica.inventario.tests.utils import crear_item

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _filas_xlsx(response):
    """Filas de datos (sin cabecera) del xlsx de la respuesta."""
    ws = load_workbook(io.BytesIO(response.content)).active
    return list(ws.iter_rows(min_row=2, values_only=True))


def _encabezados_xlsx(response):
    ws = load_workbook(io.BytesIO(response.content)).active
    return list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))


class _BaseReportesTest(TestCase):
    """Staff de logística logueado + inventario sembrado: un artículo bajo
    mínimo, stock en dos bodegas, un préstamo otorgado VENCIDO y uno recibido
    vigente."""

    def setUp(self):
        self.client = Client(HTTP_HOST='logistica.testserver')
        self.user = User.objects.create_user(username='logis', password='pass')
        self.user.groups.add(Group.objects.get(name=GRUPO_STAFF_LOGISTICA))
        self.client.login(username='logis', password='pass')

        cat = Categoria.objects.create(nombre='Papelería')
        self.bodega_a = Bodega.objects.create(nombre='Principal')
        self.bodega_b = Bodega.objects.create(nombre='Anexa')
        # stock_minimo=10: con 3 de la entrada + 4 del préstamo recibido el
        # total queda en 7 → bajo mínimo (el mínimo es GLOBAL, suma de bodegas).
        self.marcador = crear_item(categoria=cat, referencia='Marcador',
                                   grado=3, stock_minimo=10,
                                   valor_unitario=2000)
        self.resma = crear_item(categoria=cat, referencia='Resma', grado=4)
        registrar_entrada(bodega=self.bodega_a,
                          lineas=[(self.marcador, 3), (self.resma, 10)],
                          usuario=self.user, proveedor='ACME')

        hoy = timezone.localdate()
        tercero = Tercero.objects.create(nombre='Asesor Uno')
        # OTORGADO vencido (compromiso ayer): nos deben 2 resmas.
        self.prestamo_otorgado = crear_prestamo(
            tercero=tercero, fecha_compromiso=hoy - timedelta(days=1),
            lineas=[(self.resma, self.bodega_a, 2)], usuario=self.user)
        # RECIBIDO vigente (compromiso mañana): nos prestan 4 marcadores
        # que entran a la bodega anexa.
        self.prestamo_recibido = crear_prestamo(
            tercero=tercero, fecha_compromiso=hoy + timedelta(days=1),
            lineas=[(self.marcador, self.bodega_b, 4)], usuario=self.user,
            direccion=Prestamo.Direccion.RECIBIDO)


class ContextProcessorTest(_BaseReportesTest):

    def test_contadores_en_el_area(self):
        r = self.client.get('/')
        self.assertEqual(r.context['inv_bajo_minimo_count'], 1)
        self.assertEqual(r.context['inv_prestamos_vencidos_count'], 1)

    def test_badge_de_prestamos_vencidos_en_el_menu(self):
        r = self.client.get('/')
        self.assertContains(r, 'nav-badge')

    def test_fuera_del_area_no_se_calcula(self):
        u = User.objects.create_user(username='finan', password='pass')
        u.groups.add(Group.objects.get(name=GRUPO_STAFF_FINANCIERA))
        c = Client(HTTP_HOST='financiera.testserver')
        c.login(username='finan', password='pass')
        r = c.get('/')
        self.assertNotIn('inv_bajo_minimo_count', r.context)

    def test_sin_alertas_no_hay_badges(self):
        # Sin borrar préstamos (el ledger los protege con FK PROTECT): basta
        # quitar el mínimo y correr el compromiso al futuro.
        Item.objects.update(stock_minimo=0)
        Prestamo.objects.update(
            fecha_compromiso=timezone.localdate() + timedelta(days=30))
        r = self.client.get('/')
        self.assertEqual(r.context['inv_bajo_minimo_count'], 0)
        self.assertEqual(r.context['inv_prestamos_vencidos_count'], 0)


class DashboardTest(_BaseReportesTest):

    def test_tarjetas_del_dashboard(self):
        r = self.client.get('/')
        self.assertEqual(r.context['items_activos'], 2)
        # 3 + 10 (entrada) − 2 (préstamo otorgado) + 4 (recibido) = 15
        self.assertEqual(r.context['unidades_totales'], 15)
        self.assertEqual(r.context['otorgados_abiertos'], 1)
        self.assertEqual(r.context['otorgados_vencidos'], 1)
        self.assertEqual(r.context['recibidos_abiertos'], 1)
        self.assertEqual(r.context['recibidos_vencidos'], 0)
        self.assertContains(r, 'Nos deben')
        self.assertContains(r, 'Debemos devolver')

    def test_ultimos_movimientos_en_el_dashboard(self):
        r = self.client.get('/')
        # entrada (2 líneas) + préstamo + recibido = 4 movimientos
        self.assertEqual(len(r.context['ultimos_movimientos']), 4)
        self.assertContains(r, 'Últimos movimientos')


class GatesExportsTest(_BaseReportesTest):

    URLS = ['/stock/exportar/', '/movimientos/exportar/', '/prestamos/exportar/']

    def test_anonimo_no_exporta(self):
        c = Client(HTTP_HOST='logistica.testserver')
        for url in self.URLS:
            r = c.post(url)
            self.assertEqual(r.status_code, 302, url)

    def test_usuario_de_otra_area_no_exporta(self):
        u = User.objects.create_user(username='finan2', password='pass')
        u.groups.add(Group.objects.get(name=GRUPO_STAFF_FINANCIERA))
        c = Client(HTTP_HOST='logistica.testserver')
        c.login(username='finan2', password='pass')
        for url in self.URLS:
            r = c.post(url)
            self.assertEqual(r.status_code, 302, url)

    def test_get_no_permitido(self):
        for url in self.URLS:
            r = self.client.get(url)
            self.assertEqual(r.status_code, 405, url)


class ExportStockTest(_BaseReportesTest):
    """El export es PIVOTADO: una fila por (material, bodega) con una columna
    por grado. Índices de columna (0-based): 0 Categoría, 1 Referencia,
    2 Bodega, 3 Unidad, 4+g el grado g, 16 Total, 17 Mínimo, 18 Valor unitario,
    19 Valor total."""

    COL_GRADO_0, COL_TOTAL = 4, 16

    def _col_grado(self, grado):
        return self.COL_GRADO_0 + grado

    def test_cabecera_pivotada(self):
        r = self.client.post('/stock/exportar/')
        cabecera = _encabezados_xlsx(r)
        self.assertEqual(cabecera[:4], ['Categoría', 'Referencia', 'Bodega',
                                        'Unidad'])
        self.assertEqual(cabecera[self.COL_GRADO_0:self.COL_TOTAL],
                         [f'{g}°' for g in GRADOS])
        self.assertEqual(cabecera[self.COL_TOTAL], 'Total')

    def test_exporta_todas_las_bodegas(self):
        r = self.client.post('/stock/exportar/')
        self.assertEqual(r['Content-Type'], XLSX_MIME)
        filas = _filas_xlsx(r)
        # Una fila por (material, bodega): marcador@A, marcador@B, resma@A
        self.assertEqual(len(filas), 3)
        self.assertEqual({(f[0], f[1], f[2]) for f in filas},
                         {('Papelería', 'Marcador', 'Principal'),
                          ('Papelería', 'Marcador', 'Anexa'),
                          ('Papelería', 'Resma', 'Principal')})

    def test_cantidad_en_la_columna_de_su_grado(self):
        r = self.client.post('/stock/exportar/', {'bodega': self.bodega_a.pk})
        filas = {f[1]: f for f in _filas_xlsx(r)}
        # El marcador es de 3° y la resma de 4°: cada cantidad en su columna.
        self.assertEqual(filas['Marcador'][self._col_grado(3)], 3)
        # La resma entró con 10 y salieron 2 en el préstamo otorgado.
        self.assertEqual(filas['Resma'][self._col_grado(4)], 8)
        # Un grado que el material no tiene va vacío, no en 0.
        self.assertIsNone(filas['Marcador'][self._col_grado(4)])

    def test_filtro_de_bodega(self):
        r = self.client.post('/stock/exportar/', {'bodega': self.bodega_b.pk})
        filas = _filas_xlsx(r)
        self.assertEqual(len(filas), 1)
        self.assertEqual((filas[0][1], filas[0][2]), ('Marcador', 'Anexa'))
        self.assertEqual(filas[0][self._col_grado(3)], 4)

    def test_valor_total_referencial(self):
        r = self.client.post('/stock/exportar/', {'bodega': self.bodega_b.pk})
        fila = _filas_xlsx(r)[0]
        # 4 marcadores × $2000 (el total de la fila, no de una celda)
        self.assertEqual(fila[self.COL_TOTAL], 4)
        self.assertEqual(fila[18], 2000)
        self.assertEqual(fila[19], 8000)


class ExportMovimientosTest(_BaseReportesTest):

    def test_exporta_el_historico_completo(self):
        r = self.client.post('/movimientos/exportar/')
        self.assertEqual(r['Content-Type'], XLSX_MIME)
        self.assertEqual(len(_filas_xlsx(r)), Movimiento.objects.count())

    def test_filtro_por_tipo(self):
        r = self.client.post('/movimientos/exportar/', {'tipos': ['ENTRADA']})
        filas = _filas_xlsx(r)
        self.assertEqual(len(filas), 2)  # las dos líneas de la entrada
        self.assertTrue(all(f[1] == 'Entrada' for f in filas))

    def test_filtro_por_rango_de_fechas(self):
        manana = (timezone.localdate() + timedelta(days=1)).isoformat()
        r = self.client.post('/movimientos/exportar/', {'desde': manana})
        self.assertEqual(len(_filas_xlsx(r)), 0)

    def test_tipo_invalido_se_ignora(self):
        r = self.client.post('/movimientos/exportar/', {'tipos': ['NOEXISTE']})
        # Tipos inválidos se descartan → sin filtro → histórico completo.
        self.assertEqual(len(_filas_xlsx(r)), Movimiento.objects.count())


class ExportPrestamosTest(_BaseReportesTest):

    def test_exporta_ambas_direcciones(self):
        r = self.client.post('/prestamos/exportar/')
        self.assertEqual(r['Content-Type'], XLSX_MIME)
        self.assertEqual(len(_filas_xlsx(r)), 2)

    def test_filtro_por_direccion(self):
        r = self.client.post('/prestamos/exportar/',
                             {'direcciones': ['OTORGADO']})
        filas = _filas_xlsx(r)
        self.assertEqual(len(filas), 1)
        self.assertEqual(filas[0][0], self.prestamo_otorgado.pk)
        self.assertEqual(filas[0][2], 'Prestamos nosotros')

    def test_solo_vencidos(self):
        r = self.client.post('/prestamos/exportar/', {'solo_vencidos': '1'})
        filas = _filas_xlsx(r)
        self.assertEqual(len(filas), 1)
        self.assertEqual(filas[0][7], 'Sí')

    def test_filtro_por_estado_sin_resultados(self):
        r = self.client.post('/prestamos/exportar/', {'estados': ['CERRADO']})
        self.assertEqual(len(_filas_xlsx(r)), 0)

    def test_totales_por_prestamo(self):
        r = self.client.post('/prestamos/exportar/',
                             {'direcciones': ['OTORGADO']})
        fila = _filas_xlsx(r)[0]
        # prestado=2, devuelto=0, pendiente=2
        self.assertEqual((fila[8], fila[9], fila[10]), (2, 0, 2))
