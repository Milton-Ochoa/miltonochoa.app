import io
import os
from datetime import date

from django.contrib import messages
from django.contrib.auth.models import Group, User
from django.db import models, transaction
from django.db.models import ProtectedError
from django.db.models.functions import Coalesce
from django.http import FileResponse, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.areas import GRUPO_STAFF_LOGISTICA

from .adjuntos import validar_adjunto
from .forms import (BodegaForm, CategoriaForm, EntradaForm, MaterialForm,
                    PrestamoForm, SalidaForm, TerceroForm, TrasladoForm,
                    parsear_clave_material, parsear_lineas_material)
from .models import (GRADOS, AdjuntoEntrada, Bodega, BodegaUsuario, Categoria,
                     Entrada, Item, Movimiento, Prestamo, Salida, Stock,
                     Tercero, Traslado)
from .permisos import BodegaNoPermitida, exigir_bodega, solo_logistica
from .pivote import agrupar_materiales, agrupar_materiales_por_bodega
from .services import (ErrorDevolucion, StockInsuficiente, crear_prestamo,
                       items_bajo_minimo, kardex, registrar_ajuste,
                       registrar_devolucion, registrar_entrada,
                       registrar_salida, registrar_traslado)


@solo_logistica
def home(request):
    """Dashboard del inventario: tarjetas de stock y préstamos + últimos
    movimientos. "Nos deben" = préstamos OTORGADOS no cerrados; "debemos
    devolver" = RECIBIDOS no cerrados (la dirección invierte quién retiene
    el material)."""
    hoy = timezone.localdate()
    abiertos = Prestamo.objects.exclude(estado=Prestamo.Estado.CERRADO)
    otorgados = abiertos.filter(direccion=Prestamo.Direccion.OTORGADO)
    recibidos = abiertos.filter(direccion=Prestamo.Direccion.RECIBIDO)
    return render(request, 'inventario/home.html', {
        'items_activos': Item.objects.filter(activo=True).count(),
        'unidades_totales': (Stock.objects.filter(item__activo=True)
                             .aggregate(t=models.Sum('cantidad'))['t'] or 0),
        'items_alerta': list(items_bajo_minimo()),
        'otorgados_abiertos': otorgados.count(),
        'otorgados_vencidos': otorgados.filter(fecha_compromiso__lt=hoy).count(),
        'recibidos_abiertos': recibidos.count(),
        'recibidos_vencidos': recibidos.filter(fecha_compromiso__lt=hoy).count(),
        'ultimos_movimientos': (Movimiento.objects
                                .select_related('item', 'item__categoria',
                                                'bodega', 'creado_por')
                                .order_by('-creado_en', '-id')[:10]),
    })


def _form_a_messages(request, form):
    """Vuelca los errores de un form a `messages.error` (los catálogos usan
    modales con POST-redirect, no re-render del form con errores)."""
    for campo, errores in form.errors.items():
        etiqueta = form.fields[campo].label if campo in form.fields else None
        for error in errores:
            messages.error(request, f'{etiqueta}: {error}' if etiqueta else error)


# ---------------------------------------------------------------------------
# Artículos
# ---------------------------------------------------------------------------

@solo_logistica
def items_lista(request):
    """Catálogo pivotado: una fila por MATERIAL (categoría, referencia) con sus
    12 grados en columnas — es como el usuario lo maneja en su Excel. El Item
    por grado sigue siendo la unidad real (cada celda enlaza a su kardex)."""
    items = (Item.objects.select_related('categoria')
             .annotate(stock_total=Coalesce(models.Sum('stocks__cantidad'), 0))
             .order_by('categoria__nombre', 'referencia', 'grado'))
    bajo_minimo_ids = set(items_bajo_minimo().values_list('pk', flat=True))
    materiales = agrupar_materiales(items)
    for fila in materiales:
        # El mínimo es por item, pero la alerta se resume en la fila para no
        # obligar a leer las 12 celdas.
        fila['alerta'] = any(c['item'] and c['item'].pk in bajo_minimo_ids
                             for c in fila['celdas'])
    return render(request, 'inventario/items_lista.html', {
        'materiales': materiales,
        'grados': GRADOS,
        'form': MaterialForm(),
        'hay_categorias': Categoria.objects.exists(),
        'bajo_minimo_ids': bajo_minimo_ids,
    })


@require_POST
@solo_logistica
def material_guardar(request):
    """Alta/edición de un material completo: los 12 grados de una vez.

    El alta es eager (crea los 12 `Item` aunque su stock nazca en 0) porque
    los documentos necesitan el Item ya resuelto al capturar cantidades por
    grado. La edición actualiza los campos compartidos de TODO el grupo.
    """
    original = parsear_clave_material(request.POST.get('material'))
    form = MaterialForm(request.POST, original=original)
    if not form.is_valid():
        _form_a_messages(request, form)
        return redirect('log_items_lista')

    datos = form.cleaned_data
    compartidos = {
        'unidad_medida': datos['unidad_medida'],
        'descripcion': datos['descripcion'],
        'stock_minimo': datos['stock_minimo'],
        'valor_unitario': datos['valor_unitario'],
        'activo': datos['activo'],
    }
    etiqueta = f'{datos["categoria"].nombre} {datos["referencia"]}'.strip()
    if original is None:
        with transaction.atomic():
            Item.objects.bulk_create([
                Item(categoria=datos['categoria'], referencia=datos['referencia'],
                     grado=grado, **compartidos) for grado in GRADOS])
        messages.success(request, f'Material "{etiqueta}" creado con sus '
                                  f'{len(GRADOS)} grados.')
    else:
        cat_id, referencia = original
        grupo = Item.objects.filter(categoria_id=cat_id, referencia=referencia)
        if not grupo.exists():
            messages.error(request, 'El material ya no existe.')
            return redirect('log_items_lista')
        grupo.update(categoria=datos['categoria'],
                     referencia=datos['referencia'], **compartidos)
        messages.success(request, f'Material "{etiqueta}" actualizado.')
    return redirect('log_items_lista')


# ---------------------------------------------------------------------------
# Catálogos: bodegas, categorías, terceros
# ---------------------------------------------------------------------------

@solo_logistica
def bodegas(request):
    if request.method == 'POST':
        if request.POST.get('accion') == 'toggle':
            bodega = get_object_or_404(Bodega, pk=request.POST.get('bodega_id'))
            if bodega.activa:
                # Guard del soft-delete: una bodega con existencias no se puede
                # desactivar (quedaría stock "invisible" para los documentos).
                total = bodega.stocks.aggregate(t=models.Sum('cantidad'))['t'] or 0
                if total > 0:
                    messages.error(
                        request,
                        f'No se puede desactivar "{bodega.nombre}": aún tiene '
                        f'{total} unidades en existencia. Trasládalas o ajústalas primero.')
                    return redirect('log_bodegas')
                # Segundo guard: si alguien la tiene asignada, desactivarla lo
                # dejaría sin poder escribir en NINGUNA bodega.
                asignados = list(bodega.usuarios_asignados
                                 .select_related('usuario')
                                 .values_list('usuario__username', flat=True))
                if asignados:
                    messages.error(
                        request,
                        f'No se puede desactivar "{bodega.nombre}": está asignada a '
                        f'{len(asignados)} usuario(s) ({", ".join(asignados)}). '
                        f'Reasígnalos primero en Bodegas por usuario.')
                    return redirect('log_bodegas')
                bodega.activa = False
                messages.success(request, f'Bodega "{bodega.nombre}" desactivada.')
            else:
                bodega.activa = True
                messages.success(request, f'Bodega "{bodega.nombre}" reactivada.')
            bodega.save(update_fields=['activa'])
        else:
            bodega_id = request.POST.get('bodega_id') or None
            instancia = get_object_or_404(Bodega, pk=bodega_id) if bodega_id else None
            form = BodegaForm(request.POST, instance=instancia)
            if form.is_valid():
                bodega = form.save()
                verbo = 'actualizada' if instancia else 'creada'
                messages.success(request, f'Bodega "{bodega.nombre}" {verbo}.')
            else:
                _form_a_messages(request, form)
        return redirect('log_bodegas')

    lista = (Bodega.objects
             .annotate(stock_total=Coalesce(models.Sum('stocks__cantidad'), 0))
             .order_by('nombre'))
    return render(request, 'inventario/bodegas.html', {'bodegas': lista})


@solo_logistica
def bodegas_usuarios(request):
    """Asigna a cada usuario de logística la bodega en la que puede ESCRIBIR.

    Solo superusuario (patrón `plantilla_eliminar`: el botón se oculta a los
    demás y la vista rechaza el POST). Sin asignación el usuario opera todas las
    bodegas, que es el comportamiento histórico.
    """
    if not request.user.is_superuser:
        messages.error(request, 'Solo el administrador puede asignar bodegas.')
        return redirect('log_bodegas')

    if request.method == 'POST':
        usuario = get_object_or_404(User, pk=request.POST.get('user_id', ''))
        bodega_id = (request.POST.get('bodega_id') or '').strip()
        if bodega_id.isdigit():
            # Solo activas: asignar una inactiva dejaría al usuario sin poder
            # escribir en ninguna parte.
            bodega = get_object_or_404(Bodega, pk=bodega_id, activa=True)
            BodegaUsuario.objects.update_or_create(
                usuario=usuario,
                defaults={'bodega': bodega, 'asignado_por': request.user})
            messages.success(
                request,
                f'{usuario.username} solo podrá registrar movimientos en '
                f'"{bodega.nombre}".')
        else:
            BodegaUsuario.objects.filter(usuario=usuario).delete()
            messages.success(
                request,
                f'{usuario.username} ya no tiene bodega asignada: vuelve a operar '
                f'todas.')
        return redirect('log_bodegas_usuarios')

    # Usuarios del área (grupo de etiqueta) + los que ya tengan asignación.
    grupo = Group.objects.filter(name=GRUPO_STAFF_LOGISTICA).first()
    usuarios = User.objects.filter(is_active=True)
    if grupo:
        usuarios = usuarios.filter(groups=grupo)
    else:
        usuarios = usuarios.filter(bodega_inventario__isnull=False)
    usuarios = (usuarios.exclude(is_superuser=True)
                .select_related('bodega_inventario__bodega')
                .order_by('username').distinct())

    return render(request, 'inventario/bodegas_usuarios.html', {
        'usuarios': usuarios,
        'bodegas': Bodega.objects.filter(activa=True).order_by('nombre'),
    })


@solo_logistica
def categorias(request):
    if request.method == 'POST':
        if request.POST.get('accion') == 'eliminar':
            categoria = get_object_or_404(Categoria, pk=request.POST.get('categoria_id'))
            try:
                categoria.delete()
                messages.success(request, f'Categoría "{categoria.nombre}" eliminada.')
            except ProtectedError:
                messages.error(
                    request,
                    f'No se puede eliminar "{categoria.nombre}": hay artículos '
                    f'en esa categoría.')
        else:
            categoria_id = request.POST.get('categoria_id') or None
            instancia = (get_object_or_404(Categoria, pk=categoria_id)
                         if categoria_id else None)
            form = CategoriaForm(request.POST, instance=instancia)
            if form.is_valid():
                categoria = form.save()
                verbo = 'actualizada' if instancia else 'creada'
                messages.success(request, f'Categoría "{categoria.nombre}" {verbo}.')
            else:
                _form_a_messages(request, form)
        return redirect('log_categorias')

    lista = Categoria.objects.annotate(n_items=models.Count('items')).order_by('nombre')
    return render(request, 'inventario/categorias.html', {'categorias': lista})


@solo_logistica
def terceros(request):
    if request.method == 'POST':
        if request.POST.get('accion') == 'toggle':
            tercero = get_object_or_404(Tercero, pk=request.POST.get('tercero_id'))
            tercero.activo = not tercero.activo
            tercero.save(update_fields=['activo'])
            verbo = 'reactivado' if tercero.activo else 'desactivado'
            messages.success(request, f'Tercero "{tercero.nombre}" {verbo}.')
        else:
            tercero_id = request.POST.get('tercero_id') or None
            instancia = (get_object_or_404(Tercero, pk=tercero_id)
                         if tercero_id else None)
            form = TerceroForm(request.POST, instance=instancia)
            if form.is_valid():
                tercero = form.save()
                verbo = 'actualizado' if instancia else 'creado'
                messages.success(request, f'Tercero "{tercero.nombre}" {verbo}.')
            else:
                _form_a_messages(request, form)
        return redirect('log_terceros')

    return render(request, 'inventario/terceros.html',
                  {'terceros': Tercero.objects.order_by('nombre')})


@require_POST
@solo_logistica
def tercero_ajax_crear(request):
    """Alta de tercero al vuelo (JSON). Lo consumen los forms de salidas (F4)
    y préstamos (F5) para no abandonar el documento a medio llenar.

    Contrato: éxito → {ok: true, tercero: {id, nombre, documento, label}};
    error → 400 + {ok: false, error: '<texto legible>'}.
    """
    form = TerceroForm(request.POST)
    if form.is_valid():
        tercero = form.save()
        return JsonResponse({'ok': True, 'tercero': {
            'id': tercero.pk,
            'nombre': tercero.nombre,
            'documento': tercero.documento,
            'label': str(tercero),
        }})
    error = '; '.join(e for lista in form.errors.values() for e in lista)
    return JsonResponse({'ok': False, 'error': error}, status=400)


# ---------------------------------------------------------------------------
# Existencias
# ---------------------------------------------------------------------------

def _items_activos():
    """Todos los items activos: el pivote necesita el juego completo de grados
    de cada material, no solo los que ya tienen existencia."""
    return (Item.objects.filter(activo=True).select_related('categoria')
            .order_by('categoria__nombre', 'referencia', 'grado'))


def _stocks_visibles():
    return (Stock.objects
            .select_related('item', 'item__categoria', 'bodega')
            .filter(item__activo=True))


@solo_logistica
def stock(request):
    """Existencias pivotadas: una fila por (material, bodega) con los 12 grados
    en columnas. Cada celda abre el ajuste de ese Item×Bodega — también las que
    están en 0 o sin fila de `Stock` todavía (el servicio la crea)."""
    filas = agrupar_materiales_por_bodega(_items_activos(), _stocks_visibles())
    alertas = list(items_bajo_minimo())
    return render(request, 'inventario/stock.html', {
        'filas': filas,
        'grados': GRADOS,
        'items_alerta': alertas,
        'bajo_minimo_ids': {item.pk for item in alertas},
        # Para el select del modal de exportar (también inactivas: tienen historial).
        'bodegas': Bodega.objects.order_by('nombre'),
    })


# ---------------------------------------------------------------------------
# Documentos: entradas, salidas, traslados, préstamos
#
# Patrón común: la cabecera la valida un form, las líneas (un material con sus
# 12 cantidades por grado) las parsea `parsear_lineas_material` —que las expande
# a una línea de servicio por grado— y el documento lo crea SIEMPRE el servicio
# de dominio (atómico: si una línea falla, nada queda escrito). En error se
# re-renderiza el form con las filas que traía el POST (`lineas_previas`) para
# no perder lo digitado; en éxito, POST-redirect al detalle con toast.
# ---------------------------------------------------------------------------

def _fila_material_blanco():
    return {'material': '', 'bodega_id': '',
            'celdas': [{'grado': g, 'valor': ''} for g in GRADOS]}


def _lineas_previas_material(post=None, *, con_bodega=False):
    """Filas crudas para pintar `_lineas_material.html`.

    Con `post` repite lo digitado (re-render tras error); sin él devuelve una
    fila en blanco. Garantiza SIEMPRE al menos una fila: el JS del parcial
    clona la primera como plantilla y sin ninguna no podría agregar líneas.
    """
    if post is None:
        return [_fila_material_blanco()]
    claves = post.getlist('linea_material')
    columnas = {g: post.getlist(f'linea_g{g}') for g in GRADOS}
    bodegas = post.getlist('linea_bodega') if con_bodega else []
    filas = []
    for i, clave in enumerate(claves):
        filas.append({
            'material': clave,
            'bodega_id': bodegas[i] if i < len(bodegas) else '',
            'celdas': [{'grado': g,
                        'valor': columnas[g][i] if i < len(columnas[g]) else ''}
                       for g in GRADOS],
        })
    return filas or [_fila_material_blanco()]


def _materiales_para_lineas():
    """Los MATERIALES activos para el select de líneas (una opción por
    (categoría, referencia), no por grado). `select_related` obligatorio: la
    etiqueta y la clave leen la categoría."""
    return agrupar_materiales(_items_activos())


@solo_logistica
def entradas_lista(request):
    lista = (Entrada.objects.select_related('bodega', 'creado_por')
             .annotate(n_lineas=models.Count('lineas', distinct=True),
                       unidades=Coalesce(models.Sum('lineas__cantidad'), 0))
             .order_by('-creado_en'))
    return render(request, 'inventario/entradas_lista.html', {'entradas': lista})


@solo_logistica
def entrada_nueva(request):
    form = EntradaForm(request.POST or None, usuario=request.user)
    if request.method == 'POST':
        try:
            if not form.is_valid():
                _form_a_messages(request, form)
                raise ValueError('')  # cae al re-render conservando las líneas
            # 2.ª barrera (la 1.ª es el queryset recortado del form).
            exigir_bodega(request.user, form.cleaned_data['bodega'])
            lineas = parsear_lineas_material(request.POST)
            entrada = registrar_entrada(
                bodega=form.cleaned_data['bodega'], lineas=lineas,
                usuario=request.user,
                proveedor=form.cleaned_data['proveedor'],
                observaciones=form.cleaned_data['observaciones'])
        except (ValueError, BodegaNoPermitida) as e:
            if str(e):
                messages.error(request, str(e))
        else:
            messages.success(request, f'Entrada #{entrada.pk} registrada.')
            return redirect('log_entradas_detalle', pk=entrada.pk)
    return render(request, 'inventario/entrada_form.html', {
        'form': form,
        'materiales': _materiales_para_lineas(),
        'grados': GRADOS,
        'lineas_previas': _lineas_previas_material(
            request.POST if request.method == 'POST' else None),
    })


@solo_logistica
def entrada_detalle(request, pk):
    entrada = get_object_or_404(
        Entrada.objects.select_related('bodega', 'creado_por')
        .prefetch_related('lineas__item__categoria', 'adjuntos__subido_por'), pk=pk)
    return render(request, 'inventario/entrada_detalle.html', {'entrada': entrada})


@require_POST
@solo_logistica
def entrada_adjunto_subir(request, pk):
    entrada = get_object_or_404(Entrada, pk=pk)
    archivo = request.FILES.get('archivo')
    if not archivo:
        messages.error(request, 'Selecciona un archivo para subir.')
        return redirect('log_entradas_detalle', pk=entrada.pk)
    error = validar_adjunto(archivo)
    if error:
        messages.error(request, error)
        return redirect('log_entradas_detalle', pk=entrada.pk)
    AdjuntoEntrada.objects.create(
        entrada=entrada, archivo=archivo,
        nombre_original=archivo.name[:255], subido_por=request.user)
    messages.success(request, 'Adjunto subido.')
    return redirect('log_entradas_detalle', pk=entrada.pk)


@require_POST
@solo_logistica
def entrada_adjunto_eliminar(request, adjunto_id):
    adjunto = get_object_or_404(AdjuntoEntrada, pk=adjunto_id)
    entrada_pk = adjunto.entrada_id
    # Primero el archivo del storage, luego la fila (patrón soportes de pago).
    adjunto.archivo.delete(save=False)
    adjunto.delete()
    messages.success(request, 'Adjunto eliminado.')
    return redirect('log_entradas_detalle', pk=entrada_pk)


@solo_logistica
def entrada_adjunto_descargar(request, adjunto_id):
    """Descarga proxiada por el backend de storage (disco o S3): el gate de
    permiso queda server-side y NUNCA se exponen URLs firmadas. `?inline=1`
    abre en pestaña; por defecto descarga."""
    adjunto = get_object_or_404(AdjuntoEntrada, pk=adjunto_id)
    nombre = os.path.basename(adjunto.archivo.name)
    return FileResponse(adjunto.archivo.open('rb'),
                        as_attachment=request.GET.get('inline') != '1',
                        filename=nombre)


@solo_logistica
def salidas_lista(request):
    lista = (Salida.objects.select_related('bodega', 'creado_por')
             .annotate(n_lineas=models.Count('lineas', distinct=True),
                       unidades=Coalesce(models.Sum('lineas__cantidad'), 0))
             .order_by('-creado_en'))
    return render(request, 'inventario/salidas_lista.html', {'salidas': lista})


@solo_logistica
def salida_nueva(request):
    form = SalidaForm(request.POST or None, usuario=request.user)
    if request.method == 'POST':
        try:
            if not form.is_valid():
                _form_a_messages(request, form)
                raise ValueError('')
            exigir_bodega(request.user, form.cleaned_data['bodega'])
            lineas = parsear_lineas_material(request.POST)
            salida = registrar_salida(
                bodega=form.cleaned_data['bodega'], lineas=lineas,
                usuario=request.user,
                tercero=form.cleaned_data['tercero'],
                motivo=form.cleaned_data['motivo'],
                observaciones=form.cleaned_data['observaciones'])
        except (ValueError, StockInsuficiente, BodegaNoPermitida) as e:
            if str(e):
                messages.error(request, str(e))
        else:
            messages.success(request, f'Salida #{salida.pk} registrada.')
            return redirect('log_salidas_detalle', pk=salida.pk)
    return render(request, 'inventario/salida_form.html', {
        'form': form,
        'materiales': _materiales_para_lineas(),
        'grados': GRADOS,
        'lineas_previas': _lineas_previas_material(
            request.POST if request.method == 'POST' else None),
    })


@solo_logistica
def salida_detalle(request, pk):
    salida = get_object_or_404(
        Salida.objects.select_related('bodega', 'tercero', 'creado_por')
        .prefetch_related('lineas__item__categoria'), pk=pk)
    return render(request, 'inventario/salida_detalle.html', {'salida': salida})


@solo_logistica
def traslados_lista(request):
    lista = (Traslado.objects
             .select_related('bodega_origen', 'bodega_destino', 'creado_por')
             .annotate(n_lineas=models.Count('lineas', distinct=True),
                       unidades=Coalesce(models.Sum('lineas__cantidad'), 0))
             .order_by('-creado_en'))
    return render(request, 'inventario/traslados_lista.html', {'traslados': lista})


@solo_logistica
def traslado_nuevo(request):
    form = TrasladoForm(request.POST or None, usuario=request.user)
    if request.method == 'POST':
        try:
            if not form.is_valid():
                _form_a_messages(request, form)
                raise ValueError('')
            # Solo el origen: el destino puede ser cualquier bodega.
            exigir_bodega(request.user, form.cleaned_data['bodega_origen'],
                          accion='trasladar desde')
            lineas = parsear_lineas_material(request.POST)
            traslado = registrar_traslado(
                bodega_origen=form.cleaned_data['bodega_origen'],
                bodega_destino=form.cleaned_data['bodega_destino'],
                lineas=lineas, usuario=request.user,
                observaciones=form.cleaned_data['observaciones'])
        except (ValueError, StockInsuficiente, BodegaNoPermitida) as e:
            if str(e):
                messages.error(request, str(e))
        else:
            messages.success(request, f'Traslado #{traslado.pk} registrado.')
            return redirect('log_traslados_detalle', pk=traslado.pk)
    return render(request, 'inventario/traslado_form.html', {
        'form': form,
        'materiales': _materiales_para_lineas(),
        'grados': GRADOS,
        'lineas_previas': _lineas_previas_material(
            request.POST if request.method == 'POST' else None),
    })


@solo_logistica
def traslado_detalle(request, pk):
    traslado = get_object_or_404(
        Traslado.objects.select_related('bodega_origen', 'bodega_destino',
                                        'creado_por')
        .prefetch_related('lineas__item__categoria'), pk=pk)
    return render(request, 'inventario/traslado_detalle.html',
                  {'traslado': traslado})


# ---------------------------------------------------------------------------
# Préstamos y devoluciones (Fase 5)
# ---------------------------------------------------------------------------

@solo_logistica
def prestamos_lista(request):
    lista = (Prestamo.objects.select_related('tercero', 'creado_por')
             .annotate(prestado=Coalesce(models.Sum('lineas__cantidad_prestada'), 0),
                       pendiente_total=Coalesce(
                           models.Sum(models.F('lineas__cantidad_prestada')
                                      - models.F('lineas__cantidad_devuelta')), 0))
             .order_by('-creado_en'))
    return render(request, 'inventario/prestamos_lista.html', {'prestamos': lista})


@solo_logistica
def prestamo_nuevo(request):
    form = PrestamoForm(request.POST or None)
    if request.method == 'POST':
        try:
            if not form.is_valid():
                _form_a_messages(request, form)
                raise ValueError('')  # cae al re-render conservando las líneas
            lineas = parsear_lineas_material(request.POST, con_bodega=True)
            prestamo = crear_prestamo(
                tercero=form.cleaned_data['tercero'],
                fecha_compromiso=form.cleaned_data['fecha_compromiso'],
                lineas=lineas, usuario=request.user,
                direccion=form.cleaned_data['direccion'],
                observaciones=form.cleaned_data['observaciones'])
        except (ValueError, StockInsuficiente) as e:
            if str(e):
                messages.error(request, str(e))
        else:
            messages.success(request, f'Préstamo #{prestamo.pk} registrado.')
            return redirect('log_prestamos_detalle', pk=prestamo.pk)
    return render(request, 'inventario/prestamo_form.html', {
        'form': form,
        'materiales': _materiales_para_lineas(),
        'grados': GRADOS,
        'bodegas': Bodega.objects.filter(activa=True).order_by('nombre'),
        'lineas_previas': _lineas_previas_material(
            request.POST if request.method == 'POST' else None,
            con_bodega=True),
    })


@solo_logistica
def prestamo_detalle(request, pk):
    prestamo = get_object_or_404(
        Prestamo.objects.select_related('tercero', 'creado_por')
        .prefetch_related('lineas__item__categoria', 'lineas__bodega',
                          'devoluciones__creado_por',
                          'devoluciones__movimientos__item__categoria',
                          'devoluciones__movimientos__bodega'), pk=pk)
    return render(request, 'inventario/prestamo_detalle.html',
                  {'prestamo': prestamo})


@require_POST
@solo_logistica
def prestamo_devolver(request, pk):
    """Devolución (parcial o total) desde el modal del detalle. El POST trae
    listas paralelas `dev_linea_id`/`dev_cantidad` (una fila por línea del
    préstamo); las cantidades vacías o en 0 se ignoran — devolver "algo de
    algunas líneas" es el caso normal de la devolución parcial."""
    prestamo = get_object_or_404(Prestamo, pk=pk)
    lineas_por_pk = {str(l.pk): l for l in prestamo.lineas.all()}
    lineas = []
    try:
        for linea_id, cant in zip(request.POST.getlist('dev_linea_id'),
                                  request.POST.getlist('dev_cantidad')):
            cant = cant.strip()
            if not cant or cant == '0':
                continue
            linea = lineas_por_pk.get(linea_id)
            if linea is None:
                raise ErrorDevolucion('La línea no pertenece a este préstamo.')
            try:
                cantidad = int(cant)
            except ValueError:
                raise ErrorDevolucion(
                    f'Cantidad inválida para "{linea.item.nombre}".')
            lineas.append((linea, cantidad))
        devolucion = registrar_devolucion(
            prestamo=prestamo, lineas=lineas, usuario=request.user,
            observaciones=request.POST.get('observaciones', ''))
    except (ErrorDevolucion, StockInsuficiente) as e:
        messages.error(request, str(e))
    else:
        unidades = sum(c for _, c in lineas)
        messages.success(
            request,
            f'Devolución #{devolucion.pk} registrada ({unidades} unidades).')
    return redirect('log_prestamos_detalle', pk=prestamo.pk)


# ---------------------------------------------------------------------------
# Kardex, ledger global y ajustes
# ---------------------------------------------------------------------------

def _fecha_get(request, nombre):
    try:
        return date.fromisoformat(request.GET.get(nombre, ''))
    except ValueError:
        return None


@solo_logistica
def item_kardex(request, pk):
    """Kardex de un artículo: movimientos en orden cronológico con el saldo
    por fila que dejó cada uno (`saldo_resultante`). Filtros server-side por
    bodega y rango de fechas (van por GET, son compartibles por URL)."""
    item = get_object_or_404(Item, pk=pk)
    bodega_id = request.GET.get('bodega', '')
    bodega = Bodega.objects.filter(pk=bodega_id).first() if bodega_id.isdigit() else None
    desde, hasta = _fecha_get(request, 'desde'), _fecha_get(request, 'hasta')
    movimientos = (kardex(item, bodega=bodega, desde=desde, hasta=hasta)
                   .select_related('creado_por'))
    return render(request, 'inventario/kardex.html', {
        'item': item,
        'movimientos': movimientos,
        # Todas las bodegas (también inactivas): pueden tener historial.
        'bodegas': Bodega.objects.order_by('nombre'),
        'bodega_sel': bodega,
        'desde': desde,
        'hasta': hasta,
    })


# El ledger crece sin tope; la vista muestra los últimos N y el export de la
# Fase 6 será la vía para extraer el histórico completo con filtros.
MOVIMIENTOS_MAX_FILAS = 500


@solo_logistica
def movimientos(request):
    lista = (Movimiento.objects
             .select_related('item', 'item__categoria', 'bodega', 'creado_por')
             .order_by('-creado_en', '-id')[:MOVIMIENTOS_MAX_FILAS])
    return render(request, 'inventario/movimientos.html', {
        'movimientos': lista,
        'tope': MOVIMIENTOS_MAX_FILAS,
        'tipos': Movimiento.Tipo.choices,
        'grados': GRADOS,
    })


@require_POST
@solo_logistica
def ajuste_crear(request):
    """Ajuste manual desde el modal de existencias. `registrar_ajuste` recibe
    la cantidad FINAL (absoluta) y exige motivo; el delta y el tipo de
    movimiento (AJUSTE_POS/NEG) los resuelve el servicio."""
    item_id = request.POST.get('item_id', '')
    bodega_id = request.POST.get('bodega_id', '')
    if not (item_id.isdigit() and bodega_id.isdigit()):
        messages.error(request, 'Ajuste inválido: faltan el artículo o la bodega.')
        return redirect('log_stock')
    item = get_object_or_404(Item, pk=item_id)
    bodega = get_object_or_404(Bodega, pk=bodega_id)
    try:
        nueva_cantidad = int(request.POST.get('nueva_cantidad', '').strip())
    except ValueError:
        messages.error(request, 'La nueva cantidad debe ser un número entero.')
        return redirect('log_stock')
    try:
        registrar_ajuste(item=item, bodega=bodega,
                         nueva_cantidad=nueva_cantidad, usuario=request.user,
                         motivo=request.POST.get('motivo', ''))
        messages.success(
            request,
            f'Stock de "{item.nombre}" en "{bodega.nombre}" ajustado a {nueva_cantidad}.')
    except ValueError as e:
        messages.error(request, str(e))
    return redirect('log_stock')


# ---------------------------------------------------------------------------
# Exports a Excel (Fase 6)
#
# Self-contained con openpyxl (no importan helpers de otras áreas, mismo
# trade-off que `_generar_excel_viaticos` en financiera): cabecera azul,
# bordes, zebra. Todos POST desde un modal con filtros.
# ---------------------------------------------------------------------------

def _generar_excel(*, titulo, columnas, filas, anchos, formatos=None):
    """Excel genérico en memoria. ``filas`` = lista de listas (una por fila);
    ``formatos`` = {col_1based: number_format} (p. ej. moneda)."""
    wb = Workbook()
    ws = wb.active
    ws.title = titulo
    thin = Side(style='thin', color='000000')
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)
    formatos = formatos or {}

    for ci, nombre in enumerate(columnas, start=1):
        cell = ws.cell(1, ci, nombre)
        cell.font = Font(name='Arial', size=10, bold=True, color='FF1F3864')
        cell.alignment = Alignment(horizontal='center', vertical='center',
                                   wrap_text=True)
        cell.fill = PatternFill('solid', fgColor='FFB8CCE4')
        cell.border = borde
    ws.row_dimensions[1].height = 22

    for ri, fila in enumerate(filas, start=2):
        fondo = 'FFFFFFFF' if ri % 2 == 0 else 'FFF2F6FC'
        for ci, valor in enumerate(fila, start=1):
            cell = ws.cell(ri, ci, valor)
            cell.font = Font(name='Arial', size=10)
            cell.alignment = Alignment(
                horizontal='left' if isinstance(valor, str) else 'right',
                vertical='center')
            cell.fill = PatternFill('solid', fgColor=fondo)
            cell.border = borde
            if ci in formatos:
                cell.number_format = formatos[ci]

    for ci, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = ancho
    ws.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _respuesta_xlsx(excel_bytes, nombre):
    response = HttpResponse(
        excel_bytes,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{nombre}"'
    return response


def _fecha_post(request, nombre):
    try:
        return date.fromisoformat(request.POST.get(nombre, ''))
    except ValueError:
        return None


@require_POST
@solo_logistica
def stock_exportar(request):
    """Existencias pivotadas: una fila por (material, bodega) con los 12 grados
    en columnas, igual que la pantalla y que la hoja de cálculo del usuario.
    Filtro opcional de bodega. `valor_unitario` es referencial (el kardex es de
    cantidades): se exporta junto al valor total estimado de la fila."""
    stocks = _stocks_visibles()
    bodega_id = request.POST.get('bodega', '')
    if bodega_id.isdigit():
        stocks = stocks.filter(bodega_id=bodega_id)

    filas = []
    for fila in agrupar_materiales_por_bodega(_items_activos(), stocks):
        item = fila['muestra']
        vu = item.valor_unitario
        filas.append([
            fila['categoria'].nombre, fila['referencia'], fila['bodega'].nombre,
            item.get_unidad_medida_display(),
            # Celda vacía (no en 0) cuando el material no tiene ese grado: en la
            # hoja se distingue "no existe" de "existe y está agotado".
            *[c['cantidad'] if c['item'] else '' for c in fila['celdas']],
            fila['total'], item.stock_minimo or '',
            vu if vu is not None else '',
            vu * fila['total'] if vu is not None else '',
        ])
    n_grados = len(GRADOS)
    columnas = (['Categoría', 'Referencia', 'Bodega', 'Unidad']
                + [f'{g}°' for g in GRADOS]
                + ['Total', 'Mínimo global', 'Valor unitario', 'Valor total'])
    col_valor_unitario = 4 + n_grados + 3  # tras los 12 grados, total y mínimo
    excel = _generar_excel(
        titulo='Existencias',
        columnas=columnas,
        filas=filas,
        anchos=[20, 26, 18, 12] + [6] * n_grados + [10, 12, 14, 14],
        formatos={col_valor_unitario: '"$"#,##0',
                  col_valor_unitario + 1: '"$"#,##0'})
    hoy = timezone.localdate().strftime('%Y%m%d')
    return _respuesta_xlsx(excel, f'Existencias_{hoy}.xlsx')


@require_POST
@solo_logistica
def movimientos_exportar(request):
    """Histórico completo del ledger (SIN el cap de 500 de la vista), con
    filtros de tipo (checkboxes) y rango de fechas. En orden cronológico:
    así el saldo por fila se lee como un kardex."""
    qs = (Movimiento.objects
          .select_related('item', 'item__categoria', 'bodega', 'creado_por')
          .order_by('creado_en', 'id'))
    tipos_validos = set(Movimiento.Tipo.values)
    tipos = [t for t in request.POST.getlist('tipos') if t in tipos_validos]
    if tipos:
        qs = qs.filter(tipo__in=tipos)
    desde, hasta = _fecha_post(request, 'desde'), _fecha_post(request, 'hasta')
    if desde:
        qs = qs.filter(creado_en__date__gte=desde)
    if hasta:
        qs = qs.filter(creado_en__date__lte=hasta)

    filas = [[
        timezone.localtime(m.creado_en).strftime('%d/%m/%Y %H:%M'),
        m.get_tipo_display(), m.item.categoria.nombre, m.item.referencia,
        m.item.grado_display, m.bodega.nombre,
        m.delta, m.saldo_resultante, m.detalle,
        m.creado_por.username if m.creado_por else '',
    ] for m in qs]
    excel = _generar_excel(
        titulo='Movimientos',
        columnas=['Fecha', 'Tipo', 'Categoría', 'Referencia', 'Grado',
                  'Bodega', 'Cantidad', 'Saldo', 'Detalle', 'Registró'],
        filas=filas,
        anchos=[16, 24, 20, 26, 8, 18, 10, 10, 40, 14])
    partes = [desde.strftime('%Y%m%d') if desde else 'inicio',
              hasta.strftime('%Y%m%d') if hasta else 'fin']
    return _respuesta_xlsx(excel, f'Movimientos_{partes[0]}_{partes[1]}.xlsx')


@require_POST
@solo_logistica
def prestamos_exportar(request):
    """Préstamos con totales por documento (mismas annotations que la lista).
    Filtros: dirección y estado (checkboxes; vacío = todos) y "solo vencidos"."""
    qs = (Prestamo.objects.select_related('creado_por')
          .annotate(prestado=Coalesce(models.Sum('lineas__cantidad_prestada'), 0),
                    devuelto=Coalesce(models.Sum('lineas__cantidad_devuelta'), 0))
          .order_by('-creado_en'))
    direcciones = [d for d in request.POST.getlist('direcciones')
                   if d in set(Prestamo.Direccion.values)]
    if direcciones:
        qs = qs.filter(direccion__in=direcciones)
    estados = [e for e in request.POST.getlist('estados')
               if e in set(Prestamo.Estado.values)]
    if estados:
        qs = qs.filter(estado__in=estados)
    if request.POST.get('solo_vencidos'):
        qs = (qs.exclude(estado=Prestamo.Estado.CERRADO)
              .filter(fecha_compromiso__lt=timezone.localdate()))

    filas = [[
        p.pk, timezone.localtime(p.creado_en).strftime('%d/%m/%Y'),
        p.get_direccion_display(), p.tercero_nombre, p.tercero_documento,
        p.fecha_compromiso.strftime('%d/%m/%Y'), p.get_estado_display(),
        'Sí' if p.vencido else 'No',
        p.prestado, p.devuelto, p.prestado - p.devuelto,
        p.creado_por.username if p.creado_por else '',
    ] for p in qs]
    excel = _generar_excel(
        titulo='Préstamos',
        columnas=['#', 'Fecha', 'Dirección', 'Tercero', 'Documento',
                  'Compromiso', 'Estado', 'Vencido', 'Prestado', 'Devuelto',
                  'Pendiente', 'Registró'],
        filas=filas,
        anchos=[6, 12, 16, 28, 14, 12, 14, 9, 10, 10, 10, 14])
    hoy = timezone.localdate().strftime('%Y%m%d')
    return _respuesta_xlsx(excel, f'Prestamos_{hoy}.xlsx')
