"""Tests del área financiera.

Fase 3: enrutado por subdominio, gating de acceso (superusuario / grupo
`area:financiera` / otros) y que el área programación sigue intacta tras el
refactor del chrome.

Fase 4: gestión de viáticos (listar/devolver/aprobar/pagar/editar), el badge de
pendientes y el flujo cruzado programación ↔ financiera.
"""
import shutil
import tempfile
from datetime import date

from django.test import TestCase, Client, override_settings
from django.contrib.auth.models import User, Group
from django.core.files.uploadedfile import SimpleUploadedFile

from core.areas import GRUPO_STAFF_FINANCIERA, GRUPO_STAFF_PROGRAMACION
from programacion.configuracion.models import Colegio, Profesor
from programacion.viaticos.models import GastoViatico, SolicitudViatico, SoportePago

# Soportes en disco local aislado en tmp para los tests: NUNCA tocar Supabase.
_STORAGE_LOCAL = {
    'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
    'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'},
}
_MEDIA_TMP_FIN = tempfile.mkdtemp()
_MEDIA_TMP_FIN_LEG = tempfile.mkdtemp()


class FinancieraAccesoTest(TestCase):

    def setUp(self):
        self.client = Client(HTTP_HOST='financiera.testserver')
        # Los grupos los crean las migraciones 0004/0006; en tests existen ya.
        self.grupo_fin = Group.objects.get(name=GRUPO_STAFF_FINANCIERA)
        self.grupo_prog = Group.objects.get(name=GRUPO_STAFF_PROGRAMACION)

    def test_anonimo_redirige_a_login_apex(self):
        r = self.client.get('/')
        self.assertEqual(r.status_code, 302)
        self.assertIn('/usuarios/login/', r['Location'])

    def test_superusuario_ve_inicio(self):
        User.objects.create_superuser(username='admin', password='pass')
        self.client.login(username='admin', password='pass')
        r = self.client.get('/')
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Financiera AAMO')

    def test_staff_financiera_ve_inicio(self):
        u = User.objects.create_user(username='finan', password='pass')
        u.groups.add(self.grupo_fin)
        self.client.login(username='finan', password='pass')
        r = self.client.get('/')
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Financiera AAMO')

    def test_staff_financiera_ve_viaticos_placeholder(self):
        u = User.objects.create_user(username='finan2', password='pass')
        u.groups.add(self.grupo_fin)
        self.client.login(username='finan2', password='pass')
        r = self.client.get('/viaticos/')
        self.assertEqual(r.status_code, 200)

    def test_usuario_solo_programacion_es_redirigido(self):
        """Un usuario sin acceso a financiera va al selector de área del apex,
        no se le hace logout (puede tener otra área)."""
        u = User.objects.create_user(username='prog', password='pass')
        u.groups.add(self.grupo_prog)
        self.client.login(username='prog', password='pass')
        r = self.client.get('/')
        self.assertEqual(r.status_code, 302)
        # Redirige al apex (host sin el subdominio financiera).
        self.assertNotIn('financiera.testserver', r['Location'])


class SubdominioDesconocidoTest(TestCase):

    def test_subdominio_no_registrado_da_404(self):
        # 'logistica' ya es un área registrada (Fase 1 del inventario); se usa un
        # subdominio que no existe en AREAS.
        c = Client(HTTP_HOST='desconocida.testserver')
        r = c.get('/')
        self.assertEqual(r.status_code, 404)


class ProgramacionIntactaTest(TestCase):
    """Tras el refactor del chrome (base_chrome.html), el área programación
    sigue cargando con su sidebar."""

    def setUp(self):
        self.client = Client(HTTP_HOST='programacion.testserver')

    def test_home_programacion_superusuario_ok(self):
        User.objects.create_superuser(username='admin', password='pass')
        self.client.login(username='admin', password='pass')
        r = self.client.get('/')
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Programación AAMO')


class FinancieraGestionViaticosTest(TestCase):
    """Gestión de viáticos por financiera (Fase 4): devolver / aprobar / pagar /
    editar, badge de pendientes y aislamiento entre áreas."""

    def setUp(self):
        self.client = Client(HTTP_HOST='financiera.testserver')
        self.admin = User.objects.create_superuser('fin_admin', password='pass')
        self.profesor = Profesor.objects.create(
            nombre='Ana', apellido='Gómez', documento='555', cuenta_bancaria='111-222',
        )
        self.colegio = Colegio.objects.create(
            codigo='C-9', nombre='Colegio Sur', departamento='Valle', ciudad='Cali',
        )

    def _solicitud(self, estado=SolicitudViatico.Estado.ENVIADA, **kwargs):
        s = SolicitudViatico(
            profesor=self.profesor, colegio=self.colegio,
            fecha_viaje=date(2026, 6, 1), fecha_regreso=date(2026, 6, 3),
            estado=estado, creado_por=self.admin, **kwargs,
        )
        s.aplicar_snapshot()
        s.save()
        GastoViatico.objects.create(solicitud=s, nombre='Bus', valor=10000, orden=0)
        return s

    def _login_financiera(self):
        self.client.login(username='fin_admin', password='pass')

    # ── Lista / badge ───────────────────────────────────────
    def test_lista_muestra_solicitud(self):
        self._solicitud()
        self._login_financiera()
        r = self.client.get('/viaticos/')
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Ana Gómez')

    def test_badge_cuenta_solo_enviadas(self):
        self._solicitud(estado=SolicitudViatico.Estado.ENVIADA)
        self._solicitud(estado=SolicitudViatico.Estado.APROBADA)
        self._login_financiera()
        r = self.client.get('/')
        self.assertEqual(r.context['viaticos_pendientes_count'], 1)

    def test_badge_incluye_legalizacion_enviada(self):
        """El badge suma lo pendiente de gestión: ENVIADA + LEG_ENVIADA (no el resto)."""
        self._solicitud(estado=SolicitudViatico.Estado.ENVIADA)
        self._solicitud(estado=SolicitudViatico.Estado.LEG_ENVIADA)
        self._solicitud(estado=SolicitudViatico.Estado.LEG_DEVUELTA)
        self._solicitud(estado=SolicitudViatico.Estado.FINALIZADA)
        self._login_financiera()
        r = self.client.get('/')
        self.assertEqual(r.context['viaticos_pendientes_count'], 2)

    # ── Devolver ────────────────────────────────────────────
    def test_devolver_con_motivo(self):
        s = self._solicitud()
        self._login_financiera()
        r = self.client.post(f'/viaticos/{s.pk}/devolver/', {'motivo_devolucion': 'Falta factura'})
        self.assertEqual(r.status_code, 302)
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.DEVUELTA)
        self.assertEqual(s.motivo_devolucion, 'Falta factura')
        self.assertIsNotNone(s.devuelto_en)
        self.assertEqual(s.gestionado_por, self.admin)

    def test_devolver_sin_motivo_no_cambia(self):
        s = self._solicitud()
        self._login_financiera()
        self.client.post(f'/viaticos/{s.pk}/devolver/', {'motivo_devolucion': '   '})
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.ENVIADA)

    def test_devolver_no_enviada_bloqueada(self):
        s = self._solicitud(estado=SolicitudViatico.Estado.APROBADA)
        self._login_financiera()
        self.client.post(f'/viaticos/{s.pk}/devolver/', {'motivo_devolucion': 'X'})
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.APROBADA)

    # ── Aprobar / pagar ─────────────────────────────────────
    def test_aprobar_enviada(self):
        s = self._solicitud()
        self._login_financiera()
        self.client.post(f'/viaticos/{s.pk}/aprobar/')
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.APROBADA)
        self.assertIsNotNone(s.aprobado_en)

    def test_aprobar_no_enviada_bloqueada(self):
        s = self._solicitud(estado=SolicitudViatico.Estado.APROBADA)
        self._login_financiera()
        self.client.post(f'/viaticos/{s.pk}/aprobar/')
        s.refresh_from_db()
        self.assertIsNone(s.aprobado_en)

    def test_pagar_aprobada(self):
        s = self._solicitud(estado=SolicitudViatico.Estado.APROBADA)
        self._login_financiera()
        self.client.post(f'/viaticos/{s.pk}/pagar/')
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.PAGADA)
        self.assertIsNotNone(s.pagado_en)

    def test_pagar_no_aprobada_bloqueada(self):
        s = self._solicitud(estado=SolicitudViatico.Estado.ENVIADA)
        self._login_financiera()
        self.client.post(f'/viaticos/{s.pk}/pagar/')
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.ENVIADA)

    def test_get_no_dispara_accion(self):
        """Las acciones son POST-only (require_POST): un GET no cambia el estado."""
        s = self._solicitud()
        self._login_financiera()
        r = self.client.get(f'/viaticos/{s.pk}/aprobar/')
        self.assertEqual(r.status_code, 405)
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.ENVIADA)

    # ── Editar ──────────────────────────────────────────────
    def test_editar_aprobada_reemplaza_gastos(self):
        s = self._solicitud(estado=SolicitudViatico.Estado.APROBADA)
        self._login_financiera()
        r = self.client.post(f'/viaticos/{s.pk}/editar/', {
            'profesor': self.profesor.pk, 'colegio': self.colegio.pk,
            'fecha_viaje': '2026-06-01', 'fecha_regreso': '2026-06-05',
            'gasto_nombre': ['Bus', 'Hotel'], 'gasto_valor': ['10000', '40000'],
        })
        self.assertEqual(r.status_code, 302)
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.APROBADA)  # editar no cambia estado
        self.assertEqual(s.total, 50000)
        self.assertEqual(s.gastos.count(), 2)

    def test_editar_pagada_bloqueada(self):
        s = self._solicitud(estado=SolicitudViatico.Estado.PAGADA)
        self._login_financiera()
        r = self.client.get(f'/viaticos/{s.pk}/editar/')
        self.assertEqual(r.status_code, 302)  # terminal → redirige al detalle

    def test_editar_devuelta_bloqueada(self):
        """Una DEVUELTA pertenece a programación; financiera no la edita."""
        s = self._solicitud(estado=SolicitudViatico.Estado.DEVUELTA, motivo_devolucion='X')
        self._login_financiera()
        r = self.client.get(f'/viaticos/{s.pk}/editar/')
        self.assertEqual(r.status_code, 302)

    # ── Aislamiento entre áreas ─────────────────────────────
    def test_usuario_solo_programacion_no_gestiona_financiera(self):
        s = self._solicitud()
        grupo_prog, _ = Group.objects.get_or_create(name=GRUPO_STAFF_PROGRAMACION)
        u = User.objects.create_user('solo_prog', password='pass')
        u.groups.add(grupo_prog)
        self.client.login(username='solo_prog', password='pass')
        r = self.client.post(f'/viaticos/{s.pk}/aprobar/')
        # El middleware lo manda al selector de área del apex (no gestiona).
        self.assertEqual(r.status_code, 302)
        self.assertNotIn('financiera.testserver', r['Location'])
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.ENVIADA)


class FlujoViaticosCruzadoTest(TestCase):
    """Flujo end-to-end cruzando programación ↔ financiera con el mismo dato (BD única)."""

    def setUp(self):
        self.prog = Client(HTTP_HOST='programacion.testserver')
        self.fin = Client(HTTP_HOST='financiera.testserver')
        self.admin = User.objects.create_superuser('jefe', password='pass')
        self.profesor = Profesor.objects.create(
            nombre='Leo', apellido='Díaz', documento='777', cuenta_bancaria='333-444',
        )
        self.colegio = Colegio.objects.create(
            codigo='C-3', nombre='Colegio Centro', departamento='Cundinamarca', ciudad='Bogotá',
        )

    def test_enviar_devolver_reenviar_aprobar_pagar(self):
        self.prog.login(username='jefe', password='pass')
        self.fin.login(username='jefe', password='pass')

        # 1) Programación crea/envía
        self.prog.post('/viaticos/crear/', {
            'profesor': self.profesor.pk, 'colegio': self.colegio.pk,
            'fecha_viaje': '2026-06-01', 'fecha_regreso': '2026-06-03',
            'gasto_nombre': ['Bus'], 'gasto_valor': ['10000'],
        })
        s = SolicitudViatico.objects.get()
        self.assertEqual(s.estado, SolicitudViatico.Estado.ENVIADA)

        # 2) Financiera devuelve con motivo
        self.fin.post(f'/viaticos/{s.pk}/devolver/', {'motivo_devolucion': 'Falta hotel'})
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.DEVUELTA)

        # 3) Programación ve el motivo y reenvía
        r = self.prog.get(f'/viaticos/{s.pk}/')
        self.assertContains(r, 'Falta hotel')
        self.prog.post(f'/viaticos/{s.pk}/editar/', {
            'profesor': self.profesor.pk, 'colegio': self.colegio.pk,
            'fecha_viaje': '2026-06-01', 'fecha_regreso': '2026-06-04',
            'gasto_nombre': ['Bus', 'Hotel'], 'gasto_valor': ['10000', '50000'],
            'reenviar': '1',
        })
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.ENVIADA)
        self.assertEqual(s.motivo_devolucion, '')
        self.assertEqual(s.total, 60000)

        # 4) Financiera aprueba y paga
        self.fin.post(f'/viaticos/{s.pk}/aprobar/')
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.APROBADA)
        self.fin.post(f'/viaticos/{s.pk}/pagar/')
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.PAGADA)

        # 5) Terminal: programación ya no puede editar
        r = self.prog.get(f'/viaticos/{s.pk}/editar/')
        self.assertEqual(r.status_code, 302)


@override_settings(MEDIA_ROOT=_MEDIA_TMP_FIN, STORAGES=_STORAGE_LOCAL)
class FinancieraSoporteTest(TestCase):
    """Fase 3/4: financiera sube/elimina/descarga el soporte de pago (solo en PAGADA)."""

    def setUp(self):
        self.client = Client(HTTP_HOST='financiera.testserver')
        self.admin = User.objects.create_superuser('fin_sop', password='pass')
        self.profesor = Profesor.objects.create(
            nombre='Ana', apellido='Gómez', documento='555', cuenta_bancaria='111-222',
        )
        self.colegio = Colegio.objects.create(
            codigo='C-9', nombre='Colegio Sur', departamento='Valle', ciudad='Cali',
        )

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(_MEDIA_TMP_FIN, ignore_errors=True)
        super().tearDownClass()

    def _solicitud(self, estado=SolicitudViatico.Estado.PAGADA):
        s = SolicitudViatico(
            profesor=self.profesor, colegio=self.colegio,
            fecha_viaje=date(2026, 6, 1), fecha_regreso=date(2026, 6, 3),
            estado=estado, creado_por=self.admin,
        )
        s.aplicar_snapshot()
        s.save()
        return s

    def _pdf(self, nombre='soporte.pdf', size=None):
        contenido = b'%PDF-1.4 ' + (b'x' * size if size else b'datos')
        return SimpleUploadedFile(nombre, contenido, content_type='application/pdf')

    def _login(self):
        self.client.login(username='fin_sop', password='pass')

    # ── Subir ───────────────────────────────────────────────
    def test_subir_en_pagada_crea_soporte(self):
        s = self._solicitud(SolicitudViatico.Estado.PAGADA)
        self._login()
        r = self.client.post(f'/viaticos/{s.pk}/soporte/', {'archivo': self._pdf()})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(s.soportes.count(), 1)
        soporte = s.soportes.get()
        self.assertEqual(soporte.subido_por, self.admin)
        # Nombre limpio del upload_to (no el original).
        self.assertIn('viatico-ana-gomez-2026-06-01', soporte.archivo.name)

    def test_subir_en_estado_no_pagada_rechazado(self):
        s = self._solicitud(SolicitudViatico.Estado.APROBADA)
        self._login()
        r = self.client.post(f'/viaticos/{s.pk}/soporte/', {'archivo': self._pdf()})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(s.soportes.count(), 0)

    def test_subir_extension_invalida_rechazada(self):
        s = self._solicitud()
        self._login()
        malo = SimpleUploadedFile('virus.exe', b'MZ', content_type='application/octet-stream')
        self.client.post(f'/viaticos/{s.pk}/soporte/', {'archivo': malo})
        self.assertEqual(s.soportes.count(), 0)

    def test_subir_tamano_excedido_rechazado(self):
        s = self._solicitud()
        self._login()
        grande = self._pdf(size=11 * 1024 * 1024)  # > 10 MB
        self.client.post(f'/viaticos/{s.pk}/soporte/', {'archivo': grande})
        self.assertEqual(s.soportes.count(), 0)

    # ── Gate de área ────────────────────────────────────────
    def test_usuario_programacion_no_sube(self):
        """Un usuario solo-programación no puede subir soporte en financiera (gate)."""
        s = self._solicitud()
        grupo_prog, _ = Group.objects.get_or_create(name=GRUPO_STAFF_PROGRAMACION)
        u = User.objects.create_user('solo_prog_s', password='pass')
        u.groups.add(grupo_prog)
        self.client.login(username='solo_prog_s', password='pass')
        r = self.client.post(f'/viaticos/{s.pk}/soporte/', {'archivo': self._pdf()})
        self.assertEqual(r.status_code, 302)
        self.assertNotIn('financiera.testserver', r['Location'])
        self.assertEqual(s.soportes.count(), 0)

    # ── Eliminar ────────────────────────────────────────────
    def test_eliminar_soporte_borra_fila(self):
        s = self._solicitud()
        self._login()
        self.client.post(f'/viaticos/{s.pk}/soporte/', {'archivo': self._pdf()})
        soporte = s.soportes.get()
        r = self.client.post(f'/viaticos/soporte/{soporte.pk}/eliminar/')
        self.assertEqual(r.status_code, 302)
        self.assertEqual(s.soportes.count(), 0)

    # ── Descargar ───────────────────────────────────────────
    def test_descargar_soporte_financiera(self):
        s = self._solicitud()
        self._login()
        self.client.post(f'/viaticos/{s.pk}/soporte/', {'archivo': self._pdf()})
        soporte = s.soportes.get()
        r = self.client.get(f'/viaticos/soporte/{soporte.pk}/descargar/')
        self.assertEqual(r.status_code, 200)
        self.assertIn('attachment', r['Content-Disposition'])
        self.assertIn('viatico-ana-gomez-2026-06-01.pdf', r['Content-Disposition'])


@override_settings(MEDIA_ROOT=_MEDIA_TMP_FIN_LEG, STORAGES=_STORAGE_LOCAL)
class FinancieraLegalizacionTest(TestCase):
    """Fase legalización: financiera revisa lo enviado por programación
    (devolver con motivo / finalizar) y conserva la gestión del soporte de pago
    hasta FINALIZADA, sin poder tocar los soportes de legalización."""

    def setUp(self):
        self.client = Client(HTTP_HOST='financiera.testserver')
        self.admin = User.objects.create_superuser('fin_leg', password='pass')
        self.client.login(username='fin_leg', password='pass')
        self.profesor = Profesor.objects.create(
            nombre='Ana', apellido='Gómez', documento='555', cuenta_bancaria='111-222',
        )
        self.colegio = Colegio.objects.create(
            codigo='C-9', nombre='Colegio Sur', departamento='Valle', ciudad='Cali',
        )

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(_MEDIA_TMP_FIN_LEG, ignore_errors=True)
        super().tearDownClass()

    def _solicitud(self, estado=SolicitudViatico.Estado.LEG_ENVIADA, **kwargs):
        s = SolicitudViatico(
            profesor=self.profesor, colegio=self.colegio,
            fecha_viaje=date(2026, 6, 1), fecha_regreso=date(2026, 6, 3),
            estado=estado, creado_por=self.admin, **kwargs,
        )
        s.aplicar_snapshot()
        s.save()
        return s

    def _pdf(self, nombre='archivo.pdf'):
        return SimpleUploadedFile(nombre, b'%PDF-1.4 datos', content_type='application/pdf')

    # ── Devolver legalización ───────────────────────────────
    def test_devolver_legalizacion_con_motivo(self):
        s = self._solicitud()
        r = self.client.post(f'/viaticos/{s.pk}/legalizacion/devolver/',
                             {'motivo_devolucion': 'Falta recibo de hotel'})
        self.assertEqual(r.status_code, 302)
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.LEG_DEVUELTA)
        self.assertEqual(s.motivo_devolucion, 'Falta recibo de hotel')
        self.assertIsNotNone(s.legalizacion_devuelta_en)
        self.assertEqual(s.gestionado_por, self.admin)

    def test_devolver_legalizacion_sin_motivo_no_cambia(self):
        s = self._solicitud()
        self.client.post(f'/viaticos/{s.pk}/legalizacion/devolver/', {'motivo_devolucion': '  '})
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.LEG_ENVIADA)

    def test_devolver_legalizacion_en_pagada_bloqueada(self):
        s = self._solicitud(SolicitudViatico.Estado.PAGADA)
        self.client.post(f'/viaticos/{s.pk}/legalizacion/devolver/', {'motivo_devolucion': 'X'})
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.PAGADA)

    # ── Finalizar ───────────────────────────────────────────
    def test_finalizar_leg_enviada(self):
        s = self._solicitud()
        self.client.post(f'/viaticos/{s.pk}/finalizar/')
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.FINALIZADA)
        self.assertIsNotNone(s.finalizado_en)
        self.assertEqual(s.gestionado_por, self.admin)

    def test_finalizar_en_pagada_bloqueada(self):
        """Sin legalización enviada no hay nada que cerrar."""
        s = self._solicitud(SolicitudViatico.Estado.PAGADA)
        self.client.post(f'/viaticos/{s.pk}/finalizar/')
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.PAGADA)

    # ── Soporte de pago en los estados nuevos ───────────────
    def test_subir_soporte_pago_en_leg_enviada_permitido(self):
        s = self._solicitud(SolicitudViatico.Estado.LEG_ENVIADA)
        self.client.post(f'/viaticos/{s.pk}/soporte/', {'archivo': self._pdf()})
        soporte = s.soportes.get()
        self.assertEqual(soporte.tipo, SoportePago.Tipo.PAGO)

    def test_subir_soporte_pago_en_finalizada_rechazado(self):
        s = self._solicitud(SolicitudViatico.Estado.FINALIZADA)
        self.client.post(f'/viaticos/{s.pk}/soporte/', {'archivo': self._pdf()})
        self.assertEqual(s.soportes.count(), 0)

    def test_eliminar_soporte_legalizacion_da_404(self):
        """Financiera no puede borrar soportes de legalización (son de programación)."""
        s = self._solicitud()
        soporte = SoportePago.objects.create(
            solicitud=s, tipo=SoportePago.Tipo.LEGALIZACION, archivo=self._pdf('leg.pdf'),
        )
        r = self.client.post(f'/viaticos/soporte/{soporte.pk}/eliminar/')
        self.assertEqual(r.status_code, 404)
        self.assertEqual(s.soportes.count(), 1)

    def test_eliminar_soporte_pago_en_finalizada_bloqueado(self):
        s = self._solicitud(SolicitudViatico.Estado.FINALIZADA)
        soporte = SoportePago.objects.create(solicitud=s, archivo=self._pdf('pago.pdf'))
        self.client.post(f'/viaticos/soporte/{soporte.pk}/eliminar/')
        self.assertEqual(s.soportes.count(), 1)

    def test_descarga_soporte_legalizacion_en_financiera(self):
        """Solo lectura sí: financiera ve/descarga los soportes de legalización."""
        s = self._solicitud()
        soporte = SoportePago.objects.create(
            solicitud=s, tipo=SoportePago.Tipo.LEGALIZACION, archivo=self._pdf('leg.pdf'),
        )
        r = self.client.get(f'/viaticos/soporte/{soporte.pk}/descargar/')
        self.assertEqual(r.status_code, 200)

    # ── Flujo completo del ciclo de legalización ────────────
    def test_flujo_legalizacion_completo(self):
        """PAGADA → LEG_ENVIADA → LEG_DEVUELTA → LEG_ENVIADA → FINALIZADA, cruzando áreas."""
        prog = Client(HTTP_HOST='programacion.testserver')
        prog.login(username='fin_leg', password='pass')

        s = self._solicitud(SolicitudViatico.Estado.PAGADA)
        prog.post(f'/viaticos/{s.pk}/legalizacion/soporte/', {'archivo': self._pdf()})
        prog.post(f'/viaticos/{s.pk}/legalizacion/enviar/')
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.LEG_ENVIADA)

        self.client.post(f'/viaticos/{s.pk}/legalizacion/devolver/',
                         {'motivo_devolucion': 'Falta un recibo'})
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.LEG_DEVUELTA)

        # Programación ve el motivo, corrige y reenvía.
        r = prog.get(f'/viaticos/{s.pk}/')
        self.assertContains(r, 'Falta un recibo')
        prog.post(f'/viaticos/{s.pk}/legalizacion/soporte/', {'archivo': self._pdf('recibo2.pdf')})
        prog.post(f'/viaticos/{s.pk}/legalizacion/enviar/')
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.LEG_ENVIADA)
        self.assertEqual(s.motivo_devolucion, '')

        self.client.post(f'/viaticos/{s.pk}/finalizar/')
        s.refresh_from_db()
        self.assertEqual(s.estado, SolicitudViatico.Estado.FINALIZADA)


class FinancieraExportTest(TestCase):
    """Fase 5: exportación a Excel con filtros de estado y rango de fecha de viaje."""

    XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    def setUp(self):
        self.client = Client(HTTP_HOST='financiera.testserver')
        self.admin = User.objects.create_superuser('fin_exp', password='pass')
        self.profesor = Profesor.objects.create(
            nombre='Ana', apellido='Gómez', documento='555', cuenta_bancaria='111-222',
        )
        self.colegio = Colegio.objects.create(
            codigo='C-9', nombre='Colegio Sur', departamento='Valle', ciudad='Cali',
        )

    def _solicitud(self, estado, fecha_viaje=date(2026, 6, 1)):
        s = SolicitudViatico(
            profesor=self.profesor, colegio=self.colegio,
            fecha_viaje=fecha_viaje, fecha_regreso=fecha_viaje,
            estado=estado, creado_por=self.admin,
        )
        s.aplicar_snapshot()
        s.save()
        GastoViatico.objects.create(solicitud=s, nombre='Bus', valor=10000, orden=0)
        return s

    def test_export_devuelve_xlsx(self):
        self._solicitud(SolicitudViatico.Estado.APROBADA)
        self.client.login(username='fin_exp', password='pass')
        r = self.client.post('/viaticos/exportar/', {'estados': ['APROBADA']})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], self.XLSX)
        self.assertIn('attachment', r['Content-Disposition'])
        self.assertTrue(r.content)  # bytes del .xlsx

    def test_export_get_no_permitido(self):
        self.client.login(username='fin_exp', password='pass')
        r = self.client.get('/viaticos/exportar/')
        self.assertEqual(r.status_code, 405)  # require_POST

    def test_export_default_solo_aprobada(self):
        """Sin estados marcados, el default es APROBADA → solo esas filas (1 cabecera + 1 dato)."""
        from openpyxl import load_workbook
        from io import BytesIO
        self._solicitud(SolicitudViatico.Estado.APROBADA)
        self._solicitud(SolicitudViatico.Estado.ENVIADA)
        self._solicitud(SolicitudViatico.Estado.PAGADA)
        self.client.login(username='fin_exp', password='pass')
        r = self.client.post('/viaticos/exportar/', {})  # sin 'estados'
        wb = load_workbook(BytesIO(r.content))
        ws = wb.active
        self.assertEqual(ws.max_row, 2)  # cabecera + 1 fila (solo la APROBADA)

    def test_export_filtra_por_rango_fecha_viaje(self):
        from openpyxl import load_workbook
        from io import BytesIO
        self._solicitud(SolicitudViatico.Estado.APROBADA, fecha_viaje=date(2026, 6, 1))
        self._solicitud(SolicitudViatico.Estado.APROBADA, fecha_viaje=date(2026, 8, 1))
        self.client.login(username='fin_exp', password='pass')
        r = self.client.post('/viaticos/exportar/', {
            'estados': ['APROBADA'], 'fecha_desde': '2026-07-01', 'fecha_hasta': '2026-09-01',
        })
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.active.max_row, 2)  # solo la de agosto
