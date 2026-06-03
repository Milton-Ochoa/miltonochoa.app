"""Vistas del área **financiera** para la gestión de solicitudes de viáticos.

Financiera **no crea** solicitudes (eso es programación): aquí se **ven** y se
**gestionan** las que llegan. Acciones, según la matriz de permisos:

  - ENVIADA  → devolver (con motivo) / aprobar / editar
  - DEVUELTA → financiera NO actúa (vuelve a programación para corregir y reenviar)
  - APROBADA → editar / pagar
  - PAGADA   → terminal, solo lectura para todos

Los modelos viven en `programacion.viaticos` (BD única); esta app solo los importa.
Se reutiliza el form, el parseo de gastos y el contexto del select-autorrelleno de
programación para no duplicar lógica (mismo widget de filas de gastos en la plantilla).

Gate: superusuario o miembro del grupo `area:financiera`
(`core.areas.es_personal_financiera`). El middleware filtra el acceso al subdominio;
el decorador es la segunda barrera por-vista.
"""
import io
import os
from datetime import datetime

from django.contrib import messages
from django.contrib.auth.decorators import user_passes_test
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.areas import es_personal_financiera
from programacion.viaticos.forms import SolicitudViaticoForm
from programacion.viaticos.models import GastoViatico, SolicitudViatico, SoportePago
from programacion.viaticos.views import _contexto_form, _parsear_gastos, _responder_soporte

solo_financiera = user_passes_test(es_personal_financiera, login_url='login')

# Estados en los que financiera puede editar la solicitud (campos + gastos).
# DEVUELTA queda fuera a propósito: pertenece a programación hasta que la reenvíe;
# PAGADA es terminal. (Ver matriz de permisos arriba.)
EDITABLES_FINANCIERA = {SolicitudViatico.Estado.ENVIADA, SolicitudViatico.Estado.APROBADA}

# Restricciones del soporte de pago (solo se sube en PAGADA).
SOPORTE_EXTENSIONES = {'.pdf', '.jpg', '.jpeg', '.png'}
SOPORTE_MAX_BYTES = 10 * 1024 * 1024  # 10 MB


def _validar_soporte(archivo):
    """Valida extensión y tamaño de un soporte; devuelve un mensaje de error o None."""
    ext = os.path.splitext(archivo.name)[1].lower()
    if ext not in SOPORTE_EXTENSIONES:
        return f'Tipo de archivo no permitido ({ext or "sin extensión"}). Usa PDF, JPG o PNG.'
    if archivo.size > SOPORTE_MAX_BYTES:
        return 'El archivo supera el tamaño máximo de 10 MB.'
    return None


@solo_financiera
def fin_home(request):
    """Inicio del área financiera (vacío por ahora)."""
    return render(request, 'financiera/home.html')


@solo_financiera
def fin_viaticos_lista(request):
    """Tabla de todas las solicitudes. Las `ENVIADA` son las pendientes de gestión."""
    solicitudes = (
        SolicitudViatico.objects
        .select_related('profesor', 'colegio')
        .prefetch_related('gastos')
    )
    return render(request, 'financiera/viaticos_lista.html', {'solicitudes': solicitudes})


@solo_financiera
def fin_viaticos_detalle(request, pk):
    """Detalle con gastos, total, observaciones y los datos snapshot del docente/colegio.
    Expone qué acciones permite el estado actual para pintar los botones."""
    solicitud = get_object_or_404(
        SolicitudViatico.objects.select_related('profesor', 'colegio')
        .prefetch_related('gastos', 'soportes', 'soportes__subido_por'),
        pk=pk,
    )
    return render(request, 'financiera/viaticos_detalle.html', {
        'solicitud': solicitud,
        'puede_editar': solicitud.estado in EDITABLES_FINANCIERA,
        'puede_devolver': solicitud.estado == SolicitudViatico.Estado.ENVIADA,
        'puede_aprobar': solicitud.estado == SolicitudViatico.Estado.ENVIADA,
        'puede_pagar': solicitud.estado == SolicitudViatico.Estado.APROBADA,
    })


@solo_financiera
def fin_viaticos_editar(request, pk):
    """Reedita una solicitud `ENVIADA`/`APROBADA` (campos + gastos). No cambia el estado
    ni el snapshot del cliente: el servidor recalcula el snapshot desde la FK."""
    solicitud = get_object_or_404(SolicitudViatico, pk=pk)
    if solicitud.estado not in EDITABLES_FINANCIERA:
        messages.error(request, 'Esta solicitud no se puede editar en su estado actual.')
        return redirect('fin_viaticos_detalle', pk=solicitud.pk)

    if request.method == 'POST':
        form = SolicitudViaticoForm(request.POST, instance=solicitud)
        gastos, errores_gastos = _parsear_gastos(request)
        for e in errores_gastos:
            messages.error(request, e)

        if form.is_valid() and not errores_gastos:
            with transaction.atomic():
                solicitud = form.save(commit=False)
                solicitud.aplicar_snapshot()
                solicitud.save()
                # Reemplazo total de gastos (igual que en programación).
                solicitud.gastos.all().delete()
                GastoViatico.objects.bulk_create(
                    GastoViatico(solicitud=solicitud, **g) for g in gastos
                )
            messages.success(request, 'Solicitud actualizada.')
            return redirect('fin_viaticos_detalle', pk=solicitud.pk)

        return render(request, 'financiera/viaticos_form.html',
                      {**_contexto_form(form, gastos), 'solicitud': solicitud})

    form = SolicitudViaticoForm(instance=solicitud)
    return render(request, 'financiera/viaticos_form.html',
                  {**_contexto_form(form, list(solicitud.gastos.all())), 'solicitud': solicitud})


@solo_financiera
@require_POST
def fin_viaticos_devolver(request, pk):
    """`ENVIADA → DEVUELTA` con un motivo obligatorio (vuelve a programación)."""
    solicitud = get_object_or_404(SolicitudViatico, pk=pk)
    motivo = request.POST.get('motivo_devolucion', '').strip()
    if solicitud.estado != SolicitudViatico.Estado.ENVIADA:
        messages.error(request, 'Solo se puede devolver una solicitud enviada.')
    elif not motivo:
        messages.error(request, 'Indica el motivo de la devolución.')
    else:
        solicitud.estado = SolicitudViatico.Estado.DEVUELTA
        solicitud.motivo_devolucion = motivo
        solicitud.devuelto_en = timezone.now()
        solicitud.gestionado_por = request.user
        solicitud.save()
        messages.success(request, 'Solicitud devuelta a programación.')
    return redirect('fin_viaticos_detalle', pk=solicitud.pk)


@solo_financiera
@require_POST
def fin_viaticos_aprobar(request, pk):
    """`ENVIADA → APROBADA`."""
    solicitud = get_object_or_404(SolicitudViatico, pk=pk)
    if solicitud.estado != SolicitudViatico.Estado.ENVIADA:
        messages.error(request, 'Solo se puede aprobar una solicitud enviada.')
    else:
        solicitud.estado = SolicitudViatico.Estado.APROBADA
        solicitud.aprobado_en = timezone.now()
        solicitud.gestionado_por = request.user
        solicitud.save()
        messages.success(request, 'Solicitud aprobada.')
    return redirect('fin_viaticos_detalle', pk=solicitud.pk)


@solo_financiera
@require_POST
def fin_viaticos_pagar(request, pk):
    """`APROBADA → PAGADA` (estado terminal)."""
    solicitud = get_object_or_404(SolicitudViatico, pk=pk)
    if solicitud.estado != SolicitudViatico.Estado.APROBADA:
        messages.error(request, 'Solo se puede pagar una solicitud aprobada.')
    else:
        solicitud.estado = SolicitudViatico.Estado.PAGADA
        solicitud.pagado_en = timezone.now()
        solicitud.gestionado_por = request.user
        solicitud.save()
        messages.success(request, 'Solicitud marcada como pagada.')
    return redirect('fin_viaticos_detalle', pk=solicitud.pk)


@solo_financiera
@require_POST
def fin_viaticos_subir_soporte(request, pk):
    """Adjunta un soporte de pago a una solicitud `PAGADA` (solo en ese estado)."""
    solicitud = get_object_or_404(SolicitudViatico, pk=pk)
    if solicitud.estado != SolicitudViatico.Estado.PAGADA:
        messages.error(request, 'Solo se puede adjuntar soporte a una solicitud pagada.')
        return redirect('fin_viaticos_detalle', pk=solicitud.pk)

    archivo = request.FILES.get('archivo')
    if not archivo:
        messages.error(request, 'Selecciona un archivo para subir.')
        return redirect('fin_viaticos_detalle', pk=solicitud.pk)

    error = _validar_soporte(archivo)
    if error:
        messages.error(request, error)
        return redirect('fin_viaticos_detalle', pk=solicitud.pk)

    SoportePago.objects.create(
        solicitud=solicitud,
        archivo=archivo,
        nombre_original=archivo.name,
        subido_por=request.user,
    )
    messages.success(request, 'Soporte de pago adjuntado.')
    return redirect('fin_viaticos_detalle', pk=solicitud.pk)


@solo_financiera
@require_POST
def fin_viaticos_eliminar_soporte(request, soporte_id):
    """Elimina un soporte (y su archivo en storage) subido por error."""
    soporte = get_object_or_404(SoportePago.objects.select_related('solicitud'), pk=soporte_id)
    pk = soporte.solicitud_id
    # Borrar primero el archivo del storage (S3/disco), luego la fila.
    soporte.archivo.delete(save=False)
    soporte.delete()
    messages.success(request, 'Soporte eliminado.')
    return redirect('fin_viaticos_detalle', pk=pk)


@solo_financiera
def fin_soporte_descargar(request, soporte_id):
    """Ver (``?inline=1``) o descargar un soporte desde el área financiera.
    Mismo proxy que programación; difiere solo en el gate por urlconf de subdominio."""
    soporte = get_object_or_404(SoportePago, pk=soporte_id)
    return _responder_soporte(soporte, inline=request.GET.get('inline') == '1')


# ══════════════════════════════════════════════════════════════
# EXPORTAR EXCEL DE VIÁTICOS
# ══════════════════════════════════════════════════════════════

def _generar_excel_viaticos(solicitudes):
    """Genera (en memoria) el Excel de las solicitudes dadas.

    Self-contained a propósito (no importa los helpers de ``programacion.exportar``)
    para no acoplar el área financiera a internals de programación; se limita a
    imitar su estilo (cabecera azul, bordes, total como número ``"$"#,##0``).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = 'Viáticos'

    COLS = ['#', 'Docente', 'Cédula', 'Banco', 'N° Cuenta', 'Colegio', 'Código',
            'Fecha viaje', 'Fecha regreso', 'Total', 'Estado', 'Fecha de pago']
    NUM_COLS = len(COLS)
    COL_TOTAL = 10  # índice (1-based) de la columna "Total"

    thin = Side(style='thin', color='000000')
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _celda(row, col, valor='', bold=False, fill=None, color='000000',
               h='center', v='center', fmt=None):
        cell = ws.cell(row, col, valor)
        cell.font = Font(name='Arial', size=10, bold=bold, color=color)
        cell.alignment = Alignment(horizontal=h, vertical=v, wrap_text=True)
        if fill:
            cell.fill = PatternFill('solid', fgColor=fill)
        cell.border = borde
        if fmt:
            cell.number_format = fmt
        return cell

    # Fila 1: cabeceras (azul, como en exportar/pagos)
    for ci, nombre in enumerate(COLS, start=1):
        _celda(1, ci, nombre, bold=True, fill='FFB8CCE4', color='FF1F3864')
    ws.row_dimensions[1].height = 22

    # Filas de datos
    for i, s in enumerate(solicitudes, start=2):
        fill_row = 'FFFFFFFF' if i % 2 == 0 else 'FFF2F6FC'
        _celda(i, 1, s.pk,             fill=fill_row)
        _celda(i, 2, s.docente_nombre, fill=fill_row, h='left')
        _celda(i, 3, s.docente_cedula, fill=fill_row)
        _celda(i, 4, s.docente_banco,  fill=fill_row)
        _celda(i, 5, s.docente_cuenta, fill=fill_row)
        _celda(i, 6, s.colegio_nombre, fill=fill_row, h='left')
        _celda(i, 7, s.colegio_codigo, fill=fill_row)
        _celda(i, 8, s.fecha_viaje.strftime('%d/%m/%Y') if s.fecha_viaje else '', fill=fill_row)
        _celda(i, 9, s.fecha_regreso.strftime('%d/%m/%Y') if s.fecha_regreso else '', fill=fill_row)
        _celda(i, COL_TOTAL, s.total, fill=fill_row, h='right', fmt='"$"#,##0')
        _celda(i, 11, s.get_estado_display(), fill=fill_row)
        _celda(i, 12, s.pagado_en.strftime('%d/%m/%Y') if s.pagado_en else '', fill=fill_row)
        ws.row_dimensions[i].height = 18

    # Anchos de columna
    anchos = [6, 26, 14, 18, 16, 28, 10, 14, 14, 14, 12, 14]
    for ci, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = ancho

    ws.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@solo_financiera
@require_POST
def fin_viaticos_exportar(request):
    """Descarga un ``.xlsx`` de solicitudes filtradas por estado y rango de fecha de viaje.

    Filtros del modal: ``estados`` (checkboxes; default ``['APROBADA']``) y
    ``fecha_desde``/``fecha_hasta`` opcionales (sobre ``fecha_viaje``).
    """
    estados_validos = set(SolicitudViatico.Estado.values)
    estados = [e for e in request.POST.getlist('estados') if e in estados_validos]
    if not estados:
        estados = [SolicitudViatico.Estado.APROBADA]

    solicitudes = (
        SolicitudViatico.objects
        .filter(estado__in=estados)
        .select_related('profesor', 'colegio')
        .prefetch_related('gastos')
    )

    def _fecha(clave):
        try:
            return datetime.strptime(request.POST.get(clave, ''), '%Y-%m-%d').date()
        except ValueError:
            return None

    desde, hasta = _fecha('fecha_desde'), _fecha('fecha_hasta')
    if desde:
        solicitudes = solicitudes.filter(fecha_viaje__gte=desde)
    if hasta:
        solicitudes = solicitudes.filter(fecha_viaje__lte=hasta)

    excel_bytes = _generar_excel_viaticos(solicitudes)

    partes = [desde.strftime('%Y%m%d') if desde else 'inicio',
              hasta.strftime('%Y%m%d') if hasta else 'fin']
    response = HttpResponse(
        excel_bytes,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="Viaticos_{partes[0]}_{partes[1]}.xlsx"'
    return response
