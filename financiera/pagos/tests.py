"""Tests del área financiera — pagos de clases a profesores.

Financiera **solo gestiona filas de lotes ENVIADO** por programación: marca el pago
(fija `fecha_pago`/`marcado_por`), lo desmarca (limpia esos campos y borra soportes, sin
borrar la fila) y sube/elimina soportes. El cálculo semanal y la materialización los
cubren los tests de programación; aquí se valida la gestión propia y el gate de área.
"""
import io
import shutil
import tempfile
from datetime import date, datetime, time, timedelta

from django.test import TestCase, Client, override_settings
from django.contrib.auth.models import User, Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import load_workbook

from core.areas import GRUPO_STAFF_FINANCIERA, GRUPO_STAFF_PROGRAMACION
from programacion.colegios.models import Bloque, Clase, Grado
from programacion.configuracion.models import Colegio, ColegioAnio, Profesor
from programacion.pagos.models import LotePagos, PagoRealizado, SoportePagoProfesor

# Soportes en disco local aislado en tmp: NUNCA tocar Supabase (igual que viáticos).
_STORAGE_LOCAL = {
    'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
    'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'},
}
_MEDIA_TMP_PAGOS = tempfile.mkdtemp()


def _crear_pago(colegio_anio, profesor, *, estado=LotePagos.Estado.ENVIADO, fecha=date(2025, 3, 14)):
    """Crea un lote en el estado dado con una fila base lista para financiera."""
    lote = LotePagos.objects.create(
        fecha_inicio=date(2025, 3, 10), fecha_fin=date(2025, 3, 14), estado=estado)
    return PagoRealizado.objects.create(
        lote=lote, profesor=profesor, colegio=colegio_anio,
        fecha=fecha, horas=2, valor=80000)


class FinPagosTest(TestCase):

    def setUp(self):
        self.client = Client(HTTP_HOST='financiera.testserver')
        self.grupo_fin = Group.objects.get(name=GRUPO_STAFF_FINANCIERA)
        self.grupo_prog = Group.objects.get(name=GRUPO_STAFF_PROGRAMACION)

        colegio = Colegio.objects.create(
            nombre='Colegio Central', departamento='Santander', ciudad='Bucaramanga')
        self.colegio_anio = ColegioAnio.objects.create(colegio=colegio, anio=2025)
        self.profesor = Profesor.objects.create(nombre='Ana', apellido='Pérez')

        self.finan = User.objects.create_user(username='finan', password='pass')
        self.finan.groups.add(self.grupo_fin)

    def _login_financiera(self):
        self.client.login(username='finan', password='pass')

    # ── Acceso ────────────────────────────────────────────────
    def test_lista_200_para_financiera(self):
        self._login_financiera()
        r = self.client.get('/pagos/')
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'tablaPagos')

    def test_lista_rechaza_usuario_solo_programacion(self):
        u = User.objects.create_user(username='prog', password='pass')
        u.groups.add(self.grupo_prog)
        self.client.login(username='prog', password='pass')
        r = self.client.get('/pagos/')
        self.assertNotEqual(r.status_code, 200)  # middleware lo saca del subdominio

    def test_lista_solo_muestra_lotes_enviados(self):
        self._login_financiera()
        # Lote en BORRADOR → financiera no lo ve.
        _crear_pago(self.colegio_anio, self.profesor, estado=LotePagos.Estado.BORRADOR)
        r = self.client.get('/pagos/?semana=2025-03-10&tab=pendiente')
        self.assertNotContains(r, 'Pérez')

    # ── Marcar / desmarcar ────────────────────────────────────
    def test_marcar_fija_fecha_pago(self):
        self._login_financiera()
        pago = _crear_pago(self.colegio_anio, self.profesor)
        r = self.client.post('/pagos/marcar/', {'accion': 'marcar', 'pago_id': pago.id})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json()['ok'])
        pago.refresh_from_db()
        self.assertIsNotNone(pago.fecha_pago)
        self.assertEqual(pago.marcado_por, self.finan)

    def test_marcar_es_idempotente(self):
        self._login_financiera()
        pago = _crear_pago(self.colegio_anio, self.profesor)
        self.client.post('/pagos/marcar/', {'accion': 'marcar', 'pago_id': pago.id})
        self.client.post('/pagos/marcar/', {'accion': 'marcar', 'pago_id': pago.id})
        self.assertEqual(PagoRealizado.objects.count(), 1)  # no duplica filas

    def test_marcar_rechaza_lote_no_enviado(self):
        self._login_financiera()
        pago = _crear_pago(self.colegio_anio, self.profesor, estado=LotePagos.Estado.BORRADOR)
        r = self.client.post('/pagos/marcar/', {'accion': 'marcar', 'pago_id': pago.id})
        self.assertEqual(r.status_code, 400)
        pago.refresh_from_db()
        self.assertIsNone(pago.fecha_pago)

    def test_desmarcar_limpia_pero_conserva_fila(self):
        self._login_financiera()
        pago = _crear_pago(self.colegio_anio, self.profesor)
        pago.fecha_pago = timezone.now()
        pago.marcado_por = self.finan
        pago.save()
        r = self.client.post('/pagos/marcar/', {'accion': 'desmarcar', 'pago_id': pago.id})
        self.assertTrue(r.json()['ok'])
        pago.refresh_from_db()
        self.assertIsNone(pago.fecha_pago)               # ya no pagada
        self.assertEqual(PagoRealizado.objects.count(), 1)  # la fila sigue

    def test_marcar_requiere_financiera(self):
        # Sin login → no debe marcar (redirige a login/apex).
        pago = _crear_pago(self.colegio_anio, self.profesor)
        r = self.client.post('/pagos/marcar/', {'accion': 'marcar', 'pago_id': pago.id})
        self.assertNotEqual(r.status_code, 200)
        pago.refresh_from_db()
        self.assertIsNone(pago.fecha_pago)

    # ── Exportar ──────────────────────────────────────────────
    def test_exportar_devuelve_xlsx(self):
        self._login_financiera()
        r = self.client.post('/pagos/exportar/', {
            'fecha_inicio': '2025-03-10', 'fecha_fin': '2025-03-14', 'tab': 'pendiente',
        })
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r['Content-Type'])
        self.assertIn('attachment', r['Content-Disposition'])

    def _pago_enviado_el(self, dia):
        """Fila enviada por programación en `dia` (fija `LotePagos.enviado_en`)."""
        pago = _crear_pago(self.colegio_anio, self.profesor)
        pago.lote.enviado_en = timezone.make_aware(datetime.combine(dia, time(9, 30)))
        pago.lote.save(update_fields=['enviado_en'])
        return pago

    def test_exportar_incluye_departamento_y_fecha_de_envio(self):
        """El Excel trae el departamento del colegio y la fecha en que programación
        envió el lote; VALOR queda corrido a la columna 10."""
        self._login_financiera()
        self._pago_enviado_el(date(2025, 3, 17))
        r = self.client.post('/pagos/exportar/', {'tab': 'pendiente'})
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws.cell(2, 9).value, 'DEPARTAMENTO')
        self.assertEqual(ws.cell(2, 11).value, 'FECHA DE ENVÍO')
        self.assertEqual(ws.cell(3, 9).value, 'Santander')
        self.assertEqual(ws.cell(3, 11).value, '17/03/2025')
        self.assertEqual(ws.cell(3, 10).value, 80000)

    def test_lista_muestra_fecha_de_envio(self):
        self._login_financiera()
        self._pago_enviado_el(date(2025, 3, 17))
        r = self.client.get('/pagos/?tab=pendiente')
        self.assertContains(r, '17/03/2025')

    # ── Detalle ───────────────────────────────────────────────
    def test_detalle_muestra_datos_y_desglose(self):
        self._login_financiera()
        pago = _crear_pago(self.colegio_anio, self.profesor)
        from programacion.pagos.models import ExtraPago
        ExtraPago.objects.create(pago=pago, concepto='Desplazamiento', valor=15000)
        r = self.client.get(f'/pagos/{pago.pk}/')
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Colegio Central')
        self.assertContains(r, 'Desplazamiento')  # el desglose
        self.assertContains(r, 'Soporte de pago')


@override_settings(MEDIA_ROOT=_MEDIA_TMP_PAGOS, STORAGES=_STORAGE_LOCAL)
class FinPagosSoporteTest(TestCase):
    """Subida/eliminación/descarga de soportes en disco local aislado."""

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(_MEDIA_TMP_PAGOS, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.client = Client(HTTP_HOST='financiera.testserver')
        grupo_fin = Group.objects.get(name=GRUPO_STAFF_FINANCIERA)
        self.finan = User.objects.create_user(username='finan', password='pass')
        self.finan.groups.add(grupo_fin)
        self.client.login(username='finan', password='pass')

        colegio = Colegio.objects.create(
            nombre='Colegio Central', departamento='Santander', ciudad='Bucaramanga')
        colegio_anio = ColegioAnio.objects.create(colegio=colegio, anio=2025)
        profesor = Profesor.objects.create(nombre='Ana', apellido='Pérez')
        self.pago = _crear_pago(colegio_anio, profesor)

    def _archivo(self, nombre='comprobante.pdf', contenido=b'%PDF-1.4 fake'):
        return SimpleUploadedFile(nombre, contenido, content_type='application/pdf')

    def test_subir_soporte_ok(self):
        r = self.client.post(f'/pagos/{self.pago.pk}/soporte/', {'archivo': self._archivo()})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.pago.soportes.count(), 1)
        soporte = self.pago.soportes.first()
        self.assertEqual(soporte.subido_por, self.finan)
        self.assertTrue(soporte.archivo.name.startswith('pagos/pago-ana-perez-2025-03-14'))

    def test_subir_extension_invalida_rechazada(self):
        r = self.client.post(f'/pagos/{self.pago.pk}/soporte/',
                             {'archivo': self._archivo('virus.exe', b'MZ')})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.pago.soportes.count(), 0)

    def test_eliminar_soporte(self):
        self.client.post(f'/pagos/{self.pago.pk}/soporte/', {'archivo': self._archivo()})
        soporte = self.pago.soportes.first()
        r = self.client.post(f'/pagos/soporte/{soporte.pk}/eliminar/')
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.pago.soportes.count(), 0)

    def test_descargar_soporte(self):
        self.client.post(f'/pagos/{self.pago.pk}/soporte/', {'archivo': self._archivo()})
        soporte = self.pago.soportes.first()
        r = self.client.get(f'/pagos/soporte/{soporte.pk}/descargar/')
        self.assertEqual(r.status_code, 200)

    def test_desmarcar_borra_soporte_pero_conserva_fila(self):
        # Subir un soporte y luego desmarcar: el archivo y el soporte se van; la fila queda.
        self.pago.fecha_pago = timezone.now()
        self.pago.save()
        self.client.post(f'/pagos/{self.pago.pk}/soporte/', {'archivo': self._archivo()})
        self.assertEqual(SoportePagoProfesor.objects.count(), 1)
        self.client.post('/pagos/marcar/', {'accion': 'desmarcar', 'pago_id': self.pago.id})
        self.pago.refresh_from_db()
        self.assertIsNone(self.pago.fecha_pago)
        self.assertEqual(PagoRealizado.objects.count(), 1)
        self.assertEqual(SoportePagoProfesor.objects.count(), 0)


@override_settings(MEDIA_ROOT=_MEDIA_TMP_PAGOS, STORAGES=_STORAGE_LOCAL)
class FinPagosLoteNoEnviadoTest(TestCase):
    """Regresión: financiera solo ve lo enviado. Detalle, soportes y descarga deben
    rechazar filas cuyo lote no esté ENVIADO (mismo guard que fin_pagos_marcar)."""

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(_MEDIA_TMP_PAGOS, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.client = Client(HTTP_HOST='financiera.testserver')
        grupo_fin = Group.objects.get(name=GRUPO_STAFF_FINANCIERA)
        self.finan = User.objects.create_user(username='finan', password='pass')
        self.finan.groups.add(grupo_fin)
        self.client.login(username='finan', password='pass')

        colegio = Colegio.objects.create(
            nombre='Colegio Central', departamento='Santander', ciudad='Bucaramanga')
        colegio_anio = ColegioAnio.objects.create(colegio=colegio, anio=2025)
        profesor = Profesor.objects.create(nombre='Ana', apellido='Pérez')
        self.pago = _crear_pago(colegio_anio, profesor, estado=LotePagos.Estado.BORRADOR)

    def _archivo(self):
        return SimpleUploadedFile('comprobante.pdf', b'%PDF-1.4 fake',
                                  content_type='application/pdf')

    def _soporte_borrador(self):
        return SoportePagoProfesor.objects.create(
            pago=self.pago, archivo=self._archivo(),
            nombre_original='comprobante.pdf', subido_por=self.finan)

    def test_detalle_rechaza_lote_borrador(self):
        r = self.client.get(f'/pagos/{self.pago.pk}/')
        self.assertEqual(r.status_code, 302)  # redirect a la lista, no muestra la fila

    def test_subir_soporte_rechaza_lote_borrador(self):
        r = self.client.post(f'/pagos/{self.pago.pk}/soporte/', {'archivo': self._archivo()})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.pago.soportes.count(), 0)

    def test_eliminar_soporte_rechaza_lote_borrador(self):
        soporte = self._soporte_borrador()
        r = self.client.post(f'/pagos/soporte/{soporte.pk}/eliminar/')
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.pago.soportes.count(), 1)  # sigue existiendo

    def test_descargar_soporte_rechaza_lote_borrador(self):
        soporte = self._soporte_borrador()
        r = self.client.get(f'/pagos/soporte/{soporte.pk}/descargar/')
        self.assertEqual(r.status_code, 404)


class FinPagosProyeccionTest(TestCase):
    """Proyección de pagos: cálculo puro desde clases programadas (horas × valor hora),
    SOLO LECTURA (nunca materializa lotes ni filas), con filtros y export a Excel."""

    def setUp(self):
        self.client = Client(HTTP_HOST='financiera.testserver')
        grupo_fin = Group.objects.get(name=GRUPO_STAFF_FINANCIERA)
        self.finan = User.objects.create_user(username='finan', password='pass')
        self.finan.groups.add(grupo_fin)
        self.client.login(username='finan', password='pass')

        self.hoy = date.today()
        anio = self.hoy.year

        colegio_a = Colegio.objects.create(
            nombre='Colegio Central', codigo='CC1',
            departamento='Santander', ciudad='Bucaramanga')
        colegio_b = Colegio.objects.create(
            nombre='Colegio Norte', codigo='CN1',
            departamento='Santander', ciudad='Bucaramanga')
        self.ca_a = ColegioAnio.objects.create(colegio=colegio_a, anio=anio, valor_hora=40000)
        self.ca_b = ColegioAnio.objects.create(colegio=colegio_b, anio=anio, valor_hora=50000)

        self.ana  = Profesor.objects.create(nombre='Ana', apellido='Pérez')
        self.luis = Profesor.objects.create(nombre='Luis', apellido='Gómez')

        grado = Grado.objects.create(nombre='11-1')
        # Bloque de 2 h en A y de 1 h en B → proyecciones 80.000 y 50.000.
        self.bloque_a = Bloque.objects.create(
            colegio=self.ca_a, grado=grado, hora_inicio=time(8, 0), hora_fin=time(10, 0))
        self.bloque_b = Bloque.objects.create(
            colegio=self.ca_b, grado=grado, hora_inicio=time(8, 0), hora_fin=time(9, 0))

        Clase.objects.create(colegio=self.ca_a, bloque=self.bloque_a,
                             profesor=self.ana, fecha=self.hoy + timedelta(days=5))
        Clase.objects.create(colegio=self.ca_b, bloque=self.bloque_b,
                             profesor=self.luis, fecha=self.hoy + timedelta(days=6))
        # Excluidas del cálculo (mismas reglas que los pagos reales):
        Clase.objects.create(colegio=self.ca_a, bloque=self.bloque_a, profesor=self.ana,
                             fecha=self.hoy + timedelta(days=7), cancelada=True)
        Clase.objects.create(colegio=self.ca_a, bloque=self.bloque_a, profesor=self.ana,
                             fecha=self.hoy + timedelta(days=8), es_evento=True,
                             titulo_evento='Izada de bandera')
        Clase.objects.create(colegio=self.ca_a, bloque=self.bloque_a, profesor=None,
                             fecha=self.hoy + timedelta(days=9))
        # Pasada: fuera con el default `desde=hoy`, visible si el usuario amplía el rango.
        Clase.objects.create(colegio=self.ca_a, bloque=self.bloque_a,
                             profesor=self.ana, fecha=self.hoy - timedelta(days=5))

    # ── Cálculo y exclusiones ─────────────────────────────────
    def test_calcula_horas_por_valor_hora(self):
        r = self.client.get('/pagos/proyeccion/')
        self.assertEqual(r.status_code, 200)
        filas = r.context['filas']
        self.assertEqual(len(filas), 2)  # canceladas/eventos/sin profesor/pasadas fuera
        por_prof = {f['profesor_id']: f for f in filas}
        self.assertEqual(por_prof[self.ana.id]['valor_total'], 80000)   # 2 h × 40.000
        self.assertEqual(por_prof[self.luis.id]['valor_total'], 50000)  # 1 h × 50.000
        self.assertEqual(r.context['total_valor'], 130000)
        self.assertEqual(r.context['total_horas'], 3)

    def test_pasadas_entran_solo_si_se_amplia_el_rango(self):
        desde = (self.hoy - timedelta(days=10)).isoformat()
        r = self.client.get(f'/pagos/proyeccion/?desde={desde}')
        fechas = {f['fecha'] for f in r.context['filas']}
        self.assertIn(self.hoy - timedelta(days=5), fechas)

    # ── Filtros ───────────────────────────────────────────────
    def test_filtro_por_colegio(self):
        r = self.client.get(f'/pagos/proyeccion/?colegio_id={self.ca_b.id}')
        filas = r.context['filas']
        self.assertEqual(len(filas), 1)
        self.assertEqual(filas[0]['colegio_id'], self.ca_b.id)
        self.assertEqual(r.context['total_valor'], 50000)

    def test_filtro_por_profesor(self):
        r = self.client.get(f'/pagos/proyeccion/?profesor_id={self.ana.id}')
        filas = r.context['filas']
        self.assertEqual(len(filas), 1)
        self.assertEqual(filas[0]['profesor_id'], self.ana.id)

    def test_filtro_hasta_acota(self):
        hasta = (self.hoy + timedelta(days=5)).isoformat()
        r = self.client.get(f'/pagos/proyeccion/?hasta={hasta}')
        filas = r.context['filas']
        self.assertEqual(len(filas), 1)
        self.assertEqual(filas[0]['profesor_id'], self.ana.id)

    # ── Solo lectura ──────────────────────────────────────────
    def test_no_escribe_lotes_ni_pagos(self):
        self.client.get('/pagos/proyeccion/')
        self.assertEqual(LotePagos.objects.count(), 0)
        self.assertEqual(PagoRealizado.objects.count(), 0)

    # ── Gates ─────────────────────────────────────────────────
    def test_rechaza_usuario_solo_programacion(self):
        grupo_prog = Group.objects.get(name=GRUPO_STAFF_PROGRAMACION)
        u = User.objects.create_user(username='prog', password='pass')
        u.groups.add(grupo_prog)
        self.client.logout()
        self.client.login(username='prog', password='pass')
        r = self.client.get('/pagos/proyeccion/')
        self.assertNotEqual(r.status_code, 200)

    def test_rechaza_anonimo(self):
        self.client.logout()
        r = self.client.get('/pagos/proyeccion/')
        self.assertNotEqual(r.status_code, 200)

    # ── Export ────────────────────────────────────────────────
    def test_exportar_devuelve_xlsx_con_total(self):
        r = self.client.post('/pagos/proyeccion/exportar/', {'desde': self.hoy.isoformat()})
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r['Content-Type'])
        self.assertIn('attachment', r['Content-Disposition'])
        ws = load_workbook(io.BytesIO(r.content)).active
        # título + cabecera + 2 filas + TOTAL
        self.assertEqual(ws.max_row, 5)
        self.assertEqual(ws.cell(5, 7).value, 130000)
        self.assertEqual(ws.cell(5, 5).value, 3)

    def test_exportar_respeta_filtros(self):
        r = self.client.post('/pagos/proyeccion/exportar/', {
            'desde': self.hoy.isoformat(), 'profesor_id': self.ana.id})
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws.max_row, 4)  # título + cabecera + 1 fila + TOTAL
        self.assertEqual(ws.cell(3, 7).value, 80000)

    def test_exportar_exige_post(self):
        r = self.client.get('/pagos/proyeccion/exportar/')
        self.assertEqual(r.status_code, 405)

    # ── Menú ──────────────────────────────────────────────────
    def test_menu_enlaza_proyeccion(self):
        r = self.client.get('/pagos/')
        self.assertContains(r, '/pagos/proyeccion/')
