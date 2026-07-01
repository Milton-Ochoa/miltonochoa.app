"""Vistas del área **financiera** para los pagos de clases a profesores.

Flujo por **lotes**: programación materializa y revisa las filas `PagoRealizado`
(grano `(profesor, colegio, día)`) en un `LotePagos` BORRADOR y lo **envía**;
financiera **solo ve las filas de lotes ENVIADO** (guard `_lote_enviado`) y las
**gestiona**: marca/desmarca el pago (`fecha_pago`/`marcado_por`) y sube/elimina
los soportes (comprobantes). Programación las consulta en solo lectura.

El cálculo y los helpers compartidos (`construir_contexto_pagos`,
`filas_pagos_por_tab`, `_generar_excel_pagos`) viven en `programacion.pagos.views`
y se importan con `modo='financiera'` (BD única, sin duplicar lógica).

Gate: superusuario o grupo `area:financiera` (`core.areas.es_personal_financiera`).
"""
import io
from datetime import date

from django.contrib import messages
from django.contrib.auth.decorators import user_passes_test
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.areas import es_personal_financiera
from programacion.configuracion.models import ColegioAnio, Profesor
from programacion.pagos.models import LotePagos, PagoRealizado, SoportePagoProfesor
from programacion.pagos.views import (
    construir_contexto_pagos, filas_pagos_por_tab,
    _build_filas_pagos, _generar_excel_pagos, _label_rango, _parse_fecha,
)
from programacion.viaticos.soportes import validar_soporte
from programacion.viaticos.views import _responder_soporte

solo_financiera = user_passes_test(es_personal_financiera, login_url='login')


def _lote_enviado(pago):
    """Financiera solo ve lo enviado: cualquier acceso a una fila cuyo lote no esté
    ENVIADO (BORRADOR, o histórico sin lote) se rechaza. Mismo guard que fin_pagos_marcar."""
    return pago.lote is not None and pago.lote.estado == LotePagos.Estado.ENVIADO


@solo_financiera
def fin_pagos_lista(request):
    """Tabla semanal de pagos (tabs pendiente/realizado). Mismo cálculo que
    programación (reutiliza `construir_contexto_pagos`), dibujado en el chrome
    financiera con las acciones de marcar/soporte."""
    return render(request, 'financiera/pagos.html',
                  construir_contexto_pagos(request.GET, modo='financiera'))


@solo_financiera
@require_POST
def fin_pagos_marcar(request):
    """Marca/desmarca una fila enviada como pago realizado, fijando/limpiando `fecha_pago`.

    La fila ya existe (la materializó y envió programación): financiera solo registra el
    pago, no crea filas. Solo se gestiona si pertenece a un lote ENVIADO. `desmarcar`
    limpia `fecha_pago`/`marcado_por` y borra los soportes (archivos en storage + filas),
    pero **conserva la fila** (sigue siendo parte del lote enviado)."""
    try:
        pago_id = int(request.POST.get('pago_id', ''))
    except (ValueError, TypeError):
        return JsonResponse({'ok': False, 'error': 'Datos inválidos'}, status=400)

    pago = get_object_or_404(
        PagoRealizado.objects.select_related('lote').prefetch_related('soportes'),
        pk=pago_id)
    if not _lote_enviado(pago):
        return JsonResponse(
            {'ok': False, 'error': 'El pago no está disponible para financiera.'}, status=400)

    if request.POST.get('accion', 'marcar') == 'desmarcar':
        for soporte in pago.soportes.all():
            soporte.archivo.delete(save=False)
        pago.soportes.all().delete()
        pago.fecha_pago = None
        pago.marcado_por = None
        pago.save(update_fields=['fecha_pago', 'marcado_por'])
        return JsonResponse({'ok': True, 'accion': 'desmarcado', 'id': pago.id})

    pago.fecha_pago = timezone.now()
    pago.marcado_por = request.user
    pago.save(update_fields=['fecha_pago', 'marcado_por'])
    return JsonResponse({'ok': True, 'accion': 'marcado', 'id': pago.id})


@solo_financiera
@require_POST
def fin_pagos_exportar(request):
    """Descarga el `.xlsx` del tab activo (reusa el generador de programación).

    Rango opcional: si no llega filtro de fecha, se exporta **todo** el backlog del tab
    (no la semana actual), igual que en programación."""
    fi = _parse_fecha(request.POST.get('fecha_inicio', ''))
    ff = _parse_fecha(request.POST.get('fecha_fin', ''))

    tab = request.POST.get('tab', 'pendiente')
    filas = filas_pagos_por_tab(fi, ff, tab, modo='financiera')
    excel_bytes = _generar_excel_pagos(filas, _label_rango(fi, ff))

    sufijo = 'Realizados' if tab == 'realizado' else 'Pendientes'
    label = f"{fi.strftime('%Y%m%d')}_{ff.strftime('%Y%m%d')}" if fi and ff else 'todos'
    response = HttpResponse(
        excel_bytes,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="Pagos_{sufijo}_{label}.xlsx"'
    return response


# ══════════════════════════════════════════════════════════════
# DETALLE + SOPORTES (comprobante por pago)
# ══════════════════════════════════════════════════════════════

@solo_financiera
def fin_pagos_detalle(request, pago_id):
    """Detalle de un pago realizado: datos del docente/colegio y gestión de soportes
    (subir/ver/descargar/eliminar). Espejo del detalle de viáticos en `PAGADA`."""
    pago = get_object_or_404(
        PagoRealizado.objects
        .select_related('profesor', 'colegio__colegio', 'marcado_por', 'lote', 'lote__enviado_por')
        .prefetch_related('soportes', 'soportes__subido_por'),
        pk=pago_id,
    )
    if not _lote_enviado(pago):
        messages.error(request, 'El pago no está disponible para financiera.')
        return redirect('fin_pagos_lista')
    return render(request, 'financiera/pagos_detalle.html', {'pago': pago})


@solo_financiera
@require_POST
def fin_pagos_subir_soporte(request, pago_id):
    """Adjunta un comprobante a un pago realizado (mismo validador que viáticos)."""
    pago = get_object_or_404(PagoRealizado.objects.select_related('lote'), pk=pago_id)
    if not _lote_enviado(pago):
        messages.error(request, 'El pago no está disponible para financiera.')
        return redirect('fin_pagos_lista')
    archivo = request.FILES.get('archivo')
    if not archivo:
        messages.error(request, 'Selecciona un archivo para subir.')
        return redirect('fin_pagos_detalle', pago_id=pago.pk)

    error = validar_soporte(archivo)
    if error:
        messages.error(request, error)
        return redirect('fin_pagos_detalle', pago_id=pago.pk)

    SoportePagoProfesor.objects.create(
        pago=pago, archivo=archivo, nombre_original=archivo.name, subido_por=request.user)
    messages.success(request, 'Soporte de pago adjuntado.')
    return redirect('fin_pagos_detalle', pago_id=pago.pk)


@solo_financiera
@require_POST
def fin_pagos_eliminar_soporte(request, soporte_id):
    """Elimina un soporte (y su archivo en storage) subido por error."""
    soporte = get_object_or_404(
        SoportePagoProfesor.objects.select_related('pago__lote'), pk=soporte_id)
    if not _lote_enviado(soporte.pago):
        messages.error(request, 'El pago no está disponible para financiera.')
        return redirect('fin_pagos_lista')
    pago_id = soporte.pago_id
    soporte.archivo.delete(save=False)  # primero el storage, luego la fila
    soporte.delete()
    messages.success(request, 'Soporte eliminado.')
    return redirect('fin_pagos_detalle', pago_id=pago_id)


@solo_financiera
def fin_pago_soporte_descargar(request, soporte_id):
    """Ver (`?inline=1`) o descargar un soporte. Mismo proxy server-side que viáticos."""
    soporte = get_object_or_404(
        SoportePagoProfesor.objects.select_related('pago__lote'), pk=soporte_id)
    if not _lote_enviado(soporte.pago):
        raise Http404
    return _responder_soporte(soporte, inline=request.GET.get('inline') == '1')


# ══════════════════════════════════════════════════════════════
# PROYECCIÓN — SOLO LECTURA
# ══════════════════════════════════════════════════════════════
# Costo estimado de las clases PROGRAMADAS (horas × valor hora del colegio), para que
# financiera anticipe el gasto. Es un cálculo puro desde clases (`_build_filas_pagos`):
# NUNCA materializa lotes ni filas de pago, no fija fechas de pago ni escribe en BD.

def _filtros_proyeccion(data):
    """Filtros de la proyección (GET de la página o POST del export).

    `desde` default hoy (lo programado de aquí en adelante); `hasta` y los ids de
    colegio/profesor son opcionales. Ids inválidos se ignoran (= sin filtro)."""
    desde = _parse_fecha(data.get('desde', '')) or date.today()
    hasta = _parse_fecha(data.get('hasta', ''))

    def _id(clave):
        try:
            return int(data.get(clave, ''))
        except (ValueError, TypeError):
            return None

    return desde, hasta, _id('colegio_id'), _id('profesor_id')


def _filas_proyeccion(desde, hasta, colegio_id, profesor_id):
    """Filas proyectadas: el mismo cálculo (y exclusiones: canceladas, eventos, sin
    profesor) que alimenta los pagos reales, sin persistir nada.

    `_build_filas_pagos` exige rango cerrado; sin `hasta` se proyecta TODO lo
    programado (cota `date.max`). Colegio/profesor se filtran en Python para no
    cambiar la firma del helper compartido (el volumen de clases futuras es chico)."""
    filas = _build_filas_pagos(desde, hasta or date.max)
    if colegio_id:
        filas = [f for f in filas if f['colegio_id'] == colegio_id]
    if profesor_id:
        filas = [f for f in filas if f['profesor_id'] == profesor_id]
    return filas


@solo_financiera
def fin_pagos_proyeccion(request):
    """Listado de clases programadas con su costo estimado, filtros por colegio/
    profesor/rango y totales. Solo lectura: no prepara lotes ni marca pagos."""
    desde, hasta, colegio_id, profesor_id = _filtros_proyeccion(request.GET)
    filas = _filas_proyeccion(desde, hasta, colegio_id, profesor_id)
    return render(request, 'financiera/pagos_proyeccion.html', {
        'filas':        filas,
        'desde':        desde.isoformat(),
        'hasta':        hasta.isoformat() if hasta else '',
        'colegio_id':   colegio_id,
        'profesor_id':  profesor_id,
        'total_horas':  sum(f['horas'] for f in filas),
        'total_valor':  sum(f['valor_total'] for f in filas),
        'rango_label':  _label_rango(desde, hasta),
        # Catálogos de los selects. ColegioAnio (no Colegio): la tarifa y las clases
        # cuelgan del periodo; `periodo_label` desambigua periodos del mismo colegio.
        'colegios':     ColegioAnio.objects.select_related('colegio')
                                   .order_by('colegio__nombre', '-anio'),
        'profesores':   Profesor.objects.order_by('nombre', 'apellido'),
    })


def _generar_excel_proyeccion(filas, rango_label):
    """Excel de la proyección. Self-contained a propósito (no reutiliza
    `_generar_excel_pagos`: sus columnas bancarias y el desglose sobran aquí)."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Proyección'

    COLS = ['FECHA', 'DOCENTE', 'COLEGIO', 'CODIGO', 'HORAS',
            'VALOR/HORA', 'VALOR PROYECTADO']
    NUM_COLS = len(COLS)

    thin = Side(style='thin', color='000000')
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _celda(row, col, valor='', bold=False, fill=None, color='000000',
               h='center', v='center', fmt=None):
        cell = ws.cell(row, col, valor)
        cell.font = Font(name='Arial', size=10, bold=bold, color=color)
        cell.alignment = Alignment(horizontal=h, vertical=v, wrap_text=True)
        if fill:
            cell.fill = PatternFill('solid', fgColor=fill)
        cell.border = borde
        if fmt:
            cell.number_format = fmt
        return cell

    # Fila 1: título con el rango. Deja claro en el archivo que es una estimación.
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NUM_COLS)
    _celda(1, 1, f'Proyección de pagos (estimado) — {rango_label}',
           bold=True, fill='FFD9E1F2', color='FF1F3864')
    for c in range(1, NUM_COLS + 1):   # borde/relleno de todo el rango combinado
        ws.cell(1, c).border = borde
        ws.cell(1, c).fill = PatternFill('solid', fgColor='FFD9E1F2')
    ws.row_dimensions[1].height = 20

    # Fila 2: cabeceras
    for ci, nombre in enumerate(COLS, start=1):
        _celda(2, ci, nombre, bold=True, fill='FFB8CCE4', color='FF1F3864')
    ws.row_dimensions[2].height = 22

    for i, f in enumerate(filas, start=3):
        fill_row = 'FFFFFFFF' if i % 2 == 1 else 'FFF2F6FC'
        _celda(i, 1, f"{f['fecha'].day:02d}/{f['fecha'].month:02d}/{f['fecha'].year}",
               fill=fill_row)
        _celda(i, 2, f['docente'], fill=fill_row, h='left')
        _celda(i, 3, f['colegio'], fill=fill_row, h='left')
        _celda(i, 4, f['codigo'],  fill=fill_row)
        _celda(i, 5, f['horas'],   fill=fill_row, h='right', fmt='0.##')
        _celda(i, 6, f['valor_hora'],  fill=fill_row, h='right', fmt='"$"#,##0')
        _celda(i, 7, f['valor_total'], fill=fill_row, h='right', fmt='"$"#,##0')
        ws.row_dimensions[i].height = 18

    # Fila TOTAL (horas y valor como números, para que Excel pueda operar)
    total_row = 3 + len(filas)
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    _celda(total_row, 1, 'TOTAL', bold=True, fill='FFD6DCE4', h='right')
    for c in range(1, 5):
        ws.cell(total_row, c).border = borde
        ws.cell(total_row, c).fill = PatternFill('solid', fgColor='FFD6DCE4')
    _celda(total_row, 5, sum(f['horas'] for f in filas),
           bold=True, fill='FFD6DCE4', h='right', fmt='0.##')
    _celda(total_row, 6, '', fill='FFD6DCE4')
    _celda(total_row, 7, sum(f['valor_total'] for f in filas),
           bold=True, fill='FFD6DCE4', h='right', fmt='"$"#,##0')
    ws.row_dimensions[total_row].height = 22

    anchos = [12, 26, 30, 10, 9, 13, 17]
    for ci, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = ancho

    ws.freeze_panes = 'A3'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@solo_financiera
@require_POST
def fin_pagos_proyeccion_exportar(request):
    """Descarga el `.xlsx` de la proyección con los filtros vigentes de la página."""
    desde, hasta, colegio_id, profesor_id = _filtros_proyeccion(request.POST)
    filas = _filas_proyeccion(desde, hasta, colegio_id, profesor_id)
    excel_bytes = _generar_excel_proyeccion(filas, _label_rango(desde, hasta))

    label = f"{desde.strftime('%Y%m%d')}_{hasta.strftime('%Y%m%d') if hasta else 'adelante'}"
    response = HttpResponse(
        excel_bytes,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="Proyeccion_{label}.xlsx"'
    return response
