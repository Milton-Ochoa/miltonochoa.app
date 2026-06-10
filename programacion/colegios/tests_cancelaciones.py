"""
Tests — Cancelaciones de clase (Fase 5).

Cubre: cancelación por COLEGIO (comportamiento histórico + registro vigente),
cancelación por PROFESOR (clase viva sin profesor + registro histórico),
reporte /reportes/cancelaciones/ (página, Excel, gates) y el backfill 0019.
"""
import io
from datetime import date, time
from importlib import import_module

from django.apps import apps as django_apps
from django.contrib.auth.models import User
from django.test import TestCase, Client

from programacion.configuracion.models import Colegio, ColegioAnio, Profesor, Materia
from programacion.colegios.models import (
    Bloque, Clase, Grado, HistorialCambio, CancelacionClase,
)
from programacion.pagos.views import _build_filas_pagos


class CancelacionClaseTestBase(TestCase):
    """Setup común: colegio 2026, bloque, dos profesores y un superusuario logueado."""

    def setUp(self):
        self.client = Client(HTTP_HOST='programacion.testserver')
        self.admin = User.objects.create_superuser('admin_canc', password='pass')
        self.client.login(username='admin_canc', password='pass')
        col = Colegio.objects.create(
            nombre='Col Cancelaciones', departamento='Santander', ciudad='BGA'
        )
        self.colegio = ColegioAnio.objects.create(colegio=col, anio=2026, activo=True)
        self.grado = Grado.objects.create(nombre='10-1')
        self.bloque = Bloque.objects.create(
            colegio=self.colegio, grado=self.grado,
            hora_inicio=time(8, 0), hora_fin=time(10, 0),
        )
        self.materia = Materia.objects.create(nombre='Química', color='#27ae60')
        self.profesor = Profesor.objects.create(nombre='Ana', apellido='Gómez')
        self.profesor2 = Profesor.objects.create(nombre='Beto', apellido='Lara')

    def _post_clase(self, **data):
        payload = {
            'guardar_clase': '1',
            'bloque_id':     str(self.bloque.id),
            'fecha_clase':   '2026-05-15',
            'materia':       'Química',
            'profesor':      str(self.profesor.id),
            'unidad':        '1',
            'eliminar_clase': '0',
            'material_especial': '0',
            'tipo_especial': '',
            'libro_especial_id': '',
            'enlace_personalizado': '',
        }
        payload.update(data)
        return self.client.post(
            f'/colegios/ajax/guardar-clase/{self.colegio.id}/', data=payload,
        )


class CancelacionColegioTest(CancelacionClaseTestBase):

    def test_cancelar_colegio_crea_registro_y_clase_cancelada(self):
        r = self._post_clase(cancelada='on', cancelada_por='COLEGIO',
                             comentarios='Paro de transporte')
        self.assertEqual(r.status_code, 200)
        clase = Clase.objects.get(bloque=self.bloque, fecha=date(2026, 5, 15))
        self.assertTrue(clase.cancelada)
        self.assertEqual(clase.profesor_id, self.profesor.id)  # conserva al profesor
        reg = CancelacionClase.objects.get()
        self.assertEqual(reg.tipo, CancelacionClase.Tipo.COLEGIO)
        self.assertEqual(reg.clase_id, clase.id)
        self.assertEqual(reg.profesor_id, self.profesor.id)
        self.assertEqual(reg.profesor_nombre, 'Ana Gómez')
        self.assertEqual(reg.colegio_nombre, 'Col Cancelaciones')
        self.assertEqual(reg.fecha_clase, date(2026, 5, 15))
        self.assertEqual(reg.motivo, 'Paro de transporte')
        self.assertEqual(reg.registrado_por, self.admin)

    def test_post_sin_cancelada_por_default_colegio(self):
        # Compatibilidad: un POST viejo sin el radio se comporta como hoy (colegio)
        self._post_clase(cancelada='on', comentarios='x')
        clase = Clase.objects.get(bloque=self.bloque, fecha=date(2026, 5, 15))
        self.assertTrue(clase.cancelada)
        self.assertEqual(CancelacionClase.objects.get().tipo, 'COLEGIO')

    def test_descancelar_elimina_registro_colegio(self):
        self._post_clase(cancelada='on', comentarios='Motivo X')
        self.assertEqual(CancelacionClase.objects.count(), 1)
        self._post_clase()  # checkbox off
        clase = Clase.objects.get(bloque=self.bloque, fecha=date(2026, 5, 15))
        self.assertFalse(clase.cancelada)
        self.assertEqual(CancelacionClase.objects.count(), 0)

    def test_editar_cancelada_no_duplica_registro_y_refresca_motivo(self):
        self._post_clase(cancelada='on', comentarios='Motivo viejo')
        self._post_clase(cancelada='on', comentarios='Motivo nuevo')
        reg = CancelacionClase.objects.get()  # sigue habiendo UNO
        self.assertEqual(reg.motivo, 'Motivo nuevo')

    def test_historial_incluye_detalle_de_cancelacion(self):
        self._post_clase(cancelada='on', comentarios='Lluvia')
        ultimo = HistorialCambio.objects.filter(objeto_tipo='Clase').first()
        self.assertIn('Cancelación por colegio', ultimo.detalle)
        self.assertIn('Lluvia', ultimo.detalle)


class CancelacionProfesorTest(CancelacionClaseTestBase):

    def _crear_clase(self):
        return Clase.objects.create(
            colegio=self.colegio, bloque=self.bloque, fecha=date(2026, 5, 15),
            materia=self.materia, profesor=self.profesor, unidad='1',
        )

    def test_cancelar_profesor_deja_clase_viva_sin_profesor(self):
        self._crear_clase()
        r = self._post_clase(cancelada='on', cancelada_por='PROFESOR',
                             comentarios='Se enfermó')
        self.assertEqual(r.status_code, 200)
        clase = Clase.objects.get(bloque=self.bloque, fecha=date(2026, 5, 15))
        self.assertFalse(clase.cancelada)        # la clase NO muere
        self.assertIsNone(clase.profesor_id)     # pendiente de reasignar
        reg = CancelacionClase.objects.get()
        self.assertEqual(reg.tipo, CancelacionClase.Tipo.PROFESOR)
        self.assertEqual(reg.profesor_id, self.profesor.id)   # snapshot del original
        self.assertEqual(reg.profesor_nombre, 'Ana Gómez')
        self.assertEqual(reg.motivo, 'Se enfermó')

    def test_reasignar_no_borra_el_registro(self):
        self._crear_clase()
        self._post_clase(cancelada='on', cancelada_por='PROFESOR', comentarios='x')
        # Reasignación = edición normal con el nuevo profesor
        self._post_clase(profesor=str(self.profesor2.id))
        clase = Clase.objects.get(bloque=self.bloque, fecha=date(2026, 5, 15))
        self.assertEqual(clase.profesor_id, self.profesor2.id)
        self.assertEqual(CancelacionClase.objects.count(), 1)  # histórico intacto

    def test_segunda_cancelacion_crea_segundo_registro(self):
        # A cancela → se reasigna a B → B cancela = dos registros con su snapshot
        self._crear_clase()
        self._post_clase(cancelada='on', cancelada_por='PROFESOR', comentarios='A')
        self._post_clase(profesor=str(self.profesor2.id))
        self._post_clase(profesor=str(self.profesor2.id),
                         cancelada='on', cancelada_por='PROFESOR', comentarios='B')
        regs = list(CancelacionClase.objects.order_by('registrado_en', 'id'))
        self.assertEqual(len(regs), 2)
        self.assertEqual(regs[0].profesor_id, self.profesor.id)
        self.assertEqual(regs[1].profesor_id, self.profesor2.id)

    def test_sin_fila_de_pago_mientras_este_sin_profesor(self):
        self._crear_clase()
        filas = _build_filas_pagos(date(2026, 5, 11), date(2026, 5, 17))
        self.assertEqual(len(filas), 1)  # antes de cancelar sí hay fila
        self._post_clase(cancelada='on', cancelada_por='PROFESOR', comentarios='x')
        filas = _build_filas_pagos(date(2026, 5, 11), date(2026, 5, 17))
        self.assertEqual(filas, [])      # sin profesor → sin fila de pago

    def test_historial_incluye_nombre_y_motivo(self):
        self._crear_clase()
        self._post_clase(cancelada='on', cancelada_por='PROFESOR', comentarios='Cita médica')
        ultimo = HistorialCambio.objects.filter(objeto_tipo='Clase').first()
        self.assertIn('Cancelación por profesor', ultimo.detalle)
        self.assertIn('Ana Gómez', ultimo.detalle)
        self.assertIn('Cita médica', ultimo.detalle)

    def test_eliminar_clase_conserva_registro_con_snapshot(self):
        clase = self._crear_clase()
        self._post_clase(cancelada='on', cancelada_por='PROFESOR', comentarios='x')
        Clase.objects.filter(id=clase.id).delete()
        reg = CancelacionClase.objects.get()
        self.assertIsNone(reg.clase_id)          # FK SET_NULL
        self.assertEqual(reg.profesor_nombre, 'Ana Gómez')
        self.assertEqual(reg.colegio_nombre, 'Col Cancelaciones')


class ReporteCancelacionesViewTest(CancelacionClaseTestBase):

    def _crear_registros(self):
        CancelacionClase.objects.create(
            tipo='COLEGIO', colegio=self.colegio.colegio,
            colegio_nombre='Col Cancelaciones', profesor=self.profesor,
            profesor_nombre='Ana Gómez', fecha_clase=date(2026, 3, 10),
            motivo='Semana santa', registrado_por=self.admin,
        )
        CancelacionClase.objects.create(
            tipo='PROFESOR', colegio=self.colegio.colegio,
            colegio_nombre='Col Cancelaciones', profesor=self.profesor2,
            profesor_nombre='Beto Lara', fecha_clase=date(2026, 4, 20),
            motivo='Incapacidad', registrado_por=self.admin,
        )

    def test_pagina_renderiza_registros(self):
        self._crear_registros()
        r = self.client.get('/reportes/cancelaciones/')
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Ana Gómez')
        self.assertContains(r, 'Beto Lara')
        self.assertContains(r, 'Semana santa')
        self.assertContains(r, 'Incapacidad')

    def test_excel_descarga_con_content_type(self):
        self._crear_registros()
        r = self.client.post('/reportes/cancelaciones/excel/', data={})
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r['Content-Type'])
        self.assertIn('Cancelaciones_', r['Content-Disposition'])

    def test_excel_filtra_por_tipo_y_fecha(self):
        self._crear_registros()
        from openpyxl import load_workbook
        r = self.client.post('/reportes/cancelaciones/excel/', data={
            'tipos': ['PROFESOR'],
            'fecha_desde': '2026-04-01',
            'fecha_hasta': '2026-04-30',
        })
        ws = load_workbook(io.BytesIO(r.content)).active
        filas = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(filas), 1)
        self.assertEqual(filas[0][2], 'Beto Lara')

    def test_gestor_colegio_no_accede(self):
        # El prefijo /reportes/ no está en _PERMITIDAS_COLEGIO → el middleware
        # lo redirige a su dashboard.
        from usuarios.models import UsuarioColegio
        gestor = User.objects.create_user('gestor_canc', password='pass')
        UsuarioColegio.objects.create(user=gestor, colegio=self.colegio.colegio)
        self.client.login(username='gestor_canc', password='pass')
        r = self.client.get('/reportes/cancelaciones/')
        self.assertEqual(r.status_code, 302)

    def test_profesor_no_accede(self):
        from usuarios.models import UsuarioProfesor
        prof_user = User.objects.create_user('prof_canc', password='pass')
        UsuarioProfesor.objects.create(user=prof_user, profesor=self.profesor)
        self.client.login(username='prof_canc', password='pass')
        r = self.client.get('/reportes/cancelaciones/')
        self.assertEqual(r.status_code, 302)

    def test_anonimo_redirige_a_login(self):
        self.client.logout()
        r = self.client.get('/reportes/cancelaciones/')
        self.assertEqual(r.status_code, 302)


class BackfillCancelacionesTest(TestCase):
    """La data migration 0019 ya corrió al construir la BD de tests; aquí se
    valida su lógica re-ejecutándola sobre datos nuevos."""

    def test_backfill_crea_registros_colegio(self):
        col = Colegio.objects.create(nombre='Col Backfill', departamento='S', ciudad='B')
        ca = ColegioAnio.objects.create(colegio=col, anio=2026, activo=True)
        grado = Grado.objects.create(nombre='9-1')
        bloque = Bloque.objects.create(
            colegio=ca, grado=grado, hora_inicio=time(7, 0), hora_fin=time(8, 0),
        )
        prof = Profesor.objects.create(nombre='Caro', apellido='Díaz')
        Clase.objects.create(
            colegio=ca, bloque=bloque, fecha=date(2026, 2, 2),
            profesor=prof, cancelada=True, comentarios='Histórica',
        )
        Clase.objects.create(
            colegio=ca, bloque=bloque, fecha=date(2026, 2, 3),
            profesor=prof, cancelada=False,
        )

        backfill = import_module(
            'programacion.colegios.migrations.0019_backfill_cancelaciones'
        ).backfill_cancelaciones
        backfill(django_apps, None)

        reg = CancelacionClase.objects.get()  # solo la cancelada
        self.assertEqual(reg.tipo, 'COLEGIO')
        self.assertEqual(reg.profesor_nombre, 'Caro Díaz')
        self.assertEqual(reg.colegio_nombre, 'Col Backfill')
        self.assertEqual(reg.fecha_clase, date(2026, 2, 2))
        self.assertEqual(reg.motivo, 'Histórica')
        self.assertIsNone(reg.registrado_por)
