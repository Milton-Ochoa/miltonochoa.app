import io

from django.test import TestCase, Client
from django.contrib.auth.models import User, Group
from django.core.files.uploadedfile import SimpleUploadedFile

from openpyxl import Workbook, load_workbook

from programacion.colegios.models import Grado
from programacion.monitores.models import (AsignacionMonitor, ColegioSimulacro,
                                           Monitor, Simulacro)


def _xlsx_bytes(filas, *, encabezados=('nombre', 'codigo', 'ciudad', 'departamento')):
    """Construye un .xlsx en memoria para los tests de carga masiva."""
    wb = Workbook()
    ws = wb.active
    if encabezados is not None:
        ws.append(list(encabezados))
    for fila in filas:
        ws.append(list(fila))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ConfiguracionMonitoresViewTest(TestCase):
    """CRUD del catálogo de monitores desde Configuración (programación)."""

    URL = '/monitores/configuracion/'

    def setUp(self):
        self.client = Client(HTTP_HOST='programacion.testserver')
        User.objects.create_superuser(username='admin_mon', password='pass123')
        self.client.login(username='admin_mon', password='pass123')

    def test_listado_accesible(self):
        Monitor.objects.create(nombre='Ana', apellido='Pérez')
        resp = self.client.get(self.URL)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Ana')

    def test_crear_monitor(self):
        self.client.post(self.URL, {
            'accion': 'add',
            'nombre': 'Carlos', 'apellido': 'Gómez',
            'documento': '12345', 'celular': '3001112233',
            'departamento': 'Santander', 'ciudad': 'Bucaramanga',
            'banco': Monitor.Banco.NEQUI, 'tipo_cuenta': Monitor.TipoCuenta.AHORROS,
            'cuenta_bancaria': '999',
        })
        m = Monitor.objects.get(documento='12345')
        self.assertEqual(m.nombre, 'Carlos')
        self.assertEqual(m.ciudad, 'Bucaramanga')
        self.assertTrue(m.activo)

    def test_crear_documento_duplicado_no_rompe(self):
        Monitor.objects.create(nombre='Uno', documento='555')
        self.client.post(self.URL, {
            'accion': 'add', 'nombre': 'Dos', 'documento': '555',
        })
        # El unique impide el segundo; solo queda uno.
        self.assertEqual(Monitor.objects.filter(documento='555').count(), 1)

    def test_editar_monitor(self):
        m = Monitor.objects.create(nombre='Luis', celular='300')
        self.client.post(self.URL, {
            'accion': 'edit', 'monitor_id': m.id,
            'nombre': 'Luis', 'celular': '311',
        })
        m.refresh_from_db()
        self.assertEqual(m.celular, '311')

    def test_toggle_activo(self):
        m = Monitor.objects.create(nombre='Pedro')
        self.assertTrue(m.activo)
        self.client.post(self.URL, {'accion': 'toggle_activo', 'monitor_id': m.id})
        m.refresh_from_db()
        self.assertFalse(m.activo)

    def test_eliminar_monitor(self):
        m = Monitor.objects.create(nombre='Borrar')
        self.client.post(self.URL, {'accion': 'del', 'monitor_id': m.id})
        self.assertFalse(Monitor.objects.filter(id=m.id).exists())

    def test_gate_anonimo_redirige(self):
        self.client.logout()
        resp = self.client.get(self.URL)
        self.assertNotEqual(resp.status_code, 200)


class MonitorModelTest(TestCase):
    def test_nombre_corto(self):
        m = Monitor(nombre='Juan Carlos', apellido='Pérez Gómez')
        self.assertEqual(m.nombre_corto, 'Juan Pérez')

    def test_nombre_corto_sin_apellido(self):
        m = Monitor(nombre='Ana')
        self.assertEqual(m.nombre_corto, 'Ana')


class StaffAreaAccesoTest(TestCase):
    """El staff del área (grupo area:programacion, no superusuario) accede."""

    URL = '/monitores/configuracion/'

    def setUp(self):
        self.client = Client(HTTP_HOST='programacion.testserver')
        u = User.objects.create_user(username='staff_mon', password='pass123')
        grupo, _ = Group.objects.get_or_create(name='area:programacion')
        u.groups.add(grupo)
        self.client.login(username='staff_mon', password='pass123')

    def test_staff_accede(self):
        resp = self.client.get(self.URL)
        self.assertEqual(resp.status_code, 200)


class ColegioSimulacroViewTest(TestCase):
    """CRUD + carga masiva del catálogo de colegios de simulacro."""

    URL = '/monitores/colegios/'

    def setUp(self):
        self.client = Client(HTTP_HOST='programacion.testserver')
        User.objects.create_superuser(username='admin_cs', password='pass123')
        self.client.login(username='admin_cs', password='pass123')

    def test_listado_accesible(self):
        ColegioSimulacro.objects.create(nombre='Colegio Norte')
        resp = self.client.get(self.URL)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Colegio Norte')

    def test_crear_colegio(self):
        self.client.post(self.URL, {
            'accion': 'add', 'nombre': 'Liceo Sur', 'codigo': 'LS-1',
            'departamento': 'Santander', 'ciudad': 'Bucaramanga',
        })
        c = ColegioSimulacro.objects.get(nombre='Liceo Sur')
        self.assertEqual(c.codigo, 'LS-1')
        self.assertEqual(c.ciudad, 'Bucaramanga')
        self.assertTrue(c.activo)

    def test_editar_colegio(self):
        c = ColegioSimulacro.objects.create(nombre='X', ciudad='A')
        self.client.post(self.URL, {
            'accion': 'edit', 'colegio_id': c.id, 'nombre': 'X', 'ciudad': 'B',
        })
        c.refresh_from_db()
        self.assertEqual(c.ciudad, 'B')

    def test_toggle_activo(self):
        c = ColegioSimulacro.objects.create(nombre='Toggle')
        self.client.post(self.URL, {'accion': 'toggle_activo', 'colegio_id': c.id})
        c.refresh_from_db()
        self.assertFalse(c.activo)

    def test_eliminar_colegio(self):
        c = ColegioSimulacro.objects.create(nombre='Borrar')
        self.client.post(self.URL, {'accion': 'del', 'colegio_id': c.id})
        self.assertFalse(ColegioSimulacro.objects.filter(id=c.id).exists())

    def test_gate_anonimo_redirige(self):
        self.client.logout()
        resp = self.client.get(self.URL)
        self.assertNotEqual(resp.status_code, 200)

    # ── Carga masiva ──────────────────────────────────────────
    def test_carga_masiva_crea(self):
        data = _xlsx_bytes([
            ('Colegio A', 'A1', 'Bogotá', 'Cundinamarca'),
            ('Colegio B', '', 'Cali', 'Valle del Cauca'),
        ])
        archivo = SimpleUploadedFile(
            'colegios.xlsx', data,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.client.post(self.URL, {'accion': 'cargar', 'archivo': archivo})
        self.assertEqual(ColegioSimulacro.objects.count(), 2)
        a = ColegioSimulacro.objects.get(nombre='Colegio A')
        self.assertEqual(a.codigo, 'A1')
        self.assertEqual(a.departamento, 'Cundinamarca')
        b = ColegioSimulacro.objects.get(nombre='Colegio B')
        self.assertIsNone(b.codigo)  # vacío → None

    def test_carga_masiva_idempotente(self):
        ColegioSimulacro.objects.create(nombre='Colegio A')
        data = _xlsx_bytes([
            ('Colegio A', '', '', ''),   # ya existe → omitido
            ('Colegio A', '', '', ''),   # duplicado dentro del archivo → omitido
            ('Colegio C', '', '', ''),   # nuevo
        ])
        archivo = SimpleUploadedFile('c.xlsx', data)
        self.client.post(self.URL, {'accion': 'cargar', 'archivo': archivo})
        self.assertEqual(ColegioSimulacro.objects.count(), 2)
        self.assertTrue(ColegioSimulacro.objects.filter(nombre='Colegio C').exists())

    def test_carga_masiva_columnas_desordenadas(self):
        # Encabezados en otro orden / mayúsculas → se mapean por nombre.
        data = _xlsx_bytes(
            [('Santander', 'Floridablanca', 'Colegio Z', 'CZ')],
            encabezados=('Departamento', 'Ciudad', 'Nombre', 'Codigo'))
        archivo = SimpleUploadedFile('z.xlsx', data)
        self.client.post(self.URL, {'accion': 'cargar', 'archivo': archivo})
        z = ColegioSimulacro.objects.get(nombre='Colegio Z')
        self.assertEqual(z.ciudad, 'Floridablanca')
        self.assertEqual(z.departamento, 'Santander')

    def test_carga_masiva_sin_columna_nombre(self):
        data = _xlsx_bytes([('A', 'B')], encabezados=('ciudad', 'codigo'))
        archivo = SimpleUploadedFile('bad.xlsx', data)
        self.client.post(self.URL, {'accion': 'cargar', 'archivo': archivo})
        self.assertEqual(ColegioSimulacro.objects.count(), 0)

    def test_carga_masiva_no_xlsx(self):
        archivo = SimpleUploadedFile('datos.csv', b'nombre,codigo\nA,B')
        self.client.post(self.URL, {'accion': 'cargar', 'archivo': archivo})
        self.assertEqual(ColegioSimulacro.objects.count(), 0)

    def test_descargar_plantilla(self):
        resp = self.client.post('/monitores/colegios/plantilla/')
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml', resp['Content-Type'])
        wb = load_workbook(io.BytesIO(resp.content))
        encabezados = [c.value for c in wb.active[1]]
        self.assertEqual(encabezados, ['nombre', 'codigo', 'ciudad', 'departamento'])


class SimulacroViewTest(TestCase):
    """CRUD de simulacros + asignación de monitores (Operaciones, programación)."""

    URL = '/monitores/'

    def setUp(self):
        self.client = Client(HTTP_HOST='programacion.testserver')
        User.objects.create_superuser(username='admin_sim', password='pass123')
        self.client.login(username='admin_sim', password='pass123')
        self.colegio = ColegioSimulacro.objects.create(nombre='Colegio Norte')
        self.g6 = Grado.objects.create(nombre='Sexto')
        self.g7 = Grado.objects.create(nombre='Séptimo')
        self.m1 = Monitor.objects.create(nombre='Ana', documento='111')
        self.m2 = Monitor.objects.create(nombre='Beto', documento='222')

    def test_listado_accesible(self):
        Simulacro.objects.create(colegio=self.colegio, fecha='2026-07-01')
        resp = self.client.get(self.URL)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Colegio Norte')

    def test_crear_simulacro_con_monitores_y_grados(self):
        self.client.post(self.URL, {
            'accion': 'add',
            'colegio': self.colegio.id,
            'fecha': '2026-07-10',
            'jornada': 'MANANA',
            'valor': '50000',
            'grados': [self.g6.id, self.g7.id],
            'monitores': [self.m1.id, self.m2.id],
            'observaciones': 'Simulacro de prueba',
        })
        s = Simulacro.objects.get()
        self.assertEqual(s.valor, 50000)
        self.assertEqual(s.jornada, 'MANANA')
        self.assertEqual(s.colegio_nombre, 'Colegio Norte')  # snapshot
        self.assertEqual(set(s.grados.values_list('id', flat=True)), {self.g6.id, self.g7.id})
        self.assertEqual(s.asignaciones.count(), 2)

    def test_crear_simulacro_sin_monitores(self):
        self.client.post(self.URL, {
            'accion': 'add', 'colegio': self.colegio.id,
            'fecha': '2026-07-10', 'jornada': 'TODO_DIA', 'valor': '0',
        })
        s = Simulacro.objects.get()
        self.assertEqual(s.asignaciones.count(), 0)

    def test_editar_sincroniza_monitores(self):
        s = Simulacro.objects.create(colegio=self.colegio, fecha='2026-07-10')
        AsignacionMonitor.objects.create(simulacro=s, monitor=self.m1)
        # Editar: quita m1, agrega m2.
        self.client.post(self.URL, {
            'accion': 'edit', 'simulacro_id': s.id,
            'colegio': self.colegio.id, 'fecha': '2026-07-10',
            'jornada': 'TODO_DIA', 'valor': '10000',
            'monitores': [self.m2.id],
        })
        s.refresh_from_db()
        self.assertEqual(s.valor, 10000)
        self.assertEqual(list(s.asignaciones.values_list('monitor_id', flat=True)), [self.m2.id])

    def test_eliminar_simulacro(self):
        s = Simulacro.objects.create(colegio=self.colegio, fecha='2026-07-10')
        AsignacionMonitor.objects.create(simulacro=s, monitor=self.m1)
        self.client.post(self.URL, {'accion': 'del', 'simulacro_id': s.id})
        self.assertEqual(Simulacro.objects.count(), 0)
        self.assertEqual(AsignacionMonitor.objects.count(), 0)  # CASCADE

    def test_monitor_con_asignacion_es_protegido(self):
        s = Simulacro.objects.create(colegio=self.colegio, fecha='2026-07-10')
        AsignacionMonitor.objects.create(simulacro=s, monitor=self.m1)
        from django.db.models import ProtectedError
        with self.assertRaises(ProtectedError):
            self.m1.delete()

    def test_borrar_colegio_conserva_snapshot(self):
        s = Simulacro.objects.create(colegio=self.colegio, fecha='2026-07-10')
        self.colegio.delete()  # SET_NULL
        s.refresh_from_db()
        self.assertIsNone(s.colegio_id)
        self.assertEqual(s.nombre_colegio, 'Colegio Norte')  # snapshot intacto

    def test_gate_requiere_personal(self):
        otro = Client(HTTP_HOST='programacion.testserver')
        User.objects.create_user(username='pepe', password='x')
        otro.login(username='pepe', password='x')
        resp = otro.get(self.URL)
        self.assertNotEqual(resp.status_code, 200)


class AvisoSimulacrosProximosTest(TestCase):
    """Fase 4: badge + banner + correo de simulacros próximos sin monitor."""

    URL = '/monitores/'

    def setUp(self):
        from datetime import timedelta
        from django.utils import timezone
        self.hoy = timezone.localdate()
        self.delta = timedelta
        self.client = Client(HTTP_HOST='programacion.testserver')
        User.objects.create_superuser(username='admin_av', password='pass123')
        self.client.login(username='admin_av', password='pass123')
        self.colegio = ColegioSimulacro.objects.create(nombre='Colegio Aviso')
        self.monitor = Monitor.objects.create(nombre='Vigi', documento='999')

    def _simulacro(self, dias, *, con_monitor=False):
        s = Simulacro.objects.create(
            colegio=self.colegio, fecha=self.hoy + self.delta(days=dias))
        if con_monitor:
            AsignacionMonitor.objects.create(simulacro=s, monitor=self.monitor)
        return s

    def test_selector_solo_proximos_sin_monitor(self):
        from programacion.monitores.avisos import simulacros_proximos_sin_monitor
        proximo_sin = self._simulacro(3)          # entra
        self._simulacro(3, con_monitor=True)      # tiene monitor → fuera
        self._simulacro(30)                       # lejano → fuera
        self._simulacro(-1)                        # pasado → fuera
        ids = list(simulacros_proximos_sin_monitor().values_list('id', flat=True))
        self.assertEqual(ids, [proximo_sin.id])

    def test_badge_en_contexto(self):
        self._simulacro(2)
        self._simulacro(5, con_monitor=True)
        resp = self.client.get(self.URL)
        self.assertEqual(resp.context['simulacros_sin_monitor_count'], 1)

    def test_banner_en_lista(self):
        self._simulacro(2)
        resp = self.client.get(self.URL)
        self.assertContains(resp, 'sin monitor asignado')

    def test_correo_envia_cuando_hay_proximos(self):
        from django.core import mail
        from programacion.monitores.avisos import notificar_simulacros_proximos
        self._simulacro(2)
        n = notificar_simulacros_proximos()
        self.assertEqual(n, 1)
        self.assertEqual(len(mail.outbox), 1)

    def test_correo_no_envia_sin_proximos(self):
        from django.core import mail
        from programacion.monitores.avisos import notificar_simulacros_proximos
        self._simulacro(30)  # lejano
        n = notificar_simulacros_proximos()
        self.assertEqual(n, 0)
        self.assertEqual(len(mail.outbox), 0)

    def test_command_no_rompe(self):
        from django.core.management import call_command
        self._simulacro(2)
        call_command('avisar_simulacros_proximos')  # no debe lanzar


class PagoMonitorModelTest(TestCase):
    """Fase 5: propiedades del modelo de pago (espejo de PagoRealizado)."""

    def setUp(self):
        from datetime import date
        self.colegio = ColegioSimulacro.objects.create(nombre='Col')
        self.monitor = Monitor.objects.create(nombre='Ana', documento='1')
        self.sim = Simulacro.objects.create(
            colegio=self.colegio, fecha=date(2026, 7, 8), valor=50000)

    def _pago(self, **kw):
        from programacion.monitores.models import PagoMonitor
        defaults = dict(monitor=self.monitor, simulacro=self.sim,
                        fecha=self.sim.fecha, valor=50000)
        defaults.update(kw)
        return PagoMonitor.objects.create(**defaults)

    def test_valor_base_es_el_valor(self):
        self.assertEqual(self._pago().valor_base, 50000)

    def test_total_suma_extras(self):
        from programacion.monitores.models import ExtraPagoMonitor
        p = self._pago()
        ExtraPagoMonitor.objects.create(pago=p, concepto='Transporte', valor=10000)
        self.assertEqual(p.total, 60000)
        self.assertEqual(p.total_extras, 10000)

    def test_fecha_pago_nullable_y_pagada(self):
        from django.utils import timezone
        p = self._pago()
        self.assertFalse(p.pagada)
        p.fecha_pago = timezone.now()
        p.save(update_fields=['fecha_pago'])
        self.assertTrue(p.pagada)

    def test_fila_historica_sin_lote(self):
        from django.utils import timezone
        p = self._pago(lote=None, fecha_pago=timezone.now())
        self.assertIsNone(p.lote_id)
        self.assertTrue(p.pagada)


class PrepararLoteMonitoresTest(TestCase):
    """Fase 5: materialización/envío/idempotencia de pagos de monitores."""

    def setUp(self):
        from datetime import date, timedelta
        self.colegio = ColegioSimulacro.objects.create(nombre='Col')
        self.m1 = Monitor.objects.create(nombre='Ana', documento='1')
        self.m2 = Monitor.objects.create(nombre='Beto', documento='2')
        # Simulacro en una semana conocida; ancla = lunes de esa fecha.
        self.fecha = date(2026, 7, 8)  # miércoles
        self.inicio = self.fecha - timedelta(days=self.fecha.weekday())
        self.fin = self.inicio + timedelta(days=6)  # semana completa (lunes–domingo)
        self.sim = Simulacro.objects.create(
            colegio=self.colegio, fecha=self.fecha, valor=50000)
        AsignacionMonitor.objects.create(simulacro=self.sim, monitor=self.m1)
        AsignacionMonitor.objects.create(simulacro=self.sim, monitor=self.m2)

    def test_preparar_crea_lote_y_filas(self):
        from programacion.monitores.pagos_servicios import preparar_lote_semana_monitores
        from programacion.monitores.models import LoteMonitores
        lote = preparar_lote_semana_monitores(self.inicio, self.fin)
        self.assertEqual(lote.estado, LoteMonitores.Estado.BORRADOR)
        self.assertEqual(lote.filas.count(), 2)  # un pago por monitor asignado
        self.assertEqual({f.valor for f in lote.filas.all()}, {50000})

    def test_preparar_idempotente_y_respeta_excluida(self):
        from programacion.monitores.pagos_servicios import preparar_lote_semana_monitores
        lote = preparar_lote_semana_monitores(self.inicio, self.fin)
        fila = lote.filas.first()
        fila.excluida = True
        fila.save(update_fields=['excluida'])
        preparar_lote_semana_monitores(self.inicio, self.fin)  # 2ª pasada
        self.assertEqual(lote.filas.count(), 2)  # no duplica
        fila.refresh_from_db()
        self.assertTrue(fila.excluida)           # no resucita

    def test_preparar_refresca_valor_si_cambia_el_simulacro(self):
        from programacion.monitores.pagos_servicios import preparar_lote_semana_monitores
        lote = preparar_lote_semana_monitores(self.inicio, self.fin)
        self.sim.valor = 70000
        self.sim.save(update_fields=['valor'])
        preparar_lote_semana_monitores(self.inicio, self.fin)
        self.assertEqual({f.valor for f in lote.filas.all()}, {70000})

    def test_preparar_elimina_autogenerada_si_se_quita_la_asignacion(self):
        from programacion.monitores.pagos_servicios import preparar_lote_semana_monitores
        lote = preparar_lote_semana_monitores(self.inicio, self.fin)
        AsignacionMonitor.objects.filter(simulacro=self.sim, monitor=self.m2).delete()
        preparar_lote_semana_monitores(self.inicio, self.fin)
        self.assertEqual(lote.filas.count(), 1)
        self.assertEqual(lote.filas.first().monitor_id, self.m1.id)

    def test_lote_enviado_no_se_remateriliza(self):
        from programacion.monitores.pagos_servicios import (
            preparar_lote_semana_monitores, enviar_lote_monitores)
        from programacion.monitores.models import LoteMonitores
        lote = preparar_lote_semana_monitores(self.inicio, self.fin)
        self.assertTrue(enviar_lote_monitores(lote, None))
        lote.refresh_from_db()
        self.assertEqual(lote.estado, LoteMonitores.Estado.ENVIADO)
        # Re-preparar crea un BORRADOR nuevo SIN tocar las filas congeladas.
        nuevo = preparar_lote_semana_monitores(self.inicio, self.fin)
        self.assertNotEqual(nuevo.id, lote.id)
        self.assertEqual(lote.filas.count(), 2)   # siguen en el ENVIADO
        self.assertEqual(nuevo.filas.count(), 0)  # nada que adoptar

    def test_enviar_desacopla_excluidas(self):
        from programacion.monitores.pagos_servicios import (
            preparar_lote_semana_monitores, enviar_lote_monitores)
        lote = preparar_lote_semana_monitores(self.inicio, self.fin)
        fila = lote.filas.first()
        fila.excluida = True
        fila.save(update_fields=['excluida'])
        self.assertTrue(enviar_lote_monitores(lote, None))
        self.assertEqual(lote.filas.count(), 1)   # la excluida se desacopló
        fila.refresh_from_db()
        self.assertIsNone(fila.lote_id)

    def test_enviar_sin_filas_enviables_no_marca(self):
        from programacion.monitores.pagos_servicios import (
            preparar_lote_semana_monitores, enviar_lote_monitores)
        from programacion.monitores.models import LoteMonitores
        lote = preparar_lote_semana_monitores(self.inicio, self.fin)
        lote.filas.update(excluida=True)
        self.assertFalse(enviar_lote_monitores(lote, None))
        lote.refresh_from_db()
        self.assertEqual(lote.estado, LoteMonitores.Estado.BORRADOR)

    def test_un_solo_borrador_por_semana(self):
        from django.db import IntegrityError, transaction
        from programacion.monitores.models import LoteMonitores
        LoteMonitores.objects.create(fecha_inicio=self.inicio, fecha_fin=self.fin)
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                LoteMonitores.objects.create(fecha_inicio=self.inicio, fecha_fin=self.fin)

    def test_preparar_pendientes_materializa_y_borra_vacios(self):
        from programacion.monitores.pagos_servicios import preparar_pendientes_monitores
        from programacion.monitores.models import LoteMonitores
        from datetime import date, timedelta
        # Mueve el simulacro al pasado para que entre en el backlog (<= hoy).
        self.sim.fecha = date.today() - timedelta(days=3)
        self.sim.save(update_fields=['fecha'])
        # Un borrador vacío de otra semana debe quedar barrido.
        LoteMonitores.objects.create(
            fecha_inicio=date(2020, 1, 6), fecha_fin=date(2020, 1, 10))
        n = preparar_pendientes_monitores(None)
        self.assertGreaterEqual(n, 1)
        self.assertFalse(LoteMonitores.objects.filter(
            fecha_inicio=date(2020, 1, 6)).exists())
        # La semana del simulacro sí tiene su borrador con filas.
        from programacion.monitores.models import PagoMonitor
        self.assertTrue(PagoMonitor.objects.exists())

    def test_simulacro_de_sabado_entra_en_su_semana(self):
        """Los simulacros de fin de semana caen en su lote (semana lunes–domingo)."""
        from datetime import date, timedelta
        from programacion.monitores.pagos_servicios import preparar_lote_semana_monitores
        sabado = date(2026, 7, 11)  # sábado
        lunes = sabado - timedelta(days=sabado.weekday())
        sim_sab = Simulacro.objects.create(
            colegio=self.colegio, fecha=sabado, valor=40000)
        AsignacionMonitor.objects.create(simulacro=sim_sab, monitor=self.m1)
        lote = preparar_lote_semana_monitores(lunes, lunes + timedelta(days=6))
        self.assertEqual(lote.filas.filter(simulacro=sim_sab).count(), 1)


class PagosMonitoresViewTest(TestCase):
    """Fase 6: UI de pagos de monitores en programación (lista, preparar, enviar,
    excluir, extras, detalle, export)."""

    def setUp(self):
        from datetime import date, timedelta
        self.client = Client(HTTP_HOST='programacion.testserver')
        User.objects.create_superuser(username='admin_pm', password='pass123')
        self.client.login(username='admin_pm', password='pass123')
        self.colegio = ColegioSimulacro.objects.create(nombre='Col', codigo='C1')
        self.m1 = Monitor.objects.create(nombre='Ana', apellido='Diaz', documento='1',
                                         banco='Bancolombia', tipo_cuenta='Ahorros',
                                         cuenta_bancaria='123')
        self.m2 = Monitor.objects.create(nombre='Beto', documento='2')
        # Simulacro ya ocurrido (fecha <= hoy) para que "Preparar pendientes" (backlog) lo tome.
        self.fecha = date.today() - timedelta(days=3)
        self.inicio = self.fecha - timedelta(days=self.fecha.weekday())
        self.fin = self.inicio + timedelta(days=6)
        self.sim = Simulacro.objects.create(
            colegio=self.colegio, fecha=self.fecha, valor=50000)
        AsignacionMonitor.objects.create(simulacro=self.sim, monitor=self.m1)
        AsignacionMonitor.objects.create(simulacro=self.sim, monitor=self.m2)

    def _preparar(self):
        from programacion.monitores.pagos_servicios import preparar_lote_semana_monitores
        return preparar_lote_semana_monitores(self.inicio, self.fin)

    def test_lista_get_ok(self):
        r = self.client.get('/monitores/pagos/')
        self.assertEqual(r.status_code, 200)
        self.assertTemplateUsed(r, 'monitores/pagos_monitores.html')
        # Sin pestaña "Sin informe" (monitores no tienen informe).
        self.assertNotContains(r, 'Sin informe')

    def test_preparar_crea_filas_y_aparecen(self):
        r = self.client.post('/monitores/pagos/preparar/', {'tab': 'pendiente'})
        self.assertEqual(r.status_code, 302)
        from programacion.monitores.models import PagoMonitor
        self.assertEqual(PagoMonitor.objects.count(), 2)
        r = self.client.get('/monitores/pagos/?tab=pendiente')
        self.assertContains(r, 'Ana')
        self.assertContains(r, 'Beto')

    def test_excluir_y_reincluir_toggle(self):
        lote = self._preparar()
        from programacion.monitores.models import PagoMonitor
        fila = lote.filas.first()
        self.client.post(f'/monitores/pagos/{fila.id}/excluir/', {'tab': 'pendiente'})
        fila.refresh_from_db()
        self.assertTrue(fila.excluida)
        self.client.post(f'/monitores/pagos/{fila.id}/excluir/', {'tab': 'pendiente'})
        fila.refresh_from_db()
        self.assertFalse(fila.excluida)

    def test_agregar_y_eliminar_extra(self):
        lote = self._preparar()
        from programacion.monitores.models import ExtraPagoMonitor
        fila = lote.filas.first()
        self.client.post(f'/monitores/pagos/{fila.id}/extra/',
                         {'concepto': 'Transporte', 'valor': '10000', 'tab': 'pendiente'})
        extra = ExtraPagoMonitor.objects.get(pago=fila)
        self.assertEqual(extra.valor, 10000)
        self.assertEqual(fila.total, 60000)
        self.client.post(f'/monitores/pagos/extra/{extra.id}/eliminar/', {'tab': 'pendiente'})
        self.assertFalse(ExtraPagoMonitor.objects.filter(id=extra.id).exists())

    def test_extra_no_editable_si_enviado(self):
        from programacion.monitores.pagos_servicios import enviar_lote_monitores
        from programacion.monitores.models import ExtraPagoMonitor
        lote = self._preparar()
        fila = lote.filas.first()
        enviar_lote_monitores(lote, None)  # ENVIADO → ya no editable
        self.client.post(f'/monitores/pagos/{fila.id}/extra/',
                         {'concepto': 'X', 'valor': '5000', 'tab': 'pendiente'})
        self.assertFalse(ExtraPagoMonitor.objects.filter(pago=fila).exists())

    def test_enviar_mueve_a_realizado(self):
        self._preparar()
        r = self.client.post('/monitores/pagos/enviar/', {'tab': 'pendiente'})
        self.assertEqual(r.status_code, 302)
        from programacion.monitores.models import LoteMonitores
        self.assertTrue(LoteMonitores.objects.filter(
            estado=LoteMonitores.Estado.ENVIADO).exists())
        r = self.client.get('/monitores/pagos/?tab=realizado')
        self.assertContains(r, 'Ana')

    def test_detalle_ok(self):
        lote = self._preparar()
        fila = lote.filas.first()
        r = self.client.get(f'/monitores/pagos/{fila.id}/')
        self.assertEqual(r.status_code, 200)
        self.assertTemplateUsed(r, 'monitores/pago_monitor_detalle.html')
        self.assertContains(r, self.colegio.nombre)

    def test_export_excel(self):
        self._preparar()
        r = self.client.post('/monitores/pagos/', {'tab': 'pendiente'})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(
            r['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('PagosMonitores', r['Content-Disposition'])

    def test_badge_cuenta_borrador_no_excluidas(self):
        from programacion.monitores.context_processors import pagos_monitores_por_revisar
        self._preparar()

        class _Req:
            area = 'programacion'
            es_personal_programacion = True
        ctx = pagos_monitores_por_revisar(_Req())
        self.assertEqual(ctx['pagos_monitores_por_revisar_count'], 2)

    def test_gate_no_personal(self):
        otro = Client(HTTP_HOST='programacion.testserver')
        User.objects.create_user(username='nadie', password='x')
        otro.login(username='nadie', password='x')
        r = otro.get('/monitores/pagos/')
        self.assertIn(r.status_code, (302, 403))
