"""
Pagos semanales a profesores — **revisión en programación** antes de financiera.

Calcula la liquidación semanal por fila `(profesor, colegio, día)` a partir de las
clases dictadas. Programación **prepara** el borrador de la semana (materializa las
filas en un `LotePagos` BORRADOR), las **revisa** (excluye filas, agrega costos extra
del desglose) y las **envía** a financiera (BORRADOR→ENVIADO). Solo son enviables las
filas cuyo día tiene **todos los informes completados** (`_claves_sin_informe`); las
excluidas o sin informe se desacoplan al enviar (`lote=None`) y pueden ir en un envío
posterior (N lotes ENVIADO por semana, máx. 1 BORRADOR). **Financiera solo ve las
filas de lotes ENVIADO**, ve el desglose, marca el pago y sube/elimina soportes
(financiera/pagos).

Los helpers de cálculo (`construir_contexto_pagos`, `filas_pagos_por_tab`,
`_generar_excel_pagos`, `_semana_label`) son **compartidos**: financiera los importa con
`modo='financiera'` para reusar exactamente el mismo cálculo (BD única, sin duplicar).
`_build_filas_pagos` (cálculo puro desde clases) alimenta la materialización;
`_filas_desde_lote` lee las filas ya persistidas (con su desglose).
"""
import io
from datetime import date, datetime, timedelta

from django.shortcuts import render, get_object_or_404, redirect
from django.http import HttpResponse, JsonResponse
from django.urls import reverse
from django.db.models import Count, Q
from django.utils import timezone
from django.contrib import messages
from django.contrib.auth.decorators import user_passes_test
from django.views.decorators.http import require_POST
from core.areas import es_personal_programacion
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from programacion.colegios.models import Clase
from programacion.configuracion.models import Profesor
from programacion.pagos.models import ExtraPago, LotePagos, PagoRealizado, SoportePagoProfesor
from programacion.pagos.notificaciones import notificar_pagos_enviados
from programacion.viaticos.views import _responder_soporte


MESES_ES_LARGO = ['', 'enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
                  'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']


def _tipo_cuenta_display(banco, tipo_cuenta):
    """Tipo de cuenta para la tabla/Excel: Daviplata no maneja Ahorros/Corriente,
    su producto se reporta como "Ahorros a la mano"."""
    return 'Ahorros a la mano' if banco == 'Daviplata' else tipo_cuenta


def _semana_label(fecha_inicio, fecha_fin):
    if fecha_inicio.month == fecha_fin.month:
        return (f"Semana del {fecha_inicio.day} al {fecha_fin.day} "
                f"de {MESES_ES_LARGO[fecha_inicio.month]} de {fecha_inicio.year}")
    return (f"Semana del {fecha_inicio.day} de {MESES_ES_LARGO[fecha_inicio.month]} "
            f"al {fecha_fin.day} de {MESES_ES_LARGO[fecha_fin.month]} de {fecha_fin.year}")


def _build_filas_pagos(fecha_inicio, fecha_fin):
    """
    Devuelve lista de dicts con los datos de pago agrupados por
    (fecha, profesor, colegio), sumando horas del día.
    """
    clases = list(
        Clase.objects
        .filter(
            fecha__gte=fecha_inicio, fecha__lte=fecha_fin,
            cancelada=False, es_evento=False,
        )
        .select_related(
            'profesor',
            'bloque__colegio__colegio',
        )
        .values(
            'fecha',
            'profesor_id',
            'profesor__nombre',
            'profesor__apellido',
            'profesor__documento',
            'profesor__cuenta_bancaria',
            'profesor__tipo_cuenta',
            'profesor__banco',
            'bloque__colegio_id',
            'bloque__colegio__colegio__nombre',
            'bloque__colegio__colegio__codigo',
            'bloque__colegio__valor_hora',
            'bloque__hora_inicio',
            'bloque__hora_fin',
        )
    )

    # Agrupar: (fecha, profesor_id, colegio_id) → minutos + info
    grupos = {}
    for c in clases:
        if not c['profesor_id']:
            continue
        key = (c['fecha'], c['profesor_id'], c['bloque__colegio_id'])
        hi = c['bloque__hora_inicio']
        hf = c['bloque__hora_fin']
        minutos = 0
        if hi and hf:
            minutos = max(0, (hf.hour * 60 + hf.minute) - (hi.hour * 60 + hi.minute))
        if key not in grupos:
            grupos[key] = {'minutos': 0, 'info': c}
        grupos[key]['minutos'] += minutos

    filas = []
    for (fecha, prof_id, col_id), data in sorted(grupos.items(), key=lambda x: (x[0][0], x[0][1])):
        info = data['info']
        horas = data['minutos'] / 60
        valor_hora = info['bloque__colegio__valor_hora'] or 0
        valor_total = round(horas * valor_hora)

        nombre_corto = Profesor.nombre_corto_de(
            info['profesor__nombre'], info['profesor__apellido'])

        banco       = info['profesor__banco'] or ''
        tipo_cuenta = _tipo_cuenta_display(banco, info['profesor__tipo_cuenta'] or '')

        filas.append({
            'fecha':       fecha,
            'profesor_id': prof_id,
            'colegio_id':  col_id,
            'docente':     nombre_corto,
            'documento':   info['profesor__documento'] or '',
            'num_cuenta':  info['profesor__cuenta_bancaria'] or '',
            'tipo_cuenta': tipo_cuenta,
            'banco':       banco,
            'colegio':     info['bloque__colegio__colegio__nombre'] or '',
            'codigo':      info['bloque__colegio__colegio__codigo'] or '',
            'horas':       horas,
            'valor_hora':  valor_hora,
            'valor_total': valor_total,
        })

    return filas


def _fila_desde_pago(p):
    """Construye el dict de fila (misma forma que `_build_filas_pagos`) desde una
    `PagoRealizado` persistida, añadiendo el desglose y el estado de pago.

    `valor_total` = total (valor base —posiblemente editado— + extras), para que la
    tabla/Excel sigan mostrando el monto final en la columna VALOR."""
    nombre_corto = p.profesor.nombre_corto

    banco       = p.profesor.banco or ''
    tipo_cuenta = _tipo_cuenta_display(banco, p.profesor.tipo_cuenta or '')

    extras = [{'id': e.id, 'concepto': e.concepto, 'valor': e.valor} for e in p.extras.all()]
    total_extras = sum(e['valor'] for e in extras)
    valor_base = p.valor_base

    return {
        'fecha':        p.fecha,
        'profesor_id':  p.profesor_id,
        'colegio_id':   p.colegio_id,
        'docente':      nombre_corto,
        'documento':    p.profesor.documento or '',
        'num_cuenta':   p.profesor.cuenta_bancaria or '',
        'tipo_cuenta':  tipo_cuenta,
        'banco':        banco,
        'colegio':      p.colegio.colegio.nombre or '',
        'codigo':       p.colegio.colegio.codigo or '',
        'horas':        p.horas,
        'valor_hora':   p.colegio.valor_hora or 0,
        'valor_base':   valor_base,
        'editado':      p.valor_base_editado is not None,
        'extras':       extras,
        'total_extras': total_extras,
        'valor_total':  valor_base + total_extras,
        'excluida':     p.excluida,
        'pago_id':      p.id,
        'fecha_pago':   p.fecha_pago,
        'marcado_por':  ((p.marcado_por.get_full_name() or p.marcado_por.username)
                         if p.marcado_por else '—'),
        'n_soportes':   getattr(p, 'n_soportes', p.soportes.count()),
    }


def _filas_desde_lote(lote):
    """Filas persistidas de un lote (incluye las **excluidas**, marcadas con `excluida`,
    para que programación pueda re-incluirlas). El llamador filtra excluidas si toca."""
    pagos = (
        PagoRealizado.objects
        .filter(lote=lote)
        .select_related('profesor', 'colegio__colegio', 'marcado_por')
        .prefetch_related('extras')
        .annotate(n_soportes=Count('soportes'))
        .order_by('fecha', 'profesor__nombre')
    )
    return [_fila_desde_pago(p) for p in pagos]


def _claves_sin_informe(desde=None, hasta=None):
    """Set de `(fecha, profesor_id, colegio_id)` con ≥1 clase **sin informe completado**.

    Misma clave de agrupación que `_build_filas_pagos` (el colegio sale del bloque).
    `actividades=''` cuenta como sin informe: es la señal de completitud
    (`Informe.completado`) y `guardar_informe` hace strip antes de persistir."""
    qs = (Clase.objects
          .filter(cancelada=False, es_evento=False, profesor__isnull=False)
          .filter(Q(informe__isnull=True) | Q(informe__actividades='')))
    if desde:
        qs = qs.filter(fecha__gte=desde)
    if hasta:
        qs = qs.filter(fecha__lte=hasta)
    return set(qs.values_list('fecha', 'profesor_id', 'bloque__colegio_id').distinct())


def preparar_lote_semana(inicio, fin, user=None):
    """Materializa (idempotente) el borrador de la semana canónica (lunes–domingo).

    Obtiene/crea el lote **BORRADOR** de la semana (puede convivir con N lotes ENVIADO
    de la misma semana: constraint parcial) y sincroniza sus filas con el cálculo desde
    clases:
    - crea filas para clases nuevas;
    - refresca `horas`/`valor` SOLO de filas sin override, no excluidas y no pagadas;
    - elimina filas **autogeneradas** (sin override, sin extras, no excluidas, no pagadas)
      cuya clase ya no existe (cancelada);
    - las filas de lotes **ENVIADO** quedan congeladas: ni se adoptan, ni se refrescan,
      ni se borran.

    Nunca resucita excluidas ni pisa valores editados. Adopta al lote filas existentes
    con la misma tripleta (históricas `lote=NULL` y filas desacopladas al enviar),
    respetando el unique `(profesor, colegio, fecha)`."""
    lote, _ = LotePagos.objects.get_or_create(
        fecha_inicio=inicio, fecha_fin=fin, estado=LotePagos.Estado.BORRADOR)

    calc = {(f['profesor_id'], f['colegio_id'], f['fecha']): f
            for f in _build_filas_pagos(inicio, fin)}
    existentes = {
        (p.profesor_id, p.colegio_id, p.fecha): p
        for p in PagoRealizado.objects
            .filter(fecha__gte=inicio, fecha__lte=fin)
            .select_related('lote')
            .prefetch_related('extras')
    }

    def _congelada(p):
        return p.lote_id and p.lote_id != lote.id and p.lote.estado == LotePagos.Estado.ENVIADO

    for key, f in calc.items():
        p = existentes.get(key)
        if p is None:
            PagoRealizado.objects.create(
                lote=lote,
                profesor_id=f['profesor_id'], colegio_id=f['colegio_id'], fecha=f['fecha'],
                horas=f['horas'], valor=f['valor_total'],
            )
            continue
        if _congelada(p):
            continue
        cambios = []
        if p.lote_id != lote.id:
            p.lote = lote
            cambios.append('lote')
        if (p.valor_base_editado is None and not p.excluida and p.fecha_pago is None
                and (p.horas != f['horas'] or p.valor != f['valor_total'])):
            p.horas = f['horas']
            p.valor = f['valor_total']
            cambios += ['horas', 'valor']
        if cambios:
            p.save(update_fields=cambios)

    # Limpiar filas autogeneradas de ESTE lote que ya no salen en el cálculo.
    for key, p in existentes.items():
        if key in calc or p.lote_id != lote.id:
            continue
        if (p.valor_base_editado is None and not p.excluida and p.fecha_pago is None
                and not p.extras.all()):
            p.delete()

    return lote


def enviar_lote(lote, user):
    """BORRADOR → ENVIADO. A partir de aquí financiera lo ve y programación no edita.
    El envío es **definitivo por lote** (no hay reabrir).

    Antes de marcar, **desacopla** (`lote=None`) las filas no enviables — excluidas o
    cuyo día tiene clases sin informe completado — para que el lote ENVIADO contenga
    exactamente lo enviado; un "Preparar pendientes" posterior las re-adopta a un
    BORRADOR nuevo y podrán ir en otro envío. Devuelve True si el lote quedó ENVIADO;
    False si no tenía ninguna fila enviable (sigue BORRADOR)."""
    claves_sin_informe = _claves_sin_informe(lote.fecha_inicio, lote.fecha_fin)
    no_enviables = [
        p.pk for p in lote.filas.all()
        if p.excluida or (p.fecha, p.profesor_id, p.colegio_id) in claves_sin_informe
    ]
    if no_enviables:
        PagoRealizado.objects.filter(pk__in=no_enviables).update(lote=None)
    if not lote.filas.exists():
        return False
    lote.estado = LotePagos.Estado.ENVIADO
    lote.enviado_en = timezone.now()
    lote.enviado_por = user
    lote.save(update_fields=['estado', 'enviado_en', 'enviado_por', 'actualizado_en'])
    return True


def preparar_pendientes(user=None):
    """Materializa el **backlog completo**: prepara (idempotente) el BORRADOR de todas
    las semanas con clases hasta hoy (las filas ya enviadas quedan congeladas en sus
    lotes ENVIADO; las desacopladas al enviar se re-adoptan). Así la lista muestra todo
    lo pendiente por enviar sin preparar semana por semana. Devuelve nº de semanas."""
    hoy = date.today()
    fechas = (Clase.objects
              .filter(fecha__lte=hoy, cancelada=False, es_evento=False)
              .values_list('fecha', flat=True).distinct())
    lunes_set = {f - timedelta(days=f.weekday()) for f in fechas}
    for lunes in lunes_set:
        # Semana completa (lunes–domingo): hay profesores que dictan en fin de semana
        # (refuerzos, simulacros, etc.). Con lunes–viernes esas clases caían fuera de la
        # ventana de su propia semana y NUNCA se materializaban → su informe no podía
        # llegar a "por enviar". Mismo criterio que monitores (`pagos_servicios`).
        preparar_lote_semana(lunes, lunes + timedelta(days=6), user)
    # Borradores que quedaron vacíos (todo su contenido se envió o se canceló) no
    # aportan nada al backlog: fuera.
    LotePagos.objects.filter(
        estado=LotePagos.Estado.BORRADOR, filas__isnull=True).delete()
    return len(lunes_set)


def _desglose_texto(f):
    """Texto del desglose de una fila para el Excel: ``Base: $X; Concepto: $Y``.
    Vacío si la fila no tiene costos extra (la columna VALOR ya muestra el total)."""
    extras = f.get('extras') or []
    if not extras:
        return ''
    partes = [f"Base: ${f.get('valor_base', f['valor_total']):,}"]
    partes += [f"{e['concepto']}: ${e['valor']:,}" for e in extras]
    return '; '.join(partes)


def _generar_excel_pagos(filas, semana_label):
    """Genera el Excel de pagos. Una fila por pago `(profesor, colegio, día)` con el monto
    final en VALOR; la columna DESGLOSE detalla base + costos extra cuando los hay."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Pagos'

    NUM_COLS = 10
    COLS = ['FECHA', 'DOCENTE', 'DOCUMENTO', 'N° DE CUENTA',
            'TIPO DE CUENTA', 'BANCO', 'COLEGIO', 'CODIGO', 'VALOR', 'DESGLOSE']

    HEADER_FILL  = 'FF1F3864'  # azul oscuro
    SUBHDR_FILL  = 'FFD6E4F0'  # azul claro
    TOTAL_FILL   = 'FFD6DCE4'

    thin = Side(style='thin', color='000000')
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _celda(row, col, valor='', bold=False, fill=None, color='000000',
               h='center', v='center', fmt=None, borde_=True):
        cell = ws.cell(row, col, valor)
        cell.font = Font(name='Arial', size=10, bold=bold, color=color)
        cell.alignment = Alignment(horizontal=h, vertical=v, wrap_text=True)
        if fill:
            cell.fill = PatternFill('solid', fgColor=fill)
        if borde_:
            cell.border = borde
        if fmt:
            cell.number_format = fmt
        return cell

    def _bordear_combinada(r1, c1, r2, c2, fill=None):
        """Aplica borde (y relleno) a TODAS las celdas de un rango combinado. openpyxl solo
        dibuja el borde de la celda superior-izquierda al combinar; sin esto el contorno del
        bloque combinado queda incompleto en Excel."""
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                cell = ws.cell(r, c)
                cell.border = borde
                if fill:
                    cell.fill = PatternFill('solid', fgColor=fill)

    # Fila 1: título semana
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NUM_COLS)
    _celda(1, 1, semana_label, bold=True, fill='FFD9E1F2', color='FF1F3864')
    _bordear_combinada(1, 1, 1, NUM_COLS, fill='FFD9E1F2')
    ws.row_dimensions[1].height = 20

    # Fila 2: cabeceras
    for ci, col_name in enumerate(COLS, start=1):
        _celda(2, ci, col_name, bold=True, fill='FFB8CCE4', color='FF1F3864')
    ws.row_dimensions[2].height = 22

    # Filas de datos
    MESES_ES_ABREV = ['', 'Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun',
                      'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']

    for i, f in enumerate(filas, start=3):
        fill_row = 'FFFFFFFF' if i % 2 == 1 else 'FFF2F6FC'
        fecha_str = f"{f['fecha'].day:02d}/{f['fecha'].month:02d}/{f['fecha'].year}"
        _celda(i, 1, fecha_str,       fill=fill_row, h='center')
        _celda(i, 2, f['docente'],    fill=fill_row, h='left')
        _celda(i, 3, f['documento'],  fill=fill_row, h='center')
        _celda(i, 4, f['num_cuenta'], fill=fill_row, h='center')
        _celda(i, 5, f['tipo_cuenta'],fill=fill_row, h='center')
        _celda(i, 6, f['banco'],      fill=fill_row, h='center')
        _celda(i, 7, f['colegio'],    fill=fill_row, h='left')
        _celda(i, 8, f['codigo'],     fill=fill_row, h='center')
        # Valor como número para que Excel pueda sumar
        cell_val = ws.cell(i, 9, f['valor_total'])
        cell_val.font = Font(name='Arial', size=10)
        cell_val.alignment = Alignment(horizontal='right', vertical='center')
        cell_val.fill = PatternFill('solid', fgColor=fill_row)
        cell_val.border = borde
        cell_val.number_format = '"$"#,##0'
        _celda(i, 10, _desglose_texto(f), fill=fill_row, h='left')
        ws.row_dimensions[i].height = 18

    # Fila TOTAL
    total_row = 3 + len(filas)
    ws.merge_cells(start_row=total_row, start_column=1,
                   end_row=total_row, end_column=8)
    _celda(total_row, 1, 'TOTAL', bold=True, fill=TOTAL_FILL, h='right')
    _bordear_combinada(total_row, 1, total_row, 8, fill=TOTAL_FILL)
    total_val = sum(f['valor_total'] for f in filas)
    cell_t = ws.cell(total_row, 9, total_val)
    cell_t.font = Font(name='Arial', size=10, bold=True)
    cell_t.alignment = Alignment(horizontal='right', vertical='center')
    cell_t.fill = PatternFill('solid', fgColor=TOTAL_FILL)
    cell_t.border = borde
    cell_t.number_format = '"$"#,##0'
    _celda(total_row, 10, '', fill=TOTAL_FILL)
    ws.row_dimensions[total_row].height = 22

    # Anchos de columna: auto-ajuste al contenido real (longitud máx. de la columna),
    # acotado por un mínimo (legibilidad de la cabecera) y un máximo (evita columnas
    # gigantes en DESGLOSE/COLEGIO). El ancho de openpyxl ≈ nº de caracteres.
    MIN_W, MAX_W = 8, 45

    def _texto_largo(cell):
        """Longitud del texto tal como se ve: los montos numéricos se miden ya
        formateados ($ + separadores de miles), no por su valor crudo."""
        v = cell.value
        if isinstance(v, (int, float)):
            return len(f"${v:,.0f}")
        return len(str(v or ''))

    for ci in range(1, NUM_COLS + 1):
        ancho = max((_texto_largo(ws.cell(r, ci))
                     for r in range(2, total_row + 1)), default=MIN_W)
        ws.column_dimensions[get_column_letter(ci)].width = \
            max(MIN_W, min(ancho + 2, MAX_W))

    ws.freeze_panes = 'A3'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse_fecha(s):
    try:
        return datetime.strptime(s, '%Y-%m-%d').date()
    except (ValueError, TypeError):
        return None


def _rango_de(get):
    """Rango de fechas del filtro (o `(None, None)` si no se aplicó ninguno).

    El filtro de la barra manda `semana` (=desde) y `hasta`. Si no llegan, no se restringe
    por fecha → se muestra **todo** el backlog (todos los pendientes)."""
    return _parse_fecha(get.get('semana') or get.get('desde')), _parse_fecha(get.get('hasta'))


def _filas_rango(estado, desde, hasta):
    """Filas persistidas (`PagoRealizado`) cuyo lote está en `estado`, opcionalmente
    acotadas por fecha. Mantiene el desglose vía `_fila_desde_pago`."""
    qs = (PagoRealizado.objects
          .filter(lote__estado=estado)
          .select_related('profesor', 'colegio__colegio', 'marcado_por')
          .prefetch_related('extras')
          .annotate(n_soportes=Count('soportes'))
          .order_by('fecha', 'profesor__nombre'))
    if desde:
        qs = qs.filter(fecha__gte=desde)
    if hasta:
        qs = qs.filter(fecha__lte=hasta)
    return [_fila_desde_pago(p) for p in qs]


def _suma_valor(filas):
    return sum(f['valor_total'] for f in filas if not f['excluida'])


def _label_rango(desde, hasta):
    if desde and hasta:
        return _semana_label(desde, hasta)
    if desde:
        return f'Desde el {desde.day} de {MESES_ES_LARGO[desde.month]} de {desde.year}'
    if hasta:
        return f'Hasta el {hasta.day} de {MESES_ES_LARGO[hasta.month]} de {hasta.year}'
    return 'Todos los registros'


def construir_contexto_pagos(get, *, modo='programacion'):
    """Arma el contexto de la página de pagos como **backlog** (todas las semanas).

    Sin filtro de fechas → muestra todo lo pendiente; con filtro (`semana`/`hasta`) acota.
    - `programacion`: pestaña *pendiente* = filas por **enviar** (lote BORRADOR, incluye las
      excluidas para poder re-incluirlas); pestaña *sin_informe* = filas en BORRADOR cuyo
      día tiene clases **sin informe completado** (editables pero NO enviables, para
      recordarle al docente); pestaña *realizado* = filas ya **enviadas** (lote ENVIADO;
      muestran si financiera ya las pagó).
    - `financiera`: **solo** lotes ENVIADO (sin excluidas), partidas por pago (`fecha_pago`):
      pendiente = por pagar; realizado = pagadas.
    Mantiene las claves históricas (`filas_pendientes`, `filas_realizadas`, `filas`, `total_*`)."""
    desde, hasta = _rango_de(get)
    tab = get.get('tab', 'pendiente')

    filas_sin_informe = []
    if modo == 'financiera':
        enviadas = [f for f in _filas_rango(LotePagos.Estado.ENVIADO, desde, hasta)
                    if not f['excluida']]
        filas_pendientes = [f for f in enviadas if f['fecha_pago'] is None]
        filas_realizadas = [f for f in enviadas if f['fecha_pago'] is not None]
    else:
        claves_si = _claves_sin_informe(desde, hasta)
        filas_pendientes = []                                                      # por enviar
        for f in _filas_rango(LotePagos.Estado.BORRADOR, desde, hasta):
            if (f['fecha'], f['profesor_id'], f['colegio_id']) in claves_si:
                f['sin_informe'] = True
                filas_sin_informe.append(f)
            else:
                filas_pendientes.append(f)
        filas_realizadas = _filas_rango(LotePagos.Estado.ENVIADO, desde, hasta)    # enviadas

    if tab == 'realizado':
        filas_tab = filas_realizadas
    elif tab == 'sin_informe' and modo == 'programacion':
        filas_tab = filas_sin_informe
    else:
        filas_tab = filas_pendientes

    ctx = {
        'fecha_inicio':     desde.isoformat() if desde else '',
        'fecha_fin':        hasta.isoformat() if hasta else '',
        'fecha_max':        date.today().isoformat(),
        'semana_label':     _label_rango(desde, hasta),
        'filtro_aplicado':  bool(desde or hasta),
        'tab':              tab,
        'filas_pendientes': filas_pendientes,
        'filas_realizadas': filas_realizadas,
        'filas':            filas_tab,
        'total_valor':      _suma_valor(filas_tab),
        'total_pendiente':  _suma_valor(filas_pendientes),
        'total_realizado':  _suma_valor(filas_realizadas),
        # Detalle (horas/valor-hora/extras) por pago para el modal (i)/desglose. Ambas áreas.
        'detalles_pagos':   _detalles_pagos(
            filas_pendientes + filas_sin_informe + filas_realizadas),
    }
    if modo == 'programacion':
        por_enviar = [f for f in filas_pendientes if not f['excluida']]
        ctx.update({
            'tab_label_pendiente':   'Por enviar',
            'tab_label_realizado':   'Enviados',
            # La presencia de este label habilita la tercera pestaña en los parciales
            # compartidos (financiera no la define → no la dibuja).
            'tab_label_sin_informe': 'Sin informe',
            'filas_sin_informe':     filas_sin_informe,
            'puede_enviar':          len(por_enviar) > 0,
            'total_por_enviar':      len(por_enviar),
        })
    return ctx


def _detalles_pagos(filas):
    """Mapa `pago_id → {docente, horas, valor_hora, valor_base, extras, total}` para el
    modal de detalle/gestión (se serializa con `json_script`)."""
    return {
        f['pago_id']: {
            'docente':    f['docente'],
            'horas':      f['horas'],
            'valor_hora': f['valor_hora'],
            'valor_base': f['valor_base'],
            'total':      f['valor_total'],
            'extras':     f['extras'],
            'editable':   f['fecha_pago'] is None and not f['excluida'],
        }
        for f in filas if f['pago_id']
    }


def filas_pagos_por_tab(fecha_inicio, fecha_fin, tab, *, modo='programacion'):
    """Filas para el Excel del tab, como backlog acotado al rango (vacío = todo). Excluye
    siempre las filas excluidas. Mismo `modo` que `construir_contexto_pagos`."""
    get = {'tab': tab}
    if fecha_inicio:
        get['semana'] = fecha_inicio.isoformat() if hasattr(fecha_inicio, 'isoformat') else fecha_inicio
    if fecha_fin:
        get['hasta'] = fecha_fin.isoformat() if hasattr(fecha_fin, 'isoformat') else fecha_fin
    ctx = construir_contexto_pagos(get, modo=modo)
    return [f for f in ctx['filas'] if not f['excluida']]


@user_passes_test(es_personal_programacion, login_url='login')
def pagos_lista(request):
    """GET: página de pagos (backlog, tabs por enviar/enviados). POST: descarga Excel.

    En GET materializa el backlog (idempotente) para que "Por enviar" muestre siempre
    todo lo pendiente con informe completado sin depender de un botón manual."""
    if request.method == 'GET':
        preparar_pendientes(request.user)
        return render(request, 'pagos/pagos.html',
                      construir_contexto_pagos(request.GET, modo='programacion'))

    # POST: descarga Excel del tab activo (rango opcional; vacío = todo el backlog).
    fi = _parse_fecha(request.POST.get('fecha_inicio', ''))
    ff = _parse_fecha(request.POST.get('fecha_fin', ''))
    tab = request.POST.get('tab', 'pendiente')

    filas = filas_pagos_por_tab(fi, ff, tab, modo='programacion')
    excel_bytes = _generar_excel_pagos(filas, _label_rango(fi, ff))

    sufijo = {'realizado': 'Enviados', 'sin_informe': 'SinInforme'}.get(tab, 'PorEnviar')
    rango = f"{fi.strftime('%Y%m%d')}_{ff.strftime('%Y%m%d')}" if fi and ff else 'todos'
    response = HttpResponse(
        excel_bytes,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="Pagos_{sufijo}_{rango}.xlsx"'
    return response


# ══════════════════════════════════════════════════════════════
# REVISIÓN (programación): preparar pendientes · excluir · extras · enviar
# ══════════════════════════════════════════════════════════════
# Programación revisa el backlog (todas las semanas) y lo envía a financiera. Todas las
# acciones son POST + redirect a la lista (re-renderiza con datos frescos); excluir/extras
# exigen que la fila esté en BORRADOR. El envío es definitivo (no hay reabrir).

def _volver_a_lista(request):
    """Redirige a la lista conservando el filtro de fechas (si lo hay) y la pestaña."""
    tab = request.POST.get('tab', 'pendiente')
    semana = request.POST.get('semana', '')
    hasta = request.POST.get('hasta', '')
    q = f'?tab={tab}'
    if semana:
        q += f'&semana={semana}'
    if hasta:
        q += f'&hasta={hasta}'
    return redirect(reverse('pagos_lista') + q)


def _pago_editable(pago):
    """True si la fila pertenece a un lote en BORRADOR (editable por programación)."""
    return bool(pago.lote_id and pago.lote.estado == LotePagos.Estado.BORRADOR)


@user_passes_test(es_personal_programacion, login_url='login')
@require_POST
def pagos_preparar(request):
    """Prepara el backlog: materializa todas las semanas con clases hasta hoy no enviadas."""
    n = preparar_pendientes(request.user)
    messages.success(request, f'Pendientes preparados ({n} semana(s)).')
    return _volver_a_lista(request)


@user_passes_test(es_personal_programacion, login_url='login')
@require_POST
def pagos_enviar(request):
    """Envía a financiera **todo lo visible** (lotes BORRADOR con filas no excluidas dentro
    del rango filtrado, o todo si no hay filtro). El envío es definitivo."""
    desde, hasta = _rango_de(request.POST)
    qs = PagoRealizado.objects.filter(
        lote__estado=LotePagos.Estado.BORRADOR, excluida=False)
    if desde:
        qs = qs.filter(fecha__gte=desde)
    if hasta:
        qs = qs.filter(fecha__lte=hasta)
    lote_ids = set(qs.values_list('lote_id', flat=True))
    lotes = LotePagos.objects.filter(id__in=lote_ids, estado=LotePagos.Estado.BORRADOR)
    # `enviar_lote` desacopla las no enviables (excluidas / sin informe) y devuelve False
    # si el lote quedó sin nada que enviar (sigue BORRADOR).
    enviados = [lote for lote in lotes if enviar_lote(lote, request.user)]
    if enviados:
        messages.success(request, f'Enviado a financiera ({len(enviados)} semana(s)).')
        # Aviso por correo a financiera (best-effort, igual que viáticos): un fallo de
        # correo no debe afectar el envío que ya quedó persistido.
        notificar_pagos_enviados(enviados, request)
    else:
        messages.error(request, 'No hay pagos enviables. Las filas excluidas o con '
                                'clases sin informe no se envían.')
    return _volver_a_lista(request)


@user_passes_test(es_personal_programacion, login_url='login')
@require_POST
def pagos_excluir_fila(request, pago_id):
    """Excluye o vuelve a incluir una fila del envío a financiera (toggle)."""
    pago = get_object_or_404(PagoRealizado.objects.select_related('lote'), pk=pago_id)
    if not _pago_editable(pago):
        messages.error(request, 'Esta fila ya no es editable.')
    else:
        pago.excluida = not pago.excluida
        pago.save(update_fields=['excluida'])
        messages.success(request, 'Fila excluida del envío.' if pago.excluida else 'Fila incluida.')
    return _volver_a_lista(request)


@user_passes_test(es_personal_programacion, login_url='login')
@require_POST
def pagos_agregar_extra(request, pago_id):
    """Agrega un costo extra (concepto + valor) al desglose de una fila."""
    pago = get_object_or_404(PagoRealizado.objects.select_related('lote'), pk=pago_id)
    if not _pago_editable(pago):
        messages.error(request, 'Esta fila ya no es editable.')
        return _volver_a_lista(request)
    concepto = (request.POST.get('concepto', '') or '').strip()
    raw = (request.POST.get('valor', '') or '').strip().replace('.', '').replace(',', '')
    try:
        valor = int(raw)
    except (ValueError, TypeError):
        valor = 0
    if not concepto or valor <= 0:
        messages.error(request, 'Indica un concepto y un valor mayor que cero.')
        return _volver_a_lista(request)
    ExtraPago.objects.create(pago=pago, concepto=concepto, valor=valor, orden=pago.extras.count())
    messages.success(request, 'Costo extra agregado.')
    return _volver_a_lista(request)


@user_passes_test(es_personal_programacion, login_url='login')
@require_POST
def pagos_eliminar_extra(request, extra_id):
    """Elimina un costo extra del desglose."""
    extra = get_object_or_404(ExtraPago.objects.select_related('pago__lote'), pk=extra_id)
    if not _pago_editable(extra.pago):
        messages.error(request, 'Esta fila ya no es editable.')
    else:
        extra.delete()
        messages.success(request, 'Costo extra eliminado.')
    return _volver_a_lista(request)


# ══════════════════════════════════════════════════════════════
# DETALLE DE PAGO — SOLO LECTURA (el comprobante lo gestiona financiera)
# ══════════════════════════════════════════════════════════════
# Marcar el pago y subir/eliminar soportes es responsabilidad de financiera
# (financiera/pagos). Programación solo calcula, ve el listado y consulta los
# comprobantes adjuntos en solo lectura.

@user_passes_test(es_personal_programacion, login_url='login')
def pagos_detalle(request, pago_id):
    """Detalle de solo lectura de un pago realizado: datos del docente/colegio y
    los soportes adjuntos (ver/descargar, sin subir ni eliminar)."""
    pago = get_object_or_404(
        PagoRealizado.objects
        .select_related('profesor', 'colegio__colegio', 'marcado_por')
        .prefetch_related('soportes', 'soportes__subido_por'),
        pk=pago_id,
    )
    return render(request, 'pagos/pago_detalle.html', {'pago': pago})


@user_passes_test(es_personal_programacion, login_url='login')
def soporte_descargar(request, soporte_id):
    """Ver (`?inline=1`) o descargar un soporte de pago desde programación
    (solo lectura). Mismo proxy server-side que financiera; difiere en el gate."""
    soporte = get_object_or_404(SoportePagoProfesor, pk=soporte_id)
    return _responder_soporte(soporte, inline=request.GET.get('inline') == '1')
