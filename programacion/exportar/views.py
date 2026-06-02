"""
Exportación masiva de horarios y pagos.

Genera archivos Excel con openpyxl:
- Por profesor: un Excel con todas sus clases (colegio + particulares) en el período.
- Por colegio: un Excel con la matriz grado×fecha de todos los bloques.
- Pagos: lista semanal de liquidaciones pendientes o realizadas.

El ZIP de exportación se construye completamente en memoria (BytesIO) para
evitar archivos temporales en disco — importante en Render (filesystem efímero).
"""

import io
import zipfile
from datetime import date, datetime
from collections import defaultdict

from django.shortcuts import render
from django.http import HttpResponse, JsonResponse
from django.contrib.auth.decorators import user_passes_test
from core.areas import es_personal_programacion
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from programacion.configuracion.models import Profesor, NombreLibro, Unidad, ColegioAnio
from programacion.colegios.models import Clase, Asignacion, ClaseParticular, Bloque
from programacion.colegios.utils import extraer_numero_grado, ordenar_grados
from programacion.exportar.models import PagoRealizado


# ══════════════════════════════════════════════════════════════
# UTILIDADES DE FECHAS Y ASIGNACIONES
# ══════════════════════════════════════════════════════════════

# Usamos dicts propios en lugar de strftime('%a'/'%b') porque el servidor
# Render no garantiza locale en español — strftime devolvería inglés en producción.
DIAS_ES  = {0:'Lun', 1:'Mar', 2:'Mié', 3:'Jue', 4:'Vie', 5:'Sáb', 6:'Dom'}
MESES_ES = {1:'Ene', 2:'Feb', 3:'Mar', 4:'Abr', 5:'May', 6:'Jun',
             7:'Jul', 8:'Ago', 9:'Sep', 10:'Oct', 11:'Nov', 12:'Dic'}


def _safe_url(url: str | None) -> str | None:
    """Devuelve la URL solo si es HTTP/HTTPS; de lo contrario None.

    Defensa explícita contra URIs tipo 'javascript:' o 'data:' que openpyxl
    insertaría literalmente como hiperenlace en el Excel, abriendo XSS si el
    usuario hace clic desde Excel Online.
    """
    if url and str(url).startswith(('https://', 'http://')):
        return str(url)
    return None


def _fecha_label(d):
    return f"{DIAS_ES[d.weekday()]} {d.day:02d}-{MESES_ES[d.month]}"


def _build_asignaciones_map(colegios_ids):
    """Pre-carga todas las asignaciones de los colegios dados en un único query.

    Retorna dict {(colegio_id, grado_nombre): [Asignacion, ...]}

    Se usa para resolver qué libro tiene asignado un grado en una fecha concreta
    sin hacer queries dentro de los loops de generación de Excel (evita N+1).
    """
    asig_map = defaultdict(list)
    for a in Asignacion.objects.filter(
        colegio_id__in=colegios_ids
    ).select_related('grado', 'libro'):
        asig_map[(a.colegio_id, a.grado.nombre)].append(a)
    return dict(asig_map)


def _titulo_libro(asig_map, colegio_id, grado, fecha):
    """Retorna el nombre del libro asignado al grado en la fecha dada, o '' si no hay.

    Usa el mapa pre-cargado por _build_asignaciones_map para evitar queries adicionales.
    Acepta tanto instancias de Grado como strings (para compatibilidad con .values()).
    """
    grado_nombre = grado.nombre if hasattr(grado, 'nombre') else str(grado)
    for a in asig_map.get((colegio_id, grado_nombre), []):
        if a.fecha_inicio and a.fecha_fin and a.fecha_inicio <= fecha <= a.fecha_fin:
            if a.libro:
                return a.libro.nombre
            return 'Sin Libro Asignado'
    return ''

# ══════════════════════════════════════════════════════════════
# CONSTANTES DE ESTILO
# ══════════════════════════════════════════════════════════════

COL_WIDTH    = 30
LABEL_GRAY   = 'FFD4D4D4'
FECHA_GRAY   = 'FFD0D0D0'
HEADER_DARK  = 'FFB0B3B2'
NEGRO        = 'FF000000'
BLANCO       = 'FFFFFFFF'
PART_COLOR   = 'FFFFF9C4'

# Colores pastel por materia — formato ARGB (FF + hex) para compatibilidad Excel Online
COLORES_MATERIA = {
    'matemáticas':            'FFC6EFCE',  # verde
    'biología':               'FFFFB3C6',  # rosado fuerte
    'química':                'FFD4B8E0',  # morado pastel
    'física':                 'FFF4B183',  # naranja
    'sociales':               'FFFFE699',  # amarillo
    'ciencias naturales':     'FFA9D18E',  # verde oscuro
    'inglés':                 'FF9DC3E6',  # azul
    'lectura crítica':        'FFB4C7E7',  # azul gris
    'orientación profesional':'FFF2C078',  # durazno
}
COLOR_DEFAULT  = 'FFF2F2F2'
COLOR_EVENTO   = 'FFFFD700'
COLOR_CANCELADA = 'FFFFB3B3'

_sin_borde  = Border()  # sin bordes


def _e(ws, row, col, valor='', bold=False, fill=None,
       h='center', v='center', size=10, color='000000',
       link_color=False, hyperlink=None, border=False):
    """Escribe y formatea una celda."""
    cell = ws.cell(row, col, valor)
    kw = {'name': 'Arial', 'size': size, 'bold': bold, 'color': color}
    if link_color:
        kw['color'] = '0563C1'
        kw['underline'] = 'single'
    cell.font      = Font(**kw)
    cell.alignment = Alignment(horizontal=h, vertical=v, wrap_text=True)
    if border:
        _thin_side = Side(style='thin', color='000000')
        cell.border = Border(left=_thin_side, right=_thin_side,
                             top=_thin_side, bottom=_thin_side)
    if fill:
        cell.fill = PatternFill('solid', fgColor=fill)
    if hyperlink:
        cell.hyperlink = hyperlink
        cell.font = Font(name='Arial', size=size, color='0563C1', underline='single')
    return cell


def _fila_negra(ws, row, num_cols):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row, col)
        cell.fill   = PatternFill('solid', fgColor=NEGRO)
        cell.border = _sin_borde


def _color_materia(materia):
    if not materia:
        return COLOR_DEFAULT
    return COLORES_MATERIA.get(materia.lower().strip(), COLOR_DEFAULT)


_num_grado = extraer_numero_grado


# ══════════════════════════════════════════════════════════════
# GENERADOR EXCEL — PROFESOR (sin cambios de formato)
# ══════════════════════════════════════════════════════════════

def _generar_excel_profesor(profesor, fecha_inicio, fecha_fin):
    """
    Genera el Excel de horario de un profesor: clases de colegios + clases particulares.

    Estructura: tabla transpuesta (filas=etiquetas, columnas=clases ordenadas por fecha/hora).
    Clases particulares se colorean en amarillo claro (PART_COLOR) para distinguirlas visualmente.

    N+1 resuelto: los libros de clases particulares se cargan en batch con
    NombreLibro.objects.filter(nombre__in=set_materiales) — una sola query para todos.

    El valor del campo 'titulo' para material asignado se fija a 'Material Asignado'
    y el nombre real del libro va en 'unidad', siguiendo la convención del sistema.
    """

    def _minutos(hora_str):
        try:
            partes = hora_str.split('-')[0].strip().split(':')
            return int(partes[0]) * 60 + int(partes[1])
        except Exception:
            return 9999

    entradas = []

    clases_colegio = list(
        Clase.objects
        .filter(profesor=profesor, fecha__gte=fecha_inicio, fecha__lte=fecha_fin,
                cancelada=False, es_evento=False)
        .select_related('bloque__colegio__colegio', 'bloque__grado', 'materia', 'libro_especial')
    )

    colegios_ids = list({c.bloque.colegio_id for c in clases_colegio})
    asig_map   = _build_asignaciones_map(colegios_ids)

    # Pre-cargar solo las unidades de libros relevantes (no toda la tabla)
    libros_ids = set()
    for asigs in asig_map.values():
        for a in asigs:
            libros_ids.add(a.libro_id)
    materiales_particulares = set(
        ClaseParticular.objects.filter(
            profesor=profesor, fecha__gte=fecha_inicio, fecha__lte=fecha_fin,
        ).values_list('material', flat=True)
    )
    materiales_particulares.discard(None)
    materiales_particulares.discard('')
    if materiales_particulares:
        libros_ids |= set(
            NombreLibro.objects.filter(nombre__in=materiales_particulares)
            .values_list('id', flat=True)
        )

    libros_map = {
        (u.libro.nombre, u.materia.nombre, str(u.numero)): u
        for u in Unidad.objects.filter(libro_id__in=libros_ids).select_related('libro', 'materia')
    } if libros_ids else {}

    for c in clases_colegio:
        colegio  = c.bloque.colegio
        grado    = c.bloque.grado.nombre
        if c.libro_especial:
            titulo = 'Material Asignado' #  c.libro_especial.nombre
        else:
            titulo = _titulo_libro(asig_map, colegio.id, grado, c.fecha)
        
        if c.unidad == 'S':
            unidad_c = 'Socialización de Simulacro'
        elif titulo == 'Material Asignado':
            unidad_c = c.libro_especial.nombre
        else:
            unidad_c = str(c.unidad) if c.unidad else ''
        material = 'Socialización de Simulacro' if unidad_c == 'Socialización de Simulacro' else titulo or 'Sin libro asignado'
        entradas.append({
            'fecha':          c.fecha,
            'hora':           c.bloque.hora,
            'minutos':        c.bloque.hora_inicio.hour * 60 + c.bloque.hora_inicio.minute,
            'colegio_nombre': colegio.nombre,
            'ciudad':         colegio.ciudad or '',
            'mapa_link':      colegio.mapa_link or '',
            'grado':          grado,
            'material':       material,
            'materia':        c.materia.nombre if c.materia_id else '',
            'unidad':         unidad_c,
            'tipo':           'colegio',
            'libro_obj':      libros_map.get((titulo, c.materia.nombre if c.materia_id else '', unidad_c)),
        })

    for p in ClaseParticular.objects.filter(
        profesor=profesor, fecha__gte=fecha_inicio, fecha__lte=fecha_fin
    ).select_related('materia', 'grado'):
        unidad    = str(p.unidad) if p.unidad else ''
        mat_nombre = p.materia.nombre if p.materia_id else ''
        libro_obj = libros_map.get((p.material, mat_nombre, unidad))
        material  = p.material or 'Sin libro asignado'
        entradas.append({
            'fecha':          p.fecha,
            'hora':           p.hora,
            'minutos':        p.hora_inicio.hour * 60 + p.hora_inicio.minute,
            'colegio_nombre': p.estudiante,
            'ciudad':         p.ciudad or '',
            'mapa_link':      p.mapa_link or '',
            'grado':          p.grado.nombre,
            'material':       material,
            'materia':        mat_nombre,
            'unidad':         unidad,
            'tipo':           'particular',
            'libro_obj':      libro_obj,
        })

    entradas.sort(key=lambda e: (e['fecha'], e['minutos']))

    wb = Workbook()
    ws = wb.active
    ws.title = profesor.nombre_corto[:31]

    if not entradas:
        _e(ws, 1, 1, f'{profesor.nombre} — Sin clases en el período seleccionado', bold=True)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # Fila 1: nombre + fechas
    _e(ws, 1, 1, profesor.nombre, fill=HEADER_DARK)

    for ci, e in enumerate(entradas, start=2):
        _e(ws, 1, ci, _fecha_label(e['fecha']), bold=True, fill=FECHA_GRAY)

    # Combinar celdas del mismo día
    ci_inicio = 2
    for i, e in enumerate(entradas):
        ci_actual = i + 2
        fecha_sig = entradas[i + 1]['fecha'] if i < len(entradas) - 1 else None
        if fecha_sig != e['fecha']:
            if ci_inicio < ci_actual:
                ws.merge_cells(start_row=1, start_column=ci_inicio,
                               end_row=1,   end_column=ci_actual)
                cell = ws.cell(1, ci_inicio, _fecha_label(e['fecha']))
                cell.font      = Font(name='Arial', size=10, bold=True)
                cell.fill      = PatternFill('solid', fgColor=FECHA_GRAY)
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            ci_inicio = ci_actual + 1

    labels = ['Colegio', 'Ciudad', 'Hora', 'Grado', 'Material', 'Asignatura', 'Unidad']
    for row_i, label in enumerate(labels, start=2):
        _e(ws, row_i, 1, label, bold=True, fill=LABEL_GRAY)

    for ci, e in enumerate(entradas, start=2):
        fill = PART_COLOR if e['tipo'] == 'particular' else None
        safe_link = _safe_url(e['mapa_link'])
        if safe_link:
            _e(ws, 2, ci, e['colegio_nombre'], fill=fill, hyperlink=safe_link)
        else:
            _e(ws, 2, ci, e['colegio_nombre'], fill=fill)

        _e(ws, 3, ci, e['ciudad'],   fill=fill)
        _e(ws, 4, ci, e['hora'],     fill=fill)
        _e(ws, 5, ci, e['grado'],    fill=fill)
        _e(ws, 6, ci, e['material'], fill=fill)
        _e(ws, 7, ci, e['materia'],  fill=fill)

        unidad = e['unidad']
        if unidad.isdigit() and e['libro_obj']:
            libro_obj  = e['libro_obj']
            nom_u      = libro_obj.nombre or ''
            unidad_txt = f'{unidad}. {nom_u}' if nom_u else unidad
            safe_libro_link = _safe_url(libro_obj.link)
            if safe_libro_link:
                _e(ws, 8, ci, unidad_txt, fill=fill, hyperlink=safe_libro_link)
            else:
                _e(ws, 8, ci, unidad_txt, fill=fill)
        else:
            _e(ws, 8, ci, unidad, fill=fill)

    for ci in range(1, 2 + len(entradas)):
        ws.column_dimensions[get_column_letter(ci)].width = COL_WIDTH

    # Bordes en todas las celdas con datos
    _thin = Side(style='thin', color='000000')
    _borde = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
    for row in range(1, 9):
        for col in range(1, 2 + len(entradas)):
            ws.cell(row, col).border = _borde

    # Altura fija — filas 1-8
    ws.row_dimensions[1].height = 22  # fechas
    for r in range(2, 9):
        ws.row_dimensions[r].height = 40  # datos

    ws.freeze_panes = 'B1'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════
# GENERADOR EXCEL — COLEGIO (formato nuevo)
# ══════════════════════════════════════════════════════════════

def _generar_excel_colegio(colegio, fecha_inicio, fecha_fin):
    """
    Estructura por fila:
    ┌────────────────────────────────────────────────────┐
    │ Fila 1 │ A1:A2 merged │ B1:B2 merged │ C1:C2 merged │ D1 vacío │ E1 vacío...
    │ Fila 2 │  (merged)    │  (merged)    │  (merged)    │ D2=fecha │ E2=fecha...
    │ Fila 3+│ A merged     │ B=grado      │ C=hora       │ profesor │ profesor...
    │        │  (todo el    │  (merged por │  (merged 2   │ contenido│ contenido
    │        │   colegio)   │   grado)     │   filas)     │          │
    └────────────────────────────────────────────────────┘
    Al final de cada grado: fila negra separadora.
    """

    # ── Datos ─────────────────────────────────────────────────
    clases_all = list(
        Clase.objects
        .filter(colegio=colegio, fecha__gte=fecha_inicio, fecha__lte=fecha_fin)
        .select_related('bloque__grado', 'profesor', 'materia', 'libro_especial')
        .order_by('fecha', 'bloque__hora_inicio')
    )
    fechas = sorted(set(c.fecha for c in clases_all))

    if not fechas:
        wb = Workbook()
        ws = wb.active
        ws.title = colegio.nombre[:31]
        _e(ws, 1, 1, f'{colegio.nombre} — Sin clases en el período seleccionado', bold=True)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # Matriz: bloque_id → fecha → clase
    matriz = defaultdict(dict)
    for c in clases_all:
        matriz[c.bloque_id][c.fecha] = c

    # Bloques agrupados por grado
    bloques_all = list(Bloque.objects.filter(colegio=colegio).select_related('grado').order_by('grado__nombre', 'hora_inicio'))
    grados_dict = defaultdict(list)
    for b in bloques_all:
        grados_dict[b.grado.nombre].append(b)
    grados_ordenados = ordenar_grados(grados_dict.keys())

    asig_map = _build_asignaciones_map([colegio.id])

    # Pre-cargar solo las unidades de libros asignados al colegio
    libros_ids = {a.libro_id for asigs in asig_map.values() for a in asigs}
    libros_map = {
        (u.libro.nombre, u.materia.nombre, str(u.numero)): u
        for u in Unidad.objects.filter(libro_id__in=libros_ids).select_related('libro', 'materia')
    } if libros_ids else {}

    # ── Workbook ──────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = colegio.nombre[:31]

    num_cols = 3 + len(fechas)

    # ── Filas 1-2: cabecera ───────────────────────────────────
    # A1:A2 merged
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    _e(ws, 1, 1, '🏫', bold=True, fill=LABEL_GRAY, size=14)

    # B1:B2 merged
    ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
    _e(ws, 1, 2, 'GRADO - SALÓN', bold=True, fill=LABEL_GRAY)

    # C1:C2 merged
    ws.merge_cells(start_row=1, start_column=3, end_row=2, end_column=3)
    _e(ws, 1, 3, '⏱️', bold=True, fill=LABEL_GRAY, size=14)

    # Columnas de fechas: fila 1 NEGRO (vacía), fila 2 con fecha
    for ci, fecha in enumerate(fechas, start=4):
        _e(ws, 1, ci, '', fill=NEGRO)
        _e(ws, 2, ci, _fecha_label(fecha), bold=True, fill=FECHA_GRAY)

    # ── Filas de datos (fila 3 en adelante) ───────────────────
    current_row  = 3
    first_data_row = 3

    for grado in grados_ordenados:
        bloques_grado  = sorted(grados_dict[grado], key=lambda x: x.hora_inicio)
        start_row_grado = current_row

        for bloque in bloques_grado:
            row_prof = current_row      # fila del profesor
            row_cont = current_row + 1  # fila del contenido

            # Col A: estilos (se mergea al final)
            _e(ws, row_prof, 1, '', fill=LABEL_GRAY)
            _e(ws, row_cont, 1, '', fill=LABEL_GRAY)

            # Col B: estilos (se mergea al final de grado)
            _e(ws, row_prof, 2, '', fill=LABEL_GRAY)
            _e(ws, row_cont, 2, '', fill=LABEL_GRAY)

            # Col C: hora — merge 2 filas
            ws.merge_cells(start_row=row_prof, start_column=3,
                           end_row=row_cont,   end_column=3)
            _e(ws, row_prof, 3, bloque.hora, bold=True, fill=LABEL_GRAY)

            # Columnas de fechas
            for ci, fecha in enumerate(fechas, start=4):
                clase = matriz[bloque.id].get(fecha)

                if clase is None:
                    # Sin clase → negro
                    _e(ws, row_prof, ci, '', fill=NEGRO)
                    _e(ws, row_cont, ci, '', fill=NEGRO)

                elif clase.cancelada:
                    ws.merge_cells(start_row=row_prof, start_column=ci,
                                   end_row=row_cont,   end_column=ci)
                    _e(ws, row_prof, ci, 'CANCELADA', bold=True, fill=COLOR_CANCELADA)

                elif clase.es_evento:
                    ws.merge_cells(start_row=row_prof, start_column=ci,
                                   end_row=row_cont,   end_column=ci)
                    _e(ws, row_prof, ci, clase.titulo_evento or 'Evento',
                       bold=True, fill=COLOR_EVENTO)

                else:
                    materia  = clase.materia.nombre if clase.materia_id else ''
                    color    = _color_materia(materia)
                    unidad   = str(clase.unidad) if clase.unidad else ''
                    nombre_profe = clase.profesor.nombre_corto if clase.profesor else ''

                    # Fila profesor — sin negrita, sin color de fondo
                    _e(ws, row_prof, ci, nombre_profe)

                    # Fila contenido — con color de materia
                    if unidad == 'S':
                        texto_s = (f'{materia} | Socialización de Simulacro')
                        _e(ws, row_cont, ci, texto_s, fill=color)
                    elif unidad.isdigit():
                        if clase.libro_especial:
                            titulo = clase.libro_especial.nombre
                        else:
                            titulo = _titulo_libro(asig_map, colegio.id, grado, fecha)
                        libro_obj = libros_map.get((titulo, materia, unidad))
                        nom_u     = libro_obj.nombre if libro_obj else ''
                        unidad_txt = f'U.{unidad}: {nom_u}' if nom_u else f'U.{unidad}'
                        texto = (f'{titulo} | {materia} | {unidad_txt}'
                                 if titulo else f'{materia} | {unidad_txt}')
                        _e(ws, row_cont, ci, texto, fill=color)
                    else:
                        _e(ws, row_cont, ci, unidad or '', fill=color)

            current_row += 2  # avanzar 2 filas (prof + contenido)
            # Altura fija para que el texto se muestre correctamente
            ws.row_dimensions[current_row - 2].height = 15  # fila profesor
            ws.row_dimensions[current_row - 1].height = 40  # fila contenido

        end_row_grado = current_row - 1

        # Merge B para el grado completo
        ws.merge_cells(start_row=start_row_grado, start_column=2,
                       end_row=end_row_grado,     end_column=2)
        _e(ws, start_row_grado, 2, grado, bold=True, fill=LABEL_GRAY)

        # Fila separadora negra
        _fila_negra(ws, current_row, num_cols)
        current_row += 1

    last_data_row = current_row - 1

    # Merge A para todo el bloque de datos (nombre colegio rotado)
    if last_data_row >= first_data_row:
        ws.merge_cells(start_row=first_data_row, start_column=1,
                       end_row=last_data_row,    end_column=1)
        cell_a           = ws.cell(first_data_row, 1)
        cell_a.value     = colegio.nombre
        cell_a.font      = Font(name='Arial', size=11, bold=True)
        cell_a.alignment = Alignment(horizontal='center', vertical='center',
                                      wrap_text=True, text_rotation=90)
        cell_a.fill      = PatternFill('solid', fgColor=LABEL_GRAY)

    # ── Tabla resumen docentes ────────────────────────────────
    current_row += 1  # fila en blanco

    clases_normales = [c for c in clases_all
                       if not c.cancelada and not c.es_evento and c.materia_id and c.profesor]

    conteo = defaultdict(lambda: defaultdict(int))
    ids_profes = set()
    for c in clases_normales:
        conteo[(c.profesor_id, c.materia.nombre)][c.bloque.grado.nombre] += 1
        ids_profes.add(c.profesor_id)

    if ids_profes:
        tabla_inicio_row = current_row
        # Cabecera resumen
        _e(ws, current_row, 1, 'DOCENTES',  bold=True, fill='FCE4D6', border=True)
        _e(ws, current_row, 2, 'DOCUMENTO', bold=True, fill='FCE4D6', border=True)
        _e(ws, current_row, 3, 'MATERIA',   bold=True, fill='FCE4D6', border=True)
        for ci_g, grado in enumerate(grados_ordenados, start=4):
            _e(ws, current_row, ci_g, f'Clases {grado}', bold=True, fill='FCE4D6', border=True)
        current_row += 1

        profes_map = {p.id: p for p in Profesor.objects.filter(id__in=ids_profes)}

        # Agrupar materias por profesor manteniendo orden
        profe_materias = defaultdict(list)
        for prof_id, materia in sorted(
            conteo.keys(),
            key=lambda k: (profes_map[k[0]].nombre_corto, k[1])
        ):
            profe_materias[prof_id].append(materia)

        for prof_id in sorted(profe_materias.keys(),
                               key=lambda pid: profes_map[pid].nombre_corto):
            materias         = profe_materias[prof_id]
            profe            = profes_map[prof_id]
            start_row_profe  = current_row

            for materia in materias:
                _e(ws, current_row, 3, materia, border=True)
                for ci_g, grado in enumerate(grados_ordenados, start=4):
                    cnt = conteo[(prof_id, materia)].get(grado, 0)
                    _e(ws, current_row, ci_g, cnt if cnt else '', border=True)
                current_row += 1

            end_row_profe = current_row - 1

            # Merge nombre y documento si tiene más de 1 materia
            if end_row_profe > start_row_profe:
                ws.merge_cells(start_row=start_row_profe, start_column=1,
                               end_row=end_row_profe,     end_column=1)
                ws.merge_cells(start_row=start_row_profe, start_column=2,
                               end_row=end_row_profe,     end_column=2)
                # Bordes en todas las celdas del rango — todos los lados
                _thin_side = Side(style='thin', color='000000')
                _borde_full = Border(left=_thin_side, right=_thin_side,
                                     top=_thin_side, bottom=_thin_side)
                for r in range(start_row_profe, end_row_profe + 1):
                    for col in [1, 2]:
                        ws.cell(r, col).border = _borde_full

            _e(ws, start_row_profe, 1, profe.nombre_corto, v='center', border=True)
            _e(ws, start_row_profe, 2, profe.documento or '', v='center', border=True)

    # ── Ancho de columnas ────────────────────────────────────
    # Columnas A, B, C: autoajuste según contenido
    for ci, col_letter in enumerate(['A', 'B', 'C'], start=1):
        max_len = 0
        for row in ws.iter_rows(min_col=ci, max_col=ci):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 40)
    # Columnas de fechas: ancho fijo
    for ci in range(4, num_cols + 1):
        ws.column_dimensions[get_column_letter(ci)].width = COL_WIDTH

    # Altura fija de filas de cabecera
    ws.row_dimensions[1].height = 15
    ws.row_dimensions[2].height = 22

    ws.freeze_panes = 'D3'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parsear_fechas(fecha_inicio_str, fecha_fin_str):
    """Parsea strings de fecha o devuelve defaults del año actual."""
    anio = date.today().year
    try:
        fi = datetime.strptime(fecha_inicio_str, '%Y-%m-%d').date() if fecha_inicio_str else date(anio, 1, 1)
    except ValueError:
        fi = date(anio, 1, 1)
    try:
        ff = datetime.strptime(fecha_fin_str, '%Y-%m-%d').date() if fecha_fin_str else date(anio, 12, 31)
    except ValueError:
        ff = date(anio, 12, 31)
    return fi, ff


# ══════════════════════════════════════════════════════════════
# VISTAS
# ══════════════════════════════════════════════════════════════

@user_passes_test(es_personal_programacion, login_url='login')
def exportar_view(request):
    """
    Exportación masiva de horarios en ZIP.

    GET: muestra la página de filtros (profesores, colegios, rango de fechas).
    POST: genera el ZIP en memoria (BytesIO) y lo devuelve como attachment.

    El ZIP se construye completamente en memoria sin archivos temporales,
    requerimiento crítico para Render cuyo filesystem es efímero.
    Estructura interna: Profesores/Horario {nombre}.xlsx y Colegios/Horario {nombre}.xlsx.
    """

    if request.method == 'GET':
        anio = date.today().year
        profesores = Profesor.objects.filter(activo=True).order_by('nombre')
        colegios = (ColegioAnio.objects
                    .filter(activo=True)
                    .select_related('colegio')
                    .order_by('colegio__nombre'))
        return render(request, 'exportar/exportar.html', {
            'profesores':   profesores,
            'colegios':     colegios,
            'fecha_inicio': date(anio, 1, 1).isoformat(),
            'fecha_fin':    date(anio, 12, 31).isoformat(),
        })

    # ── POST: generar ZIP ──────────────────────────────────────
    tipo = request.POST.get('tipo', 'ambos')
    fecha_inicio, fecha_fin = _parsear_fechas(
        request.POST.get('fecha_inicio', ''),
        request.POST.get('fecha_fin', ''),
    )
    profesores_ids = [int(x) for x in request.POST.getlist('profesores_ids') if x.isdigit()]
    colegios_ids   = [int(x) for x in request.POST.getlist('colegios_ids') if x.isdigit()]

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zf:

        if tipo in ('profesores', 'ambos'):
            q_cl = Clase.objects.filter(
                fecha__gte=fecha_inicio, fecha__lte=fecha_fin,
                cancelada=False, es_evento=False,
            )
            q_pa = ClaseParticular.objects.filter(
                fecha__gte=fecha_inicio, fecha__lte=fecha_fin,
            )
            if profesores_ids:
                q_cl = q_cl.filter(profesor_id__in=profesores_ids)
                q_pa = q_pa.filter(profesor_id__in=profesores_ids)

            ids_prof = (
                set(q_cl.values_list('profesor_id', flat=True))
                | set(q_pa.values_list('profesor_id', flat=True))
            )
            ids_prof.discard(None)

            for profesor in Profesor.objects.filter(id__in=ids_prof).order_by('nombre'):
                excel_bytes = _generar_excel_profesor(profesor, fecha_inicio, fecha_fin)
                zf.writestr(f"Profesores/Horario {profesor.nombre_corto}.xlsx", excel_bytes)

        if tipo in ('colegios', 'ambos'):
            q_col = Clase.objects.filter(
                fecha__gte=fecha_inicio, fecha__lte=fecha_fin,
            )
            if colegios_ids:
                q_col = q_col.filter(colegio_id__in=colegios_ids)

            ids_col = set(q_col.values_list('colegio_id', flat=True).distinct())

            for colegio in (ColegioAnio.objects
                            .filter(id__in=ids_col)
                            .select_related('colegio')
                            .order_by('colegio__nombre')):
                excel_bytes = _generar_excel_colegio(colegio, fecha_inicio, fecha_fin)
                zf.writestr(f"Colegios/Horario {colegio.nombre}.xlsx", excel_bytes)

    zip_buffer.seek(0)
    label = f"{fecha_inicio.strftime('%Y%m%d')}_{fecha_fin.strftime('%Y%m%d')}"
    response = HttpResponse(zip_buffer, content_type='application/zip')
    response['Content-Disposition'] = f'attachment; filename="Exportacion_{label}.zip"'
    return response


@user_passes_test(es_personal_programacion, login_url='login')
def exportar_contar(request):
    """AJAX: cuenta clases según filtros activos (solo COUNT queries, sin generar Excels)."""
    tipo = request.GET.get('tipo', 'ambos')
    fecha_inicio, fecha_fin = _parsear_fechas(
        request.GET.get('fecha_inicio', ''),
        request.GET.get('fecha_fin', ''),
    )
    profesores_ids = [int(x) for x in request.GET.getlist('profesores_ids') if x.isdigit()]
    colegios_ids   = [int(x) for x in request.GET.getlist('colegios_ids') if x.isdigit()]

    resultado = {
        'n_profesores': 0, 'clases_colegio': 0, 'clases_particular': 0,
        'n_colegios': 0, 'clases_colegios': 0,
        'total': 0,
    }

    if tipo in ('profesores', 'ambos'):
        q_cl = Clase.objects.filter(
            fecha__gte=fecha_inicio, fecha__lte=fecha_fin,
            cancelada=False, es_evento=False,
        )
        q_pa = ClaseParticular.objects.filter(
            fecha__gte=fecha_inicio, fecha__lte=fecha_fin,
        )
        if profesores_ids:
            q_cl = q_cl.filter(profesor_id__in=profesores_ids)
            q_pa = q_pa.filter(profesor_id__in=profesores_ids)

        resultado['clases_colegio']    = q_cl.count()
        resultado['clases_particular'] = q_pa.count()

        ids_prof = (
            set(q_cl.values_list('profesor_id', flat=True))
            | set(q_pa.values_list('profesor_id', flat=True))
        )
        ids_prof.discard(None)
        resultado['n_profesores'] = len(ids_prof)

    if tipo in ('colegios', 'ambos'):
        q_col = Clase.objects.filter(
            fecha__gte=fecha_inicio, fecha__lte=fecha_fin,
        )
        if colegios_ids:
            q_col = q_col.filter(colegio_id__in=colegios_ids)

        resultado['clases_colegios'] = q_col.count()
        resultado['n_colegios'] = q_col.values('colegio_id').distinct().count()

    resultado['total'] = (resultado['clases_colegio']
                          + resultado['clases_particular']
                          + resultado['clases_colegios'])
    return JsonResponse(resultado)


# ══════════════════════════════════════════════════════════════
# EXPORTAR PAGOS
# ══════════════════════════════════════════════════════════════

MESES_ES_LARGO = ['', 'enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
                  'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']


def _semana_label(fecha_inicio, fecha_fin):
    if fecha_inicio.month == fecha_fin.month:
        return (f"Semana del {fecha_inicio.day} al {fecha_fin.day} "
                f"de {MESES_ES_LARGO[fecha_inicio.month]} de {fecha_inicio.year}")
    return (f"Semana del {fecha_inicio.day} de {MESES_ES_LARGO[fecha_inicio.month]} "
            f"al {fecha_fin.day} de {MESES_ES_LARGO[fecha_fin.month]} de {fecha_fin.year}")


def _build_filas_pagos(fecha_inicio, fecha_fin):
    """
    Devuelve lista de dicts con los datos de pago agrupados por
    (fecha, profesor, colegio), sumando horas del día.
    """
    from collections import defaultdict

    clases = list(
        Clase.objects
        .filter(
            fecha__gte=fecha_inicio, fecha__lte=fecha_fin,
            cancelada=False, es_evento=False,
        )
        .select_related(
            'profesor',
            'bloque__colegio__colegio',
        )
        .values(
            'fecha',
            'profesor_id',
            'profesor__nombre',
            'profesor__apellido',
            'profesor__documento',
            'profesor__cuenta_bancaria',
            'profesor__tipo_cuenta',
            'profesor__banco',
            'bloque__colegio_id',
            'bloque__colegio__colegio__nombre',
            'bloque__colegio__colegio__codigo',
            'bloque__colegio__valor_hora',
            'bloque__hora_inicio',
            'bloque__hora_fin',
        )
    )

    # Agrupar: (fecha, profesor_id, colegio_id) → minutos + info
    grupos = {}
    for c in clases:
        if not c['profesor_id']:
            continue
        key = (c['fecha'], c['profesor_id'], c['bloque__colegio_id'])
        hi = c['bloque__hora_inicio']
        hf = c['bloque__hora_fin']
        minutos = 0
        if hi and hf:
            minutos = max(0, (hf.hour * 60 + hf.minute) - (hi.hour * 60 + hi.minute))
        if key not in grupos:
            grupos[key] = {'minutos': 0, 'info': c}
        grupos[key]['minutos'] += minutos

    filas = []
    for (fecha, prof_id, col_id), data in sorted(grupos.items(), key=lambda x: (x[0][0], x[0][1])):
        info = data['info']
        horas = data['minutos'] / 60
        valor_hora = info['bloque__colegio__valor_hora'] or 0
        valor_total = round(horas * valor_hora)

        nombre   = info['profesor__nombre'] or ''
        apellido = info['profesor__apellido'] or ''
        pn = nombre.split()[0] if nombre else ''
        pa = apellido.split()[0] if apellido else ''
        nombre_corto = f"{pn} {pa}".strip()

        # Tipo de cuenta: "Ahorros a la mano" si Daviplata, else valor normal
        banco      = info['profesor__banco'] or ''
        tipo_raw   = info['profesor__tipo_cuenta'] or ''
        tipo_cuenta = 'Ahorros a la mano' if banco == 'Daviplata' else tipo_raw

        filas.append({
            'fecha':       fecha,
            'profesor_id': prof_id,
            'colegio_id':  col_id,
            'docente':     nombre_corto,
            'documento':   info['profesor__documento'] or '',
            'num_cuenta':  info['profesor__cuenta_bancaria'] or '',
            'tipo_cuenta': tipo_cuenta,
            'banco':       banco,
            'colegio':     info['bloque__colegio__colegio__nombre'] or '',
            'codigo':      info['bloque__colegio__colegio__codigo'] or '',
            'horas':       horas,
            'valor_hora':  valor_hora,
            'valor_total': valor_total,
        })

    return filas


def _generar_excel_pagos(filas, semana_label):
    """Genera el Excel de pagos con el formato de la imagen."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Pagos'

    NUM_COLS = 9
    COLS = ['FECHA', 'DOCENTE', 'DOCUMENTO', 'N° DE CUENTA',
            'TIPO DE CUENTA', 'BANCO', 'COLEGIO', 'CODIGO', 'VALOR']

    HEADER_FILL  = 'FF1F3864'  # azul oscuro
    SUBHDR_FILL  = 'FFD6E4F0'  # azul claro
    TOTAL_FILL   = 'FFD6DCE4'

    thin = Side(style='thin', color='000000')
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _celda(row, col, valor='', bold=False, fill=None, color='000000',
               h='center', v='center', fmt=None, borde_=True):
        cell = ws.cell(row, col, valor)
        cell.font = Font(name='Arial', size=10, bold=bold, color=color)
        cell.alignment = Alignment(horizontal=h, vertical=v, wrap_text=True)
        if fill:
            cell.fill = PatternFill('solid', fgColor=fill)
        if borde_:
            cell.border = borde
        if fmt:
            cell.number_format = fmt
        return cell

    # Fila 1: título semana
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NUM_COLS)
    _celda(1, 1, semana_label, bold=True, fill='FFD9E1F2', color='FF1F3864',
           borde_=False)
    ws.row_dimensions[1].height = 20

    # Fila 2: cabeceras
    for ci, col_name in enumerate(COLS, start=1):
        _celda(2, ci, col_name, bold=True, fill='FFB8CCE4', color='FF1F3864')
    ws.row_dimensions[2].height = 22

    # Filas de datos
    MESES_ES_ABREV = ['', 'Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun',
                      'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']

    for i, f in enumerate(filas, start=3):
        fill_row = 'FFFFFFFF' if i % 2 == 1 else 'FFF2F6FC'
        fecha_str = f"{f['fecha'].day:02d}/{f['fecha'].month:02d}/{f['fecha'].year}"
        _celda(i, 1, fecha_str,       fill=fill_row, h='center')
        _celda(i, 2, f['docente'],    fill=fill_row, h='left')
        _celda(i, 3, f['documento'],  fill=fill_row, h='center')
        _celda(i, 4, f['num_cuenta'], fill=fill_row, h='center')
        _celda(i, 5, f['tipo_cuenta'],fill=fill_row, h='center')
        _celda(i, 6, f['banco'],      fill=fill_row, h='center')
        _celda(i, 7, f['colegio'],    fill=fill_row, h='left')
        _celda(i, 8, f['codigo'],     fill=fill_row, h='center')
        # Valor como número para que Excel pueda sumar
        cell_val = ws.cell(i, 9, f['valor_total'])
        cell_val.font = Font(name='Arial', size=10)
        cell_val.alignment = Alignment(horizontal='right', vertical='center')
        cell_val.fill = PatternFill('solid', fgColor=fill_row)
        cell_val.border = borde
        cell_val.number_format = '"$"#,##0'
        ws.row_dimensions[i].height = 18

    # Fila TOTAL
    total_row = 3 + len(filas)
    ws.merge_cells(start_row=total_row, start_column=1,
                   end_row=total_row, end_column=8)
    _celda(total_row, 1, 'TOTAL', bold=True, fill=TOTAL_FILL, h='right')
    total_val = sum(f['valor_total'] for f in filas)
    cell_t = ws.cell(total_row, 9, total_val)
    cell_t.font = Font(name='Arial', size=10, bold=True)
    cell_t.alignment = Alignment(horizontal='right', vertical='center')
    cell_t.fill = PatternFill('solid', fgColor=TOTAL_FILL)
    cell_t.border = borde
    cell_t.number_format = '"$"#,##0'
    ws.row_dimensions[total_row].height = 22

    # Anchos de columna
    anchos = [12, 22, 14, 18, 18, 16, 30, 10, 14]
    for ci, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = ancho

    ws.freeze_panes = 'A3'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@user_passes_test(es_personal_programacion, login_url='login')
def exportar_pagos_view(request):
    """GET: página de pagos (tabs pendiente/realizado). POST: descarga Excel."""
    from datetime import timedelta

    hoy = date.today()
    lunes   = hoy - timedelta(days=hoy.weekday())
    viernes = lunes + timedelta(days=4)

    if request.method == 'GET':
        ayer = hoy - timedelta(days=1)
        semana_str = request.GET.get('semana', '')
        hasta_str  = request.GET.get('hasta', '')
        if semana_str:
            try:
                lunes = datetime.strptime(semana_str, '%Y-%m-%d').date()
                viernes = (datetime.strptime(hasta_str, '%Y-%m-%d').date()
                           if hasta_str else lunes + timedelta(days=4))
            except ValueError:
                pass
        # Clamp: no permitir fechas futuras
        viernes = min(viernes, ayer)
        lunes   = min(lunes, ayer)

        tab = request.GET.get('tab', 'pendiente')

        # Claves ya pagadas en el rango
        pagados_keys = set(
            PagoRealizado.objects
            .filter(fecha__gte=lunes, fecha__lte=viernes)
            .values_list('profesor_id', 'colegio_id', 'fecha')
        )

        todas_filas = _build_filas_pagos(lunes, viernes)

        filas_pendientes = []
        filas_realizadas = []
        for f in todas_filas:
            key = (f['profesor_id'], f['colegio_id'], f['fecha'])
            if key in pagados_keys:
                filas_realizadas.append(f)
            else:
                filas_pendientes.append(f)

        # Pagos realizados enriquecidos con fecha_pago y marcado_por
        pagos_db = {
            (p.profesor_id, p.colegio_id, p.fecha): p
            for p in PagoRealizado.objects
                .filter(fecha__gte=lunes, fecha__lte=viernes)
                .select_related('marcado_por')
        }
        for f in filas_realizadas:
            pago = pagos_db.get((f['profesor_id'], f['colegio_id'], f['fecha']))
            f['fecha_pago']  = pago.fecha_pago  if pago else None
            f['marcado_por'] = pago.marcado_por.get_full_name() or pago.marcado_por.username if pago and pago.marcado_por else '—'
            f['pago_id']     = pago.id if pago else None

        filas_tab = filas_pendientes if tab == 'pendiente' else filas_realizadas

        return render(request, 'exportar/pagos.html', {
            'fecha_inicio':      lunes.isoformat(),
            'fecha_fin':         viernes.isoformat(),
            'fecha_max':         ayer.isoformat(),
            'semana_label':      _semana_label(lunes, viernes),
            'tab':               tab,
            'filas_pendientes':  filas_pendientes,
            'filas_realizadas':  filas_realizadas,
            'filas':             filas_tab,
            'total_valor':       sum(f['valor_total'] for f in filas_tab),
            'total_pendiente':   sum(f['valor_total'] for f in filas_pendientes),
            'total_realizado':   sum(f['valor_total'] for f in filas_realizadas),
        })

    # POST: descarga Excel del tab activo
    fi_str = request.POST.get('fecha_inicio', '')
    ff_str = request.POST.get('fecha_fin', '')
    tab    = request.POST.get('tab', 'pendiente')
    try:
        fi = datetime.strptime(fi_str, '%Y-%m-%d').date()
        ff = datetime.strptime(ff_str, '%Y-%m-%d').date()
    except ValueError:
        fi, ff = lunes, viernes

    pagados_keys = set(
        PagoRealizado.objects
        .filter(fecha__gte=fi, fecha__lte=ff)
        .values_list('profesor_id', 'colegio_id', 'fecha')
    )
    todas = _build_filas_pagos(fi, ff)
    if tab == 'realizado':
        filas = [f for f in todas
                 if (f['profesor_id'], f['colegio_id'], f['fecha']) in pagados_keys]
    else:
        filas = [f for f in todas
                 if (f['profesor_id'], f['colegio_id'], f['fecha']) not in pagados_keys]

    semana_label = _semana_label(fi, ff)
    excel_bytes  = _generar_excel_pagos(filas, semana_label)

    sufijo = 'Realizados' if tab == 'realizado' else 'Pendientes'
    label  = f"{fi.strftime('%Y%m%d')}_{ff.strftime('%Y%m%d')}"
    response = HttpResponse(
        excel_bytes,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = (
        f'attachment; filename="Pagos_{sufijo}_{label}.xlsx"'
    )
    return response


@user_passes_test(es_personal_programacion, login_url='login')
def ajax_marcar_pago(request):
    """
    Marca o desmarca una fila (profesor, colegio, fecha) como pago realizado.

    get_or_create respeta el constraint único (profesor, colegio, fecha) sin lanzar
    IntegrityError si el mismo pago se marca dos veces (idempotente).
    La acción 'desmarcar' elimina el registro para revertir el marcado.
    """
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'Método no permitido'}, status=405)

    accion      = request.POST.get('accion', 'marcar')  # 'marcar' | 'desmarcar'
    profesor_id = request.POST.get('profesor_id', '')
    colegio_id  = request.POST.get('colegio_id', '')
    fecha_str   = request.POST.get('fecha', '')

    try:
        profesor_id = int(profesor_id)
        colegio_id  = int(colegio_id)
        fecha_obj   = datetime.strptime(fecha_str, '%Y-%m-%d').date()
    except (ValueError, TypeError):
        return JsonResponse({'ok': False, 'error': 'Datos inválidos'}, status=400)

    if accion == 'desmarcar':
        deleted, _ = PagoRealizado.objects.filter(
            profesor_id=profesor_id,
            colegio_id=colegio_id,
            fecha=fecha_obj,
        ).delete()
        return JsonResponse({'ok': True, 'accion': 'desmarcado', 'deleted': deleted})

    # marcar
    try:
        horas = float(request.POST.get('horas', 0))
        valor = int(request.POST.get('valor', 0))
    except (ValueError, TypeError):
        return JsonResponse({'ok': False, 'error': 'Valor/horas inválidos'}, status=400)

    pago, created = PagoRealizado.objects.get_or_create(
        profesor_id=profesor_id,
        colegio_id=colegio_id,
        fecha=fecha_obj,
        defaults={
            'horas':       horas,
            'valor':       valor,
            'marcado_por': request.user,
        },
    )
    return JsonResponse({'ok': True, 'accion': 'marcado', 'created': created, 'id': pago.id})