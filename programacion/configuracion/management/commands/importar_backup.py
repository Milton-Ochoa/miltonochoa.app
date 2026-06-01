"""Importa el backup completo de la base de datos desde el Excel de exportación.

El proyecto solo trae exportadores; este comando es el camino inverso: lee el
`.xlsx` con las 17 hojas que reflejan las tablas del sistema y las vuelca en la
base de datos. Está pensado para arrancar una instancia limpia (BD vacía +
`migrate`) con los datos reales de producción.

Uso:
    python manage.py importar_backup backups/AAMO_export.xlsx

Decisiones de diseño:
  - **PKs preservadas:** cada fila se inserta con su `id` original. Como todas
    las FKs del Excel vienen como `<campo>_id` (enteros), preservar los ids hace
    que las relaciones encajen sin necesidad de resolver por nombre.
  - **Idempotente:** usa `update_or_create(id=...)`, así que re-ejecutarlo
    actualiza en vez de duplicar.
  - **Orden de dependencias:** las hojas se procesan de catálogos a tablas que
    los referencian (ver SHEET_SPECS).
  - **Campos auto_now/auto_now_add:** Django los pisa con la hora actual al
    guardar; tras crear cada fila se reescriben con el valor real del backup vía
    `QuerySet.update()` (que no dispara `pre_save`).
  - **Transacción atómica:** cualquier fallo revierte toda la importación.

Limitaciones conocidas (el backup no las trae, se reportan al terminar):
  - Las contraseñas de los usuarios no están en el Excel → se importan como
    inutilizables (`set_unusable_password`); hay que resetearlas o usar el
    superusuario de prueba creado aparte.
  - Las M2M (`Profesor.materias`, `AlertaAuditoria.colegios_implicados`) no
    tienen hoja propia y no se reconstruyen.
"""

from datetime import date, time

import openpyxl
from django.apps import apps
from django.contrib.auth.hashers import make_password
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.models import CharField, TextField
from django.utils.dateparse import parse_datetime


# (hoja, 'app_label.Modelo', {columna: tipo_a_parsear})
# El orden importa: primero los catálogos sin dependencias, luego lo que
# referencia a otros por FK.
SHEET_SPECS = [
    ('Materias',           'configuracion.Materia',     {}),
    ('Grados',             'colegios.Grado',            {}),
    ('NombreLibro',        'configuracion.NombreLibro', {}),
    ('Colegios',           'configuracion.Colegio',     {}),
    ('Profesores',         'configuracion.Profesor',    {'fecha_nacimiento': 'date'}),
    ('Usuarios',           'auth.User',                 {'date_joined': 'datetime'}),
    ('ColegioAnio',        'configuracion.ColegioAnio', {}),
    ('Unidades',           'configuracion.Unidad',      {}),
    ('Bloques',            'colegios.Bloque',           {'hora_inicio': 'time', 'hora_fin': 'time'}),
    ('Asignaciones',       'colegios.Asignacion',       {'fecha_inicio': 'date', 'fecha_fin': 'date'}),
    ('ClasesParticulares', 'colegios.ClaseParticular',  {'fecha': 'date', 'hora_inicio': 'time', 'hora_fin': 'time'}),
    ('Informes',           'informes.Informe',          {'fecha': 'date', 'creado_en': 'datetime', 'actualizado_en': 'datetime'}),
    ('Tareas',             'pendientes.Tarea',          {'fecha_creacion': 'datetime', 'fecha_completado': 'datetime'}),
    ('AlertasAuditoria',   'auditoria.AlertaAuditoria', {'detectado': 'datetime', 'resuelto_en': 'datetime'}),
    ('HistorialCambios',   'colegios.HistorialCambio',  {'fecha': 'datetime'}),
    ('UsuarioColegio',     'usuarios.UsuarioColegio',   {}),
    ('UsuarioProfesor',    'usuarios.UsuarioProfesor',  {}),
]

# Campos gestionados por Django (auto_now_add / auto_now) que hay que reescribir
# después de crear la fila para conservar la marca de tiempo original del backup.
AUTO_FIELDS = {
    'pendientes.Tarea':         ['fecha_creacion'],
    'auditoria.AlertaAuditoria': ['detectado'],
    'colegios.HistorialCambio': ['fecha'],
    'informes.Informe':         ['creado_en', 'actualizado_en'],
}


def _parse(kind, value):
    """Convierte el valor crudo del Excel (siempre str/num/bool/None) al tipo Python."""
    if value is None or value == '':
        return None
    if kind == 'date':
        return date.fromisoformat(str(value)[:10])
    if kind == 'time':
        return time.fromisoformat(str(value))
    if kind == 'datetime':
        return parse_datetime(str(value))
    return value


def _coerce_none(model, field_name, value):
    """None en un char/text no-nullable rompe el INSERT; lo cambia por ''."""
    if value is not None or field_name.endswith('_id'):
        return value
    try:
        field = model._meta.get_field(field_name)
    except Exception:
        return value
    if not field.null and isinstance(field, (CharField, TextField)):
        return ''
    return value


def _iter_rows(ws):
    """Genera un dict {columna: valor} por fila de datos (salta el encabezado)."""
    it = ws.iter_rows(values_only=True)
    try:
        header = list(next(it))
    except StopIteration:
        return
    for row in it:
        yield dict(zip(header, row))


class Command(BaseCommand):
    help = 'Importa el backup completo (AAMO_export.xlsx) a la base de datos.'

    def add_arguments(self, parser):
        parser.add_argument('archivo', help='Ruta al .xlsx de backup (ej: backups/AAMO_export.xlsx)')

    def handle(self, *args, **options):
        ruta = options['archivo']
        try:
            wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        except FileNotFoundError:
            raise CommandError(f'No se encontró el archivo: {ruta}')

        faltantes = [s for s, _, _ in SHEET_SPECS if s not in wb.sheetnames]
        if faltantes:
            raise CommandError(f'Al Excel le faltan hojas esperadas: {", ".join(faltantes)}')

        resumen = []
        with transaction.atomic():
            for sheet, model_path, types in SHEET_SPECS:
                model = apps.get_model(*model_path.split('.'))
                es_user = model_path == 'auth.User'
                auto = AUTO_FIELDS.get(model_path, [])
                creados = actualizados = 0

                for row in _iter_rows(wb[sheet]):
                    pk = row.pop('id')
                    defaults = {}
                    for col, raw in row.items():
                        val = _parse(types.get(col), raw)
                        defaults[col] = _coerce_none(model, col, val)

                    if es_user:
                        # El backup no incluye contraseñas: se dejan inutilizables.
                        defaults['password'] = make_password(None)

                    obj, created = model.objects.update_or_create(id=pk, defaults=defaults)
                    creados += created
                    actualizados += not created

                    # Reescribir las marcas de tiempo que Django acaba de pisar.
                    if auto:
                        model.objects.filter(pk=obj.pk).update(
                            **{f: defaults[f] for f in auto if defaults.get(f) is not None}
                        )

                total = model.objects.count()
                resumen.append((sheet, creados, actualizados, total))
                self.stdout.write(
                    f'  {sheet:<20} creados={creados:<5} actualizados={actualizados:<5} total_bd={total}'
                )

        self.stdout.write(self.style.SUCCESS('\nImportación completada (transacción confirmada).'))
        total_filas = sum(c + a for _, c, a, _ in resumen)
        self.stdout.write(self.style.SUCCESS(f'{total_filas} filas procesadas en {len(resumen)} tablas.'))
        self.stdout.write(self.style.WARNING(
            'Nota: las contraseñas se importaron como inutilizables; usa el superusuario '
            'de prueba o resetea contraseñas. Las M2M (materias de profesor, colegios '
            'implicados en alertas) no vienen en el backup.'
        ))
